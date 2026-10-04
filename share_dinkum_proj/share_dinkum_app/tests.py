"""Test suite for share_dinkum_app. Run with: python manage.py test share_dinkum_app"""
import io
import json
import shutil
import sqlite3
import pickle
import tempfile
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch, MagicMock
from django.utils import timezone

import openpyxl
import pandas as pd

from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase, TransactionTestCase
from django.urls import reverse
from django.db import IntegrityError, transaction
from djmoney.money import Money

from share_dinkum_app.constants import DEFAULT_CURRENCY, CGT_DISCOUNT_RATE, CGT_DISCOUNT_THRESHOLD_DAYS
from share_dinkum_app.models import (
    AppUser,
    FiscalYearType,
    FiscalYear,
    Account,
    Market,
    Instrument,
    ExchangeRate,
    CurrentExchangeRate,
    Buy,
    Sell,
    Parcel,
    SellAllocation,
    ShareSplit,
    CostBaseAdjustment,
    CostBaseAdjustmentAllocation,
    LogEntry,
    InstrumentPriceHistory,
    Dividend,
    Distribution,
    DataExport,
    CGTReturnSnapshot,
    CGTReturnSnapshotRow,
    AttributionStatement,
    AttributionComponent,
    ResidencyPeriod,
    CPIIndex,
    InstrumentValuation,
    CapitalLossCarryForward,
    LodgedSnapshotError,
)
from share_dinkum_app.utils.currency import add_currencies
from share_dinkum_app.utils.filefield_operations import user_directory_path, process_filefield
from share_dinkum_app.decorators import safe_property
from share_dinkum_app.reports import (
    RealisedCapitalGainReport,
    OpenParcelReport,
    CGTBasisChangeReport,
    CGTEventReport,
    CGTScheduleReport,
)
from django.apps import apps
from share_dinkum_app import choices
from share_dinkum_app.choices import CGTAssetCategory
from django.contrib import admin
from share_dinkum_app import (
    cgt, column_help, constants, dashboard, excelinterface, loading, reports, version, yfinanceinterface)
from share_dinkum_app.management.commands import make_fake_data, make_import_template


# --- Test data factories (minimal objects for isolation) ---


def create_fiscal_year_type(description='Australian Tax Year', start_month=7, start_day=1):
    return FiscalYearType.objects.create(
        description=description,
        start_month=start_month,
        start_day=start_day,
    )


def create_user(username='testuser', password='testpass123'):
    return AppUser.objects.create_user(username=username, password=password)


def create_account(owner=None, currency=DEFAULT_CURRENCY, description='Test Account',
                   fy_type=None):
    if owner is None:
        owner = create_user()
    if fy_type is None:
        fy_type = create_fiscal_year_type()
    return Account.objects.create(
        owner=owner,
        description=description,
        currency=currency,
        fiscal_year_type=fy_type,
    )


def create_market(account=None, code='ASX', suffix='AX'):
    if account is None:
        account = create_account()
    return Market.objects.create(account=account, code=code, suffix=suffix)


def create_instrument(account=None, market=None, name='BHP', currency=DEFAULT_CURRENCY):
    if account is None:
        account = create_account()
    if market is None:
        market = create_market(account=account)
    return Instrument.objects.create(
        account=account,
        market=market,
        name=name,
        description=f'{name} description',
        currency=currency,
    )


def create_exchange_rate(account, convert_from, convert_to, rate=Decimal('1.5'), exchange_date=None):
    if exchange_date is None:
        exchange_date = date.today()
    return ExchangeRate.objects.create(
        account=account,
        convert_from=convert_from,
        convert_to=convert_to,
        date=exchange_date,
        exchange_rate_multiplier=rate,
    )


# =============================================================================
# Utils: currency
# =============================================================================


class AddCurrenciesTests(TestCase):
    """Tests for share_dinkum_app.utils.currency.add_currencies."""

    def test_empty_returns_zero_in_default_currency(self):
        result = add_currencies()
        self.assertEqual(result.amount, 0)
        self.assertEqual(str(result.currency), DEFAULT_CURRENCY)

    def test_single_zero_returns_zero(self):
        result = add_currencies(Money(0, 'AUD'))
        self.assertEqual(result.amount, 0)
        self.assertEqual(str(result.currency), DEFAULT_CURRENCY)

    def test_multiple_zeros_returns_zero(self):
        result = add_currencies(Money(0, 'AUD'), Money(0, 'USD'))
        self.assertEqual(result.amount, 0)

    def test_single_nonzero_returns_same(self):
        result = add_currencies(Money(100, 'AUD'))
        self.assertEqual(result.amount, 100)
        self.assertEqual(str(result.currency), 'AUD')

    def test_same_currency_sums(self):
        result = add_currencies(Money(10, 'AUD'), Money(20, 'AUD'), Money(5, 'AUD'))
        self.assertEqual(result.amount, 35)
        self.assertEqual(str(result.currency), 'AUD')

    def test_ignores_zero_amounts_in_mix(self):
        result = add_currencies(Money(0, 'AUD'), Money(10, 'AUD'), Money(0, 'AUD'))
        self.assertEqual(result.amount, 10)

    def test_different_currencies_raises(self):
        with self.assertRaises(ValueError) as ctx:
            add_currencies(Money(10, 'AUD'), Money(20, 'USD'))
        self.assertIn('Cannot add different currencies', str(ctx.exception))

    def test_non_money_raises(self):
        with self.assertRaises(TypeError) as ctx:
            add_currencies(10, Money(20, 'AUD'))
        self.assertIn('Expected Money', str(ctx.exception))


# =============================================================================
# Utils: filefield_operations
# =============================================================================


class UserDirectoryPathTests(TestCase):
    """Tests for user_directory_path upload_to helper."""

    def test_with_account_and_date(self):
        acc = create_account()
        obj = type('Obj', (), {})()
        obj.account = acc
        obj.date = date(2024, 6, 15)
        path = user_directory_path(obj, 'statement.pdf')
        self.assertIn(str(acc.id), path)
        self.assertIn('2024-06-15', path)
        self.assertIn('statement.pdf', path)

    def test_with_instrument(self):
        acc = create_account()
        obj = type('Obj', (), {})()
        obj.account = acc
        obj.instrument = type('Inst', (), {'name': 'BHP'})()
        obj.date = date(2024, 1, 1)
        path = user_directory_path(obj, 'file.xlsx')
        self.assertIn('BHP', path)


class ProcessFilefieldTests(TestCase):
    """Tests for process_filefield (minimal: no files on disk)."""

    def test_none_returns_none(self):
        self.assertIsNone(process_filefield(None))

    def test_empty_string_returns_none(self):
        self.assertIsNone(process_filefield(''))


# =============================================================================
# yfinanceinterface
# =============================================================================


class ToSnakeCaseTests(TestCase):
    """Tests for yfinanceinterface.to_snake_case."""

    def test_lowercase_unchanged(self):
        self.assertEqual(yfinanceinterface.to_snake_case('open'), 'open')

    def test_camel_case_converted(self):
        # to_snake_case only replaces non-alphanumeric and lowercases; it does not split CamelCase
        self.assertEqual(yfinanceinterface.to_snake_case('StockSplits'), 'stocksplits')

    def test_mixed_case_with_space_becomes_underscore(self):
        self.assertEqual(yfinanceinterface.to_snake_case('Stock Splits'), 'stock_splits')

    def test_non_alnum_replaced_with_underscore(self):
        self.assertEqual(yfinanceinterface.to_snake_case('Close'), 'close')
        self.assertEqual(yfinanceinterface.to_snake_case('High'), 'high')


@patch('share_dinkum_app.yfinanceinterface.yf')
class GetExchangeRateTests(TestCase):
    """Tests for yfinanceinterface.get_exchange_rate."""

    def test_returns_decimal_when_history_has_data(self, mock_yf):
        mock_ticker = MagicMock()
        mock_yf.Ticker.return_value = mock_ticker
        mock_ticker.history.return_value = pd.DataFrame({'Close': [1.55]})
        result = yfinanceinterface.get_exchange_rate('USD', 'AUD', exchange_date=date(2024, 1, 15))
        # Decimal(float) can have float noise; compare quantized or as float
        self.assertEqual(result.quantize(Decimal('0.01')), Decimal('1.55'))
        mock_ticker.history.assert_called_once()

    def test_returns_none_on_exception(self, mock_yf):
        mock_ticker = MagicMock()
        mock_yf.Ticker.return_value = mock_ticker
        mock_ticker.history.side_effect = Exception('network error')
        result = yfinanceinterface.get_exchange_rate('USD', 'AUD', exchange_date=date(2024, 1, 15))
        self.assertIsNone(result)

    def test_returns_none_when_history_empty(self, mock_yf):
        mock_ticker = MagicMock()
        mock_yf.Ticker.return_value = mock_ticker
        mock_ticker.history.return_value = pd.DataFrame({'Close': []})
        result = yfinanceinterface.get_exchange_rate('USD', 'AUD', exchange_date=date(2024, 1, 15))
        self.assertIsNone(result)

    def test_uses_today_when_exchange_date_none(self, mock_yf):
        mock_ticker = MagicMock()
        mock_yf.Ticker.return_value = mock_ticker
        mock_ticker.history.return_value = pd.DataFrame({'Close': [2.0]})
        result = yfinanceinterface.get_exchange_rate('USD', 'AUD', exchange_date=None)
        self.assertEqual(result.quantize(Decimal('0.01')), Decimal('2.00'))


@patch('share_dinkum_app.yfinanceinterface.yf')
class GetExchangeRateHistoryTests(TestCase):
    """Tests for yfinanceinterface.get_exchange_rate_history."""

    def test_returns_dataframe_with_expected_columns_on_success(self, mock_yf):
        mock_ticker = MagicMock()
        mock_yf.Ticker.return_value = mock_ticker
        # history() returns DataFrame with Date index; code resets index and renames cols to snake_case
        df = pd.DataFrame({
            'Open': [1.0], 'High': [1.1], 'Low': [0.9], 'Close': [1.05],
            'Volume': [1000], 'Stock Splits': [0],
        }, index=pd.DatetimeIndex([pd.Timestamp('2024-01-15')]))
        df.index.name = 'Date'
        mock_ticker.history.return_value = df.copy()
        result = yfinanceinterface.get_exchange_rate_history('USD', 'AUD', start_date=date(2024, 1, 1))
        self.assertFalse(result.empty)
        self.assertIn('convert_from', result.columns)
        self.assertIn('convert_to', result.columns)
        self.assertIn('date', result.columns)
        self.assertIn('exchange_rate_multiplier', result.columns)
        self.assertEqual(result['convert_from'].iloc[0], 'USD')
        self.assertEqual(result['convert_to'].iloc[0], 'AUD')

    def test_returns_empty_dataframe_on_exception(self, mock_yf):
        mock_ticker = MagicMock()
        mock_yf.Ticker.return_value = mock_ticker
        mock_ticker.history.side_effect = Exception('api error')
        result = yfinanceinterface.get_exchange_rate_history('USD', 'AUD', start_date=date(2024, 1, 1))
        self.assertTrue(result.empty)
        self.assertIsInstance(result, pd.DataFrame)


@patch('share_dinkum_app.yfinanceinterface.yf')
class GetInstrumentPriceHistoryTests(TestCase):
    """Tests for yfinanceinterface.get_instrument_price_history."""

    def test_returns_dataframe_with_expected_columns_on_success(self, mock_yf):
        mock_ticker = MagicMock()
        mock_yf.Ticker.return_value = mock_ticker
        instrument = MagicMock()
        instrument.yfinance_ticker_code = 'BHP.AX'
        df = pd.DataFrame({
            'Open': [50.0], 'High': [51.0], 'Low': [49.0], 'Close': [50.5],
            'Volume': [1000000], 'Stock Splits': [0],
        }, index=pd.DatetimeIndex([pd.Timestamp('2024-01-15')]))
        df.index.name = 'Date'
        mock_ticker.history.return_value = df.copy()
        result = yfinanceinterface.get_instrument_price_history(instrument, start_date=date(2024, 1, 1))
        self.assertFalse(result.empty)
        self.assertIn('instrument', result.columns)
        self.assertIn('date', result.columns)
        self.assertIn('close', result.columns)
        self.assertIn('open', result.columns)
        self.assertIn('volume', result.columns)
        self.assertIn('stock_splits', result.columns)
        mock_ticker.history.assert_called_once_with(start='2024-01-01', auto_adjust=False)


    def test_returns_empty_dataframe_on_exception(self, mock_yf):
        mock_ticker = MagicMock()
        mock_yf.Ticker.return_value = mock_ticker
        mock_ticker.history.side_effect = Exception('api error')
        instrument = MagicMock()
        instrument.yfinance_ticker_code = 'BHP.AX'
        result = yfinanceinterface.get_instrument_price_history(instrument, start_date=date(2024, 1, 1))
        self.assertTrue(result.empty)
        self.assertIsInstance(result, pd.DataFrame)
        mock_ticker.history.assert_called_once_with(start='2024-01-01', auto_adjust=False)



    def test_end_date_is_inclusive(self, mock_yf):
        """Fetched to today, so later splits can be undone, then cut at the end date."""
        mock_ticker = MagicMock()
        mock_yf.Ticker.return_value = mock_ticker
        instrument = MagicMock()
        instrument.yfinance_ticker_code = 'BHP.AX'
        df = pd.DataFrame({
            'Open': [50.0, 25.0], 'High': [51.0, 26.0], 'Low': [49.0, 24.0],
            'Close': [50.5, 25.5], 'Volume': [1000000, 2000000], 'Stock Splits': [0, 2.0],
        }, index=pd.DatetimeIndex([pd.Timestamp('2024-01-31'), pd.Timestamp('2024-03-01')]))
        df.index.name = 'Date'
        mock_ticker.history.return_value = df.copy()
        result = yfinanceinterface.get_instrument_price_history(
            instrument,
            start_date=date(2024, 1, 1),
            end_date=date(2024, 1, 31),
        )
        self.assertEqual(list(result['date']), [date(2024, 1, 31)])
        # Yahoo had already halved it for the 2-for-1 split in March.
        self.assertEqual(result['close'].iloc[0], Decimal('101'))
        mock_ticker.history.assert_called_once_with(
            start='2024-01-01', end=(date.today() + timedelta(days=1)).isoformat(),
            auto_adjust=False)

    def test_an_end_date_of_today_is_forwarded(self, mock_yf):
        mock_ticker = MagicMock()
        mock_yf.Ticker.return_value = mock_ticker
        instrument = MagicMock()
        instrument.yfinance_ticker_code = 'BHP.AX'
        mock_ticker.history.return_value = pd.DataFrame()
        yfinanceinterface.get_instrument_price_history(
            instrument, start_date=date(2024, 1, 1), end_date=date.today())
        mock_ticker.history.assert_called_once_with(
            start='2024-01-01', end=(date.today() + timedelta(days=1)).isoformat(),
            auto_adjust=False)

    def test_prices_are_stored_as_traded(self, mock_yf):
        """Yahoo divides earlier prices by later splits; each is multiplied back.

        The rows are from NVDA around its 10-for-1 split on 10 June 2024, as Yahoo returns them
        unadjusted: the 7 June close comes back as 120.888, not the 1,208.88 it traded at.
        """
        mock_ticker = MagicMock()
        mock_yf.Ticker.return_value = mock_ticker
        instrument = MagicMock()
        instrument.yfinance_ticker_code = 'NVDA'
        df = pd.DataFrame({
            'Open': [121.0, 120.0], 'High': [122.0, 122.5], 'Low': [119.0, 119.5],
            'Close': [120.888, 121.79], 'Adj Close': [120.54, 121.44],
            'Volume': [412386000, 314162700], 'Dividends': [0.0, 0.0],
            'Stock Splits': [0.0, 10.0],
        }, index=pd.DatetimeIndex([pd.Timestamp('2024-06-07'), pd.Timestamp('2024-06-10')]))
        df.index.name = 'Date'
        mock_ticker.history.return_value = df.copy()

        result = yfinanceinterface.get_instrument_price_history(
            instrument, start_date=date(2024, 6, 7))

        self.assertEqual(list(result['close']), [Decimal('1208.88'), Decimal('121.79')])
        self.assertEqual(list(result['volume']), [41238600, 314162700])
        self.assertNotIn('adj_close', result.columns)


@patch('share_dinkum_app.yfinanceinterface.yf')
class GetCurrentPriceTests(TestCase):
    """Tests for yfinanceinterface.get_current_price."""

    def test_returns_decimal_when_current_price_exists(self, mock_yf):
        mock_ticker = MagicMock()
        mock_yf.Ticker.return_value = mock_ticker
        mock_ticker.info = {'currentPrice': 150.25}
        instrument = MagicMock()
        instrument.yfinance_ticker_code = 'BHP.AX'
        result = yfinanceinterface.get_current_price(instrument)
        self.assertEqual(result.quantize(Decimal('0.01')), Decimal('150.25'))

    def test_falls_back_to_regular_market_price(self, mock_yf):
        mock_ticker = MagicMock()
        mock_yf.Ticker.return_value = mock_ticker
        mock_ticker.info = {'regularMarketPrice': 148.50}
        instrument = MagicMock()
        instrument.yfinance_ticker_code = 'BHP.AX'
        result = yfinanceinterface.get_current_price(instrument)
        self.assertEqual(result.quantize(Decimal('0.01')), Decimal('148.50'))

    def test_prefers_current_price_over_regular_market_price(self, mock_yf):
        mock_ticker = MagicMock()
        mock_yf.Ticker.return_value = mock_ticker
        mock_ticker.info = {'currentPrice': 150.25, 'regularMarketPrice': 148.50}
        instrument = MagicMock()
        instrument.yfinance_ticker_code = 'BHP.AX'
        result = yfinanceinterface.get_current_price(instrument)
        self.assertEqual(result.quantize(Decimal('0.01')), Decimal('150.25'))

    def test_returns_none_when_no_price_available(self, mock_yf):
        mock_ticker = MagicMock()
        mock_yf.Ticker.return_value = mock_ticker
        mock_ticker.info = {}
        instrument = MagicMock()
        instrument.yfinance_ticker_code = 'BHP.AX'
        result = yfinanceinterface.get_current_price(instrument)
        self.assertIsNone(result)

    def test_returns_none_on_exception(self, mock_yf):
        mock_ticker = MagicMock()
        mock_yf.Ticker.return_value = mock_ticker
        mock_ticker.info = property(lambda self: (_ for _ in ()).throw(Exception('network error')))
        type(mock_ticker).info = property(lambda self: (_ for _ in ()).throw(Exception('network error')))
        instrument = MagicMock()
        instrument.yfinance_ticker_code = 'BHP.AX'
        result = yfinanceinterface.get_current_price(instrument)
        self.assertIsNone(result)


# =============================================================================
# Decorators
# =============================================================================



class SafePropertyTests(TestCase):
    """Tests for @safe_property decorator."""

    def test_returns_none_when_adding(self):
        class Model:
            _state = type('State', (), {'adding': True})()
            def get_val(self):
                return 42
        Model.get_val = safe_property(Model.get_val)
        m = Model()
        self.assertIsNone(m.get_val)

    def test_returns_value_when_not_adding(self):
        class Model:
            _state = type('State', (), {'adding': False})()
            def get_val(self):
                return 42
        Model.get_val = safe_property(Model.get_val)
        m = Model()
        self.assertEqual(m.get_val, 42)

    def test_fget_has_safe_property_marker(self):
        def fn(self):
            return 1
        wrapped = safe_property(fn)
        self.assertTrue(getattr(wrapped.fget, '_is_safe_property', False))


# =============================================================================
# Models: FiscalYearType & FiscalYear
# =============================================================================


class FiscalYearTypeTests(TestCase):
    """Tests for FiscalYearType model."""

    def test_classify_date_same_calendar_year(self):
        fy_type = create_fiscal_year_type(start_month=7, start_day=1)
        fy, created = fy_type.classify_date(date(2024, 8, 1))
        self.assertTrue(created)
        self.assertEqual(fy.start_year, 2024)
        self.assertEqual(fy.fiscal_year_type, fy_type)

    def test_classify_date_previous_calendar_year(self):
        fy_type = create_fiscal_year_type(start_month=7, start_day=1)
        fy, created = fy_type.classify_date(date(2024, 3, 1))
        self.assertTrue(created)
        self.assertEqual(fy.start_year, 2023)

    def test_classify_date_get_or_create_reuse(self):
        fy_type = create_fiscal_year_type()
        fy1, c1 = fy_type.classify_date(date(2024, 8, 1))
        fy2, c2 = fy_type.classify_date(date(2024, 9, 1))
        self.assertFalse(c2)
        self.assertEqual(fy1.id, fy2.id)

    def test_str(self):
        fy_type = create_fiscal_year_type(description='AU Tax')
        self.assertIn('AU Tax', str(fy_type))

    def test_classify_date_boundaries_default_parameters(self):
        fy_type = FiscalYearType.objects.create(
            description='Australian Tax Year',
            start_month=7,
            start_day=1,
        )

        scenarios = [
            (date(2024, 6, 30), 2023),
            (date(2024, 7, 1), 2024),
            (date(2024, 12, 31), 2024),
            (date(2025, 1, 1), 2024),
        ]

        encountered_years = {}

        for input_date, expected_year in scenarios:
            fiscal_year, _ = fy_type.classify_date(input_date)
            self.assertEqual(
                fiscal_year.start_year,
                expected_year,
                f"{input_date} should map to fiscal year starting {expected_year}",
            )
            if expected_year in encountered_years:
                self.assertEqual(encountered_years[expected_year], fiscal_year.id)
            else:
                encountered_years[expected_year] = fiscal_year.id

    def test_classify_date_calendar_year_boundaries(self):
        fy_type = FiscalYearType.objects.create(
            description='Calendar Year',
            start_month=1,
            start_day=1,
        )

        scenarios = [
            (date(2023, 12, 31), 2023),
            (date(2024, 1, 1), 2024),
            (date(2024, 2, 29), 2024),
            (date(2024, 12, 31), 2024),
            (date(2025, 1, 1), 2025),
        ]

        encountered_years = {}

        for input_date, expected_year in scenarios:
            fiscal_year, _ = fy_type.classify_date(input_date)
            self.assertEqual(
                fiscal_year.start_year,
                expected_year,
                f"{input_date} should map to fiscal year starting {expected_year}",
            )
            if expected_year in encountered_years:
                self.assertEqual(encountered_years[expected_year], fiscal_year.id)
            else:
                encountered_years[expected_year] = fiscal_year.id

        fy_mid_year, _ = fy_type.classify_date(date(2024, 6, 30))
        fy_late_year, _ = fy_type.classify_date(date(2024, 11, 15))
        self.assertEqual(fy_mid_year.id, fy_late_year.id)



class FiscalYearTests(TestCase):
    """Tests for FiscalYear model."""

    def test_start_date_end_date_july_year(self):
        fy_type = create_fiscal_year_type(start_month=7, start_day=1)
        fy = FiscalYear.objects.create(fiscal_year_type=fy_type, start_year=2024)
        self.assertEqual(fy.start_date, date(2024, 7, 1))
        self.assertEqual(fy.end_date, date(2025, 6, 30))

    def test_get_name_financial_year(self):
        fy_type = create_fiscal_year_type(start_month=7, start_day=1)
        fy = FiscalYear.objects.create(fiscal_year_type=fy_type, start_year=2024)
        self.assertEqual(fy.get_name(), 'FY2024/25')

    def test_get_name_calendar_year(self):
        fy_type = FiscalYearType.objects.create(
            description='Calendar',
            start_month=1,
            start_day=1,
        )
        fy = FiscalYear.objects.create(fiscal_year_type=fy_type, start_year=2024)
        self.assertEqual(fy.get_name(), '2024')

    def test_calendar_year_start_and_end_dates_cover_full_year(self):
        fy_type = FiscalYearType.objects.create(
            description='Calendar Year',
            start_month=1,
            start_day=1,
        )
        fy = FiscalYear.objects.create(fiscal_year_type=fy_type, start_year=2024)

        self.assertEqual(fy.start_date, date(2024, 1, 1))
        self.assertEqual(fy.end_date, date(2024, 12, 31))
        self.assertEqual(
            (fy.end_date - fy.start_date).days + 1,
            366,
            "Leap year should span all 366 days",
        )

        fy_non_leap = FiscalYear.objects.create(fiscal_year_type=fy_type, start_year=2023)
        self.assertEqual(fy_non_leap.start_date, date(2023, 1, 1))
        self.assertEqual(fy_non_leap.end_date, date(2023, 12, 31))
        self.assertEqual(
            (fy_non_leap.end_date - fy_non_leap.start_date).days + 1,
            365,
            "Non-leap fiscal year should span 365 days",
        )




# =============================================================================
# Models: AppUser, Account
# =============================================================================


class AppUserTests(TestCase):
    """Tests for AppUser model."""

    def test_save_sets_blank_first_last_name(self):
        user = AppUser(username='u', email='u@test.com')
        user.set_password('x')
        user.save()
        self.assertEqual(user.first_name, '')
        self.assertEqual(user.last_name, '')


class AccountTests(TestCase):
    """Tests for Account model."""

    def test_str(self):
        acc = create_account()
        self.assertIn(acc.description, str(acc))
        self.assertIn(acc.currency, str(acc))

    def test_portfolio_value_converted_empty_is_zero(self):
        acc = create_account()
        self.assertEqual(acc.portfolio_value_converted.amount, 0)
        self.assertEqual(str(acc.portfolio_value_converted.currency), acc.currency)


# =============================================================================
# Models: Exchange rates (with mocked yfinance)
# =============================================================================


@patch('share_dinkum_app.models.yfinanceinterface.get_exchange_rate')
class ExchangeRateTests(TestCase):
    """Tests for ExchangeRate and AbstractExchangeRate (with yfinance mocked)."""

    def test_apply_same_currency(self, mock_get_rate):
        acc = create_account()
        rate = create_exchange_rate(acc, 'AUD', 'AUD', rate=Decimal('1'))
        m = Money(100, 'AUD')
        result = rate.apply(m)
        self.assertEqual(result.amount, 100)
        self.assertEqual(str(result.currency), 'AUD')

    def test_apply_converts(self, mock_get_rate):
        acc = create_account()
        rate = create_exchange_rate(acc, 'USD', 'AUD', rate=Decimal('1.5'))
        m = Money(100, 'USD')
        result = rate.apply(m)
        self.assertEqual(result.amount, Decimal('150'))
        self.assertEqual(str(result.currency), 'AUD')

    def test_apply_wrong_currency_raises(self, mock_get_rate):
        acc = create_account()
        rate = create_exchange_rate(acc, 'USD', 'AUD', rate=Decimal('1.5'))
        m = Money(100, 'AUD')
        with self.assertRaises(AssertionError):
            rate.apply(m)

    def test_update_current_creates_current_rate(self, mock_get_rate):
        acc = create_account()
        hist = create_exchange_rate(acc, 'USD', 'AUD', rate=Decimal('1.6'))
        current = hist.update_current()
        self.assertIsNotNone(current)
        self.assertEqual(current.exchange_rate_multiplier, Decimal('1.6'))
        self.assertEqual(CurrentExchangeRate.objects.filter(
            account=acc, convert_from='USD', convert_to='AUD'
        ).count(), 1)

    def test_get_or_create_creates_with_mock_rate(self, mock_get_rate):
        mock_get_rate.return_value = Decimal('1.55')
        acc = create_account()
        rate = ExchangeRate.get_or_create(
            account=acc,
            convert_from='USD',
            convert_to='AUD',
            exchange_date=date(2024, 1, 15),
        )
        self.assertEqual(rate.exchange_rate_multiplier, Decimal('1.55'))
        mock_get_rate.assert_called_once()

    def test_a_rate_that_could_not_be_fetched_is_fetched_again_later(self, mock_get_rate):
        acc = create_account()
        mock_get_rate.return_value = None
        stand_in = ExchangeRate.get_or_create(
            account=acc, convert_from='USD', convert_to='AUD', exchange_date=date(2015, 3, 2))
        self.assertTrue(stand_in.is_placeholder)

        mock_get_rate.return_value = Decimal('1.3')
        rate = ExchangeRate.get_or_create(
            account=acc, convert_from='USD', convert_to='AUD', exchange_date=date(2015, 3, 2))

        self.assertEqual(rate.pk, stand_in.pk)
        self.assertEqual(rate.exchange_rate_multiplier, Decimal('1.3'))
        self.assertFalse(rate.is_placeholder)

    def test_a_stand_in_rate_comes_from_the_nearest_earlier_date(self, mock_get_rate):
        acc = create_account()
        create_exchange_rate(acc, 'USD', 'AUD', rate=Decimal('1.1'), exchange_date=date(2015, 1, 2))
        create_exchange_rate(acc, 'USD', 'AUD', rate=Decimal('1.6'), exchange_date=date(2025, 1, 2))
        mock_get_rate.return_value = None

        rate = ExchangeRate.get_or_create(
            account=acc, convert_from='USD', convert_to='AUD', exchange_date=date(2015, 3, 2))

        self.assertEqual(rate.exchange_rate_multiplier, Decimal('1.1'))
        self.assertTrue(rate.is_placeholder)

    def test_the_history_refresh_replaces_a_stand_in_rate(self, mock_get_rate):
        acc = create_account()
        mock_get_rate.return_value = Decimal('1.5')
        inst = create_instrument(account=acc, name='AAPL', currency='USD')
        mock_get_rate.return_value = None
        buy = Buy.objects.create(
            account=acc, instrument=inst, date=date(2024, 1, 15), quantity=Decimal('10'),
            unit_price=Money(100, 'USD'), total_brokerage=Money(0, 'USD'),
        )
        self.assertTrue(buy.exchange_rate.is_placeholder)

        history = pd.DataFrame([{
            'convert_from': 'USD', 'convert_to': 'AUD', 'date': date(2024, 1, 15),
            'exchange_rate_multiplier': Decimal('1.4'), 'is_continuous_history': True,
        }])
        with patch('share_dinkum_app.models.yfinanceinterface.get_exchange_rate_history',
                   return_value=history):
            ExchangeRate.update_exchange_rate_history(
                account=acc, convert_from='USD', convert_to='AUD')

        buy.refresh_from_db()
        self.assertFalse(buy.exchange_rate.is_placeholder)
        self.assertEqual(buy.exchange_rate.exchange_rate_multiplier, Decimal('1.4'))
        self.assertEqual(buy.calculated_unit_price_converted.amount, Decimal('140'))
        parcel = Parcel.objects.get(buy=buy)
        self.assertEqual(parcel.calculated_total_cost_base.amount, Decimal('1400'))


# =============================================================================
# Models: Market, Instrument
# =============================================================================


class MarketTests(TestCase):
    """Tests for Market model."""

    def test_unique_per_account_code(self):
        acc = create_account()
        Market.objects.create(account=acc, code='ASX')
        with self.assertRaises(IntegrityError):
            Market.objects.create(account=acc, code='ASX')


class InstrumentTests(TestCase):
    """Tests for Instrument model."""

    def test_yfinance_ticker_code_with_suffix(self):
        acc = create_account()
        market = Market.objects.create(account=acc, code='ASX', suffix='.AX')
        inst = Instrument.objects.create(
            account=acc, market=market, name='BHP',
            description='BHP', currency='AUD',
        )
        self.assertEqual(inst.yfinance_ticker_code, 'BHP.AX')

    def test_yfinance_ticker_code_no_suffix(self):
        acc = create_account()
        market = Market.objects.create(account=acc, code='NASDAQ', suffix='')
        inst = Instrument.objects.create(
            account=acc, market=market, name='AAPL',
            description='Apple', currency='USD',
        )
        self.assertEqual(inst.yfinance_ticker_code, 'AAPL')

    def test_quantity_held_no_trades_is_zero(self):
        inst = create_instrument()
        self.assertEqual(inst.quantity_held, 0)

    def test_value_held_no_price_is_zero(self):
        inst = create_instrument()
        self.assertEqual(inst.value_held.amount, 0)

    @patch('share_dinkum_app.models.yfinanceinterface.get_exchange_rate', return_value=None)
    def test_a_foreign_instrument_can_be_added_offline(self, mock_get_rate):
        """With no rate to be had, its converted value is unknown rather than an error."""
        inst = create_instrument(name='AAPL', currency='USD')
        self.assertIsNone(inst.value_held_converted)


# =============================================================================
# Models: Buy, Sell, Parcel, SellAllocation (with signals)
# =============================================================================


class BuyAndParcelSignalsTests(TransactionTestCase):
    """Test Buy creation creates Parcel via signal."""

    def test_create_buy_creates_parcel(self):
        acc = create_account()
        market = create_market(account=acc)
        inst = create_instrument(account=acc, market=market)
        buy = Buy.objects.create(
            account=acc,
            instrument=inst,
            date=date(2024, 1, 10),
            quantity=Decimal('100'),
            unit_price=Money(50, 'AUD'),
            total_brokerage=Money(10, 'AUD'),
        )
        parcels = Parcel.objects.filter(buy=buy)
        self.assertEqual(parcels.count(), 1)
        self.assertEqual(parcels.first().parcel_quantity, Decimal('100'))


@patch('share_dinkum_app.models.yfinanceinterface.get_exchange_rate', return_value=Decimal('1.5'))
class ForeignCurrencyTradeTests(TransactionTestCase):
    """A foreign-currency trade has its rate before anything built from it is worked out.

    The parcel and the sell allocations are created by post_save signals. If the rate is
    attached after them, they see the trade unconverted: the parcel stores its cost base in
    USD, and a MIN_CGT sell subtracts an AUD cost base from USD proceeds.
    """

    def _usd_buy(self, account, instrument):
        return Buy.objects.create(
            account=account, instrument=instrument, date=date(2024, 1, 10),
            quantity=Decimal('100'), unit_price=Money(Decimal('50'), 'USD'),
            total_brokerage=Money(Decimal('10'), 'USD'),
        )

    def test_buy_parcel_cost_base_is_converted(self, mock_get_rate):
        account = create_account()
        instrument = create_instrument(account=account, name='IBIT', currency='USD')
        buy = self._usd_buy(account, instrument)

        self.assertIsNotNone(buy.exchange_rate)
        parcel = Parcel.objects.get(buy=buy)
        self.assertEqual(str(parcel.calculated_total_cost_base.currency), 'AUD')
        # (100 x 50 + 10) USD x 1.5
        self.assertEqual(parcel.calculated_total_cost_base.amount, Decimal('7515'))

    def test_min_cgt_sell_allocates_in_account_currency(self, mock_get_rate):
        account = create_account()
        instrument = create_instrument(account=account, name='IBIT', currency='USD')
        self._usd_buy(account, instrument)

        sell = Sell.objects.create(
            account=account, instrument=instrument, date=date(2025, 3, 10),
            quantity=Decimal('40'), unit_price=Money(Decimal('60'), 'USD'),
            total_brokerage=Money(Decimal('10'), 'USD'), strategy='MIN_CGT',
        )

        allocation = sell.sale_allocation.get(is_active=True)
        self.assertEqual(allocation.quantity, Decimal('40'))
        self.assertEqual(str(allocation.calculated_total_capital_gain.currency), 'AUD')


@patch('share_dinkum_app.models.yfinanceinterface.get_exchange_rate', return_value=Decimal('1.5'))
class MixedCurrencyTradeTests(TransactionTestCase):
    """Price and brokerage are each converted by their own currency's rate.

    An Australian broker commonly charges AUD brokerage on a USD trade. Converting the
    brokerage at the price's rate failed on the currency mismatch.
    """

    def setUp(self):
        self.account = create_account()
        self.instrument = create_instrument(account=self.account, name='IBIT', currency='USD')

    def _buy(self, price_currency, brokerage_currency):
        return Buy.objects.create(
            account=self.account, instrument=self.instrument, date=date(2024, 1, 10),
            quantity=Decimal('100'), unit_price=Money(Decimal('50'), price_currency),
            total_brokerage=Money(Decimal('10'), brokerage_currency),
        )

    def _cost_base(self, buy):
        return Parcel.objects.get(buy=buy).calculated_total_cost_base

    def test_a_usd_price_with_aud_brokerage(self, mock_get_rate):
        buy = self._buy('USD', 'AUD')

        self.assertEqual(str(buy.exchange_rate.convert_from), 'USD')
        # 100 x 50 USD x 1.5, plus 10 AUD as it is.
        self.assertEqual(self._cost_base(buy), Money(Decimal('7510'), 'AUD'))
        self.assertEqual(buy.total_brokerage_converted, Money(Decimal('10'), 'AUD'))

    def test_an_aud_price_with_usd_brokerage(self, mock_get_rate):
        buy = self._buy('AUD', 'USD')

        self.assertIsNone(buy.exchange_rate)
        # 100 x 50 AUD, plus 10 USD x 1.5.
        self.assertEqual(self._cost_base(buy), Money(Decimal('5015'), 'AUD'))

    def test_both_in_usd(self, mock_get_rate):
        buy = self._buy('USD', 'USD')
        self.assertEqual(self._cost_base(buy), Money(Decimal('7515'), 'AUD'))

    def test_a_sale_with_aud_brokerage_on_a_usd_price(self, mock_get_rate):
        self._buy('USD', 'USD')
        sell = Sell.objects.create(
            account=self.account, instrument=self.instrument, date=date(2025, 3, 10),
            quantity=Decimal('40'), unit_price=Money(Decimal('60'), 'USD'),
            total_brokerage=Money(Decimal('10'), 'AUD'), strategy='FIFO',
        )
        # 40 x 60 USD x 1.5, less 10 AUD.
        self.assertEqual(sell.proceeds, Money(Decimal('3590'), 'AUD'))

    def test_correcting_the_second_rate_reaches_the_trade(self, mock_get_rate):
        """The brokerage's rate is not linked to the trade, but a change to it is carried."""
        buy = self._buy('AUD', 'USD')
        rate = ExchangeRate.objects.get(
            account=self.account, convert_from='USD', convert_to='AUD', date=buy.date)

        rate.exchange_rate_multiplier = Decimal('1.6')
        rate.save()
        rate.rate_corrected()

        self.assertEqual(self._cost_base(buy), Money(Decimal('5016'), 'AUD'))


class DividendCompanyTaxRateTests(TestCase):
    """Franking credits divide by 1 less the company rate, so 100% is refused."""

    def _dividend(self, rate):
        account = create_account()
        return Dividend(
            account=account, instrument=create_instrument(account=account),
            date=date(2024, 4, 1), quantity=Decimal('100'),
            franked_amount_per_share=Money(Decimal('0.50'), 'AUD'),
            unfranked_amount_per_share=Money(0, 'AUD'),
            corporate_tax_rate_percentage=Decimal(rate),
        )

    def test_a_rate_of_100_is_refused(self):
        dividend = self._dividend('100')
        with self.assertRaisesMessage(ValidationError, 'less than 100'):
            dividend.full_clean()
        with self.assertRaisesMessage(ValueError, 'less than 100'):
            dividend.save()

    def test_a_negative_rate_is_refused(self):
        with self.assertRaisesMessage(ValueError, 'at least 0'):
            self._dividend('-1').save()

    def test_a_rate_below_100_is_accepted(self):
        dividend = self._dividend('25')
        dividend.save()
        # 50 franked at 25%: 50 x 0.25 / 0.75.
        self.assertAlmostEqual(dividend.total_franking_credits.amount, Decimal('16.6667'), places=4)


@patch('share_dinkum_app.models.yfinanceinterface.get_exchange_rate', return_value=Decimal('1.5'))
class RepairPortfolioDataTests(TransactionTestCase):
    """Records left wrong by bugs fixed in 0.3.0: the dashboard warns, the command repairs."""

    def setUp(self):
        from share_dinkum_app.dashboard import _data_check_warning
        self._warning = _data_check_warning

    def _portfolio(self, currency='USD'):
        account = create_account()
        instrument = create_instrument(account=account, name='IBIT', currency=currency)
        buy = Buy.objects.create(
            account=account, instrument=instrument, date=date(2024, 1, 10),
            quantity=Decimal('100'), unit_price=Money(Decimal('50'), currency),
            total_brokerage=Money(Decimal('10'), currency),
        )
        return account, instrument, buy

    def _unconvert(self, buy):
        """Put the buy's parcel back as the bug left it, its cost base labelled USD."""
        Parcel.objects.filter(buy=buy).update(calculated_total_cost_base_currency='USD')

    def _repair(self, *args):
        out = io.StringIO()
        call_command('repair_portfolio_data', *args, stdout=out)
        return out.getvalue()

    def test_a_clean_portfolio_is_left_alone(self, mock_get_rate):
        account, _, _ = self._portfolio()
        self.assertIsNone(self._warning(account))

    def test_an_unconverted_parcel_is_warned_about(self, mock_get_rate):
        account, _, buy = self._portfolio()
        self._unconvert(buy)

        warning = self._warning(account)
        self.assertIn('1 parcel(s) with their cost base stored in a currency', warning)
        self.assertIn('repair_portfolio_data', warning)
        self.assertIn('not affected', warning)

    def test_the_command_recalculates_the_parcel_and_clears_the_warning(self, mock_get_rate):
        account, _, buy = self._portfolio()
        self._unconvert(buy)

        self._repair()

        parcel = Parcel.objects.get(buy=buy)
        self.assertEqual(str(parcel.calculated_total_cost_base.currency), 'AUD')
        self.assertEqual(parcel.calculated_total_cost_base.amount, Decimal('7515'))
        self.assertIsNone(self._warning(account))

    def test_a_stale_sold_flag_is_recalculated(self, mock_get_rate):
        account, instrument, _ = self._portfolio(currency='AUD')
        Sell.objects.create(
            account=account, instrument=instrument, date=date(2024, 6, 1),
            quantity=Decimal('40'), unit_price=Money(60, 'AUD'),
            total_brokerage=Money(0, 'AUD'), strategy='FIFO',
        )
        Parcel.objects.filter(sale_date__isnull=False).update(calculated_is_sold=False)
        self.assertIn('stored sold flag', self._warning(account))

        self._repair()

        self.assertTrue(Parcel.objects.get(sale_date__isnull=False).calculated_is_sold)
        self.assertIsNone(self._warning(account))

    def test_an_adjustment_left_behind_by_a_split_is_carried_forward(self, mock_get_rate):
        """As the bug left it: on the replaced parcel, and on none of its descendants."""
        account, instrument, buy = self._portfolio(currency='AUD')
        CostBaseAdjustment.objects.create(
            account=account, instrument=instrument, financial_year_end_date=date(2024, 6, 30),
            cost_base_increase=Money(Decimal('50'), 'AUD'),
        )
        ShareSplit.objects.create(
            account=account, instrument=instrument, date=date(2024, 9, 1),
            quantity_before=Decimal('1'), quantity_after=Decimal('2'),
        )
        Sell.objects.create(
            account=account, instrument=instrument, date=date(2024, 12, 1),
            quantity=Decimal('50'), unit_price=Money(30, 'AUD'),
            total_brokerage=Money(0, 'AUD'), strategy='FIFO',
        )
        # Undo what split_or_consolidate now does, to recreate what the bug left.
        original = Parcel.objects.get(buy=buy, parent_parcel__isnull=True)
        for allocation in CostBaseAdjustmentAllocation.objects.filter(account=account):
            if allocation.parcel_id == original.pk:
                CostBaseAdjustmentAllocation.objects.filter(pk=allocation.pk).update(
                    deactivation_date=None, is_active=True)
            else:
                CostBaseAdjustmentAllocation.objects.filter(pk=allocation.pk).delete()
        self.assertIn('left behind by a share split', self._warning(account))

        out = self._repair()

        self.assertIn('carried 1 cost base adjustment allocation(s)', out)
        held = Parcel.objects.filter(buy=buy, deactivation_date__isnull=True)
        self.assertEqual(held.count(), 2)  # 50 sold and 150 held, of 200 after the split
        self.assertEqual(sum(p.total_adjustments.amount for p in held), Decimal('50'))
        self.assertIsNone(self._warning(account))

    def test_a_stand_in_rate_is_fetched_again(self, mock_get_rate):
        account, _, buy = self._portfolio()
        ExchangeRate.objects.filter(pk=buy.exchange_rate_id).update(
            is_placeholder=True, exchange_rate_multiplier=Decimal('1'))
        self.assertIn('standing in', self._warning(account))

        out = self._repair()

        self.assertIn('fetched 1 exchange rate(s)', out)
        buy.refresh_from_db()
        self.assertEqual(buy.calculated_unit_price_converted.amount, Decimal('75'))

    def test_a_dry_run_changes_nothing(self, mock_get_rate):
        account, _, buy = self._portfolio()
        self._unconvert(buy)

        out = self._repair('--dry-run')

        self.assertIn('1 parcel(s) with their cost base stored in a currency', out)
        self.assertIn('would repair', out)
        self.assertEqual(Parcel.with_unconverted_cost_base(account).count(), 1)

    def test_sales_needing_a_person_are_listed(self, mock_get_rate):
        account, instrument, _ = self._portfolio(currency='AUD')
        Sell.objects.create(
            account=account, instrument=instrument, date=date(2024, 6, 1),
            quantity=Decimal('150'), unit_price=Money(60, 'AUD'),
            total_brokerage=Money(0, 'AUD'), strategy='FIFO',
        )
        self.assertIn('capital gains are wrong', self._warning(account))

        out = self._repair()

        self.assertIn('IBIT on 2024-06-01: 50 of 150 units', out)

    def test_an_unconverted_adjustment_is_listed_but_not_changed(self, mock_get_rate):
        """Re-allocating would spread the adjustment again, so that is left to the operator."""
        account, instrument, _ = self._portfolio()
        adjustment = CostBaseAdjustment.objects.create(
            account=account, instrument=instrument, financial_year_end_date=date(2024, 6, 30),
            cost_base_increase=Money(Decimal('100'), 'USD'),
        )
        allocations = CostBaseAdjustmentAllocation.objects.filter(cost_base_adjustment=adjustment)
        allocations.update(cost_base_increase_currency='USD')

        self.assertIn(
            '1 cost base adjustment(s) allocated without being converted', self._warning(account))

        out = self._repair()

        self.assertIn('Delete each one and enter it again', out)
        self.assertEqual(
            set(allocations.values_list('cost_base_increase_currency', flat=True)), {'USD'})


class SellAllocationTests(TransactionTestCase):
    """Test Sell with strategy creates allocations and parcel bifurcation."""

    def test_sell_fifo_creates_allocations(self):
        acc = create_account()
        inst = create_instrument(account=acc)
        Buy.objects.create(
            account=acc,
            instrument=inst,
            date=date(2024, 1, 5),
            quantity=Decimal('100'),
            unit_price=Money(50, 'AUD'),
            total_brokerage=Money(0, 'AUD'),
        )
        Buy.objects.create(
            account=acc,
            instrument=inst,
            date=date(2024, 1, 10),
            quantity=Decimal('50'),
            unit_price=Money(52, 'AUD'),
            total_brokerage=Money(0, 'AUD'),
        )
        sell = Sell.objects.create(
            account=acc,
            instrument=inst,
            date=date(2024, 2, 1),
            quantity=Decimal('75'),
            unit_price=Money(55, 'AUD'),
            total_brokerage=Money(0, 'AUD'),
            strategy='FIFO',
        )
        allocations = SellAllocation.objects.filter(sell=sell, is_active=True)
        self.assertGreaterEqual(allocations.count(), 1)
        total_allocated = sum(a.quantity for a in allocations)
        self.assertEqual(total_allocated, Decimal('75'))

    def test_sell_manual_no_auto_allocations(self):
        acc = create_account()
        inst = create_instrument(account=acc)
        Buy.objects.create(
            account=acc,
            instrument=inst,
            date=date(2024, 1, 5),
            quantity=Decimal('100'),
            unit_price=Money(50, 'AUD'),
            total_brokerage=Money(0, 'AUD'),
        )
        sell = Sell.objects.create(
            account=acc,
            instrument=inst,
            date=date(2024, 2, 1),
            quantity=Decimal('50'),
            unit_price=Money(55, 'AUD'),
            total_brokerage=Money(0, 'AUD'),
            strategy='MANUAL',
        )
        allocations = SellAllocation.objects.filter(sell=sell)
        self.assertEqual(allocations.count(), 0)

    def _manual_sale(self, quantity='100', on=date(2024, 3, 1), instrument_name='BHP'):
        account = create_account()
        instrument = create_instrument(account=account, name=instrument_name)
        buy = Buy.objects.create(
            account=account, instrument=instrument, date=date(2024, 1, 5),
            quantity=Decimal('100'), unit_price=Money(50, 'AUD'), total_brokerage=Money(0, 'AUD'),
        )
        sell = Sell.objects.create(
            account=account, instrument=instrument, date=on, quantity=Decimal(quantity),
            unit_price=Money(55, 'AUD'), total_brokerage=Money(0, 'AUD'), strategy='MANUAL',
        )
        return account, instrument, Parcel.objects.get(buy=buy), sell

    def _refused(self, account, parcel, sell, quantity, expected):
        allocation = SellAllocation(
            account=account, parcel=parcel, sell=sell, quantity=Decimal(quantity))
        with self.assertRaisesMessage(ValidationError, expected):
            allocation.full_clean()
        with self.assertRaisesMessage(ValueError, expected):
            SellAllocation.objects.create(
                account=account, parcel=parcel, sell=sell, quantity=Decimal(quantity))

    def test_a_sold_parcel_cannot_be_allocated_again(self):
        account, instrument, parcel, first = self._manual_sale()
        SellAllocation.objects.create(
            account=account, parcel=parcel, sell=first, quantity=Decimal('100'))
        second = Sell.objects.create(
            account=account, instrument=instrument, date=date(2024, 4, 1),
            quantity=Decimal('100'), unit_price=Money(55, 'AUD'),
            total_brokerage=Money(0, 'AUD'), strategy='MANUAL',
        )

        self._refused(account, Parcel.objects.get(pk=parcel.pk), second, '100', 'unsold')
        self.assertEqual(Instrument.objects.get(pk=instrument.pk).quantity_held, Decimal('0'))

    def test_a_parcel_of_another_instrument_is_refused(self):
        account, _, parcel, _ = self._manual_sale()
        other = create_instrument(
            account=account, market=parcel.buy.instrument.market, name='RIO')
        sell = Sell.objects.create(
            account=account, instrument=other, date=date(2024, 3, 1), quantity=Decimal('10'),
            unit_price=Money(55, 'AUD'), total_brokerage=Money(0, 'AUD'), strategy='MANUAL',
        )
        self._refused(account, parcel, sell, '10', 'RIO')

    def test_a_parcel_bought_after_the_sale_is_refused(self):
        account, _, parcel, sell = self._manual_sale(on=date(2024, 1, 2))
        self._refused(account, parcel, sell, '10', 'after the sale')

    def test_allocations_cannot_add_up_to_more_than_the_sale(self):
        account, _, parcel, sell = self._manual_sale(quantity='40')
        self._refused(account, parcel, sell, '50', 'too many')

    def test_the_sold_part_of_a_parcel_is_stored_as_sold(self):
        """The parcel split off for a partial sale records itself as sold, not as held."""
        account, _, parcel, sell = self._manual_sale(quantity='40')
        allocation = SellAllocation.objects.create(
            account=account, parcel=parcel, sell=sell, quantity=Decimal('40'))

        sold = Parcel.objects.get(pk=allocation.parcel_id)
        self.assertTrue(sold.calculated_is_sold)
        self.assertEqual(sold.calculated_remaining_quantity, Decimal('0'))

    def test_deleting_an_allocation_clears_the_parcels_sale_date(self):
        """Otherwise a later adjustment weights the parcel as if it were sold that day."""
        account, _, parcel, sell = self._manual_sale(quantity='40')
        allocation = SellAllocation.objects.create(
            account=account, parcel=parcel, sell=sell, quantity=Decimal('40'))
        sold_id = allocation.parcel_id

        allocation.delete()

        self.assertIsNone(Parcel.objects.get(pk=sold_id).sale_date)


class ParcelTests(TransactionTestCase):
    """Tests for Parcel model methods (bifurcate, split, cost base)."""

    def test_remaining_quantity_after_buy(self):
        acc = create_account()
        inst = create_instrument(account=acc)
        Buy.objects.create(
            account=acc,
            instrument=inst,
            date=date(2024, 1, 5),
            quantity=Decimal('100'),
            unit_price=Money(50, 'AUD'),
            total_brokerage=Money(0, 'AUD'),
        )
        parcel = Parcel.objects.get(buy__instrument=inst)
        self.assertEqual(parcel.remaining_quantity, Decimal('100'))

    def test_split_or_consolidate_multiplier(self):
        acc = create_account()
        inst = create_instrument(account=acc)
        Buy.objects.create(
            account=acc,
            instrument=inst,
            date=date(2024, 1, 5),
            quantity=Decimal('100'),
            unit_price=Money(50, 'AUD'),
            total_brokerage=Money(0, 'AUD'),
        )
        parcel = Parcel.objects.get(buy__instrument=inst)
        new_parcel = parcel.split_or_consolidate(multiplier=Decimal('2'), date=date(2024, 2, 1))
        self.assertEqual(new_parcel.parcel_quantity, Decimal('200'))
        self.assertEqual(new_parcel.cumulative_split_multiplier, Decimal('2'))
        parcel.refresh_from_db()
        self.assertIsNotNone(parcel.deactivation_date)


# =============================================================================
# Models: ShareSplit
# =============================================================================


class ShareSplitTests(TransactionTestCase):
    """Tests for ShareSplit model."""

    def test_split_multiplier(self):
        acc = create_account()
        inst = create_instrument(account=acc)
        ss = ShareSplit.objects.create(
            account=acc,
            instrument=inst,
            quantity_before=Decimal('1'),
            quantity_after=Decimal('3'),
            date=date(2024, 3, 1),
        )
        self.assertEqual(ss.split_multiplier, Decimal('3'))

    def _holding(self, quantity='100', price='10'):
        account = create_account()
        instrument = create_instrument(account=account, name='SPL')
        buy = Buy.objects.create(
            account=account, instrument=instrument, date=date(2021, 8, 1),
            quantity=Decimal(quantity), unit_price=Money(Decimal(price), 'AUD'),
            total_brokerage=Money(0, 'AUD'),
        )
        return account, instrument, buy

    def _split(self, account, instrument, before, after, on=date(2023, 1, 1)):
        return ShareSplit.objects.create(
            account=account, instrument=instrument, date=on,
            quantity_before=Decimal(before), quantity_after=Decimal(after),
        )

    def _active_parcel(self, buy):
        return Parcel.objects.get(buy=buy, deactivation_date__isnull=True)

    def test_a_split_keeps_the_adjustments_already_allocated(self):
        account, instrument, buy = self._holding()
        CostBaseAdjustment.objects.create(
            account=account, instrument=instrument, financial_year_end_date=date(2022, 6, 30),
            cost_base_increase=Money(Decimal('50'), 'AUD'),
        )

        self._split(account, instrument, '1', '2')

        parcel = self._active_parcel(buy)
        self.assertEqual(parcel.parcel_quantity, Decimal('200'))
        self.assertEqual(parcel.total_adjustments.amount, Decimal('50'))
        self.assertEqual(parcel.total_cost_base.amount, Decimal('1050'))

    def test_a_one_for_three_consolidation_is_exact(self):
        account, instrument, buy = self._holding(quantity='3000', price='1')

        self._split(account, instrument, '3', '1')

        parcel = self._active_parcel(buy)
        self.assertEqual(parcel.parcel_quantity, Decimal('1000'))
        self.assertEqual(
            parcel.total_cost_base.amount.quantize(Decimal('0.01')), Decimal('3000.00'))

    def test_deleting_a_split_restores_the_parcels_it_split(self):
        account, instrument, buy = self._holding()
        CostBaseAdjustment.objects.create(
            account=account, instrument=instrument, financial_year_end_date=date(2022, 6, 30),
            cost_base_increase=Money(Decimal('50'), 'AUD'),
        )
        split = self._split(account, instrument, '1', '3')

        split.delete()

        parcel = self._active_parcel(buy)
        self.assertEqual(parcel.parcel_quantity, Decimal('100'))
        self.assertEqual(parcel.total_cost_base.amount, Decimal('1050'))

    def test_a_split_whose_parcels_have_since_been_sold_cannot_be_deleted(self):
        """Reversing it would leave the sale in post-split units against pre-split parcels."""
        account, instrument, buy = self._holding()
        split = self._split(account, instrument, '1', '2')
        Sell.objects.create(
            account=account, instrument=instrument, date=date(2023, 6, 1),
            quantity=Decimal('50'), unit_price=Money(Decimal('8'), 'AUD'),
            total_brokerage=Money(0, 'AUD'), strategy='FIFO',
        )

        with self.assertRaises(ValueError):
            split.delete()

        self.assertTrue(ShareSplit.objects.filter(pk=split.pk).exists())
        self.assertEqual(Instrument.objects.get(pk=instrument.pk).quantity_held, Decimal('150'))


class StructuralEditTests(TransactionTestCase):
    """What other records were derived from cannot be changed; prices can be corrected."""

    def setUp(self):
        self.account = create_account()
        self.instrument = create_instrument(account=self.account)
        self.buy = Buy.objects.create(
            account=self.account, instrument=self.instrument, date=date(2024, 1, 5),
            quantity=Decimal('100'), unit_price=Money(10, 'AUD'), total_brokerage=Money(10, 'AUD'),
        )

    def _refused(self, record, **changes):
        for name, value in changes.items():
            setattr(record, name, value)
        with self.assertRaisesMessage(ValidationError, 'Delete it and enter it again'):
            record.full_clean()
        with self.assertRaisesMessage(ValueError, 'Delete it and enter it again'):
            record.save()

    def test_a_buys_quantity_cannot_be_changed(self):
        self._refused(self.buy, quantity=Decimal('200'))
        self.assertEqual(Parcel.objects.get(buy=self.buy).parcel_quantity, Decimal('100'))

    def test_a_buys_date_cannot_be_changed(self):
        self._refused(self.buy, date=date(2024, 2, 5))

    def test_a_sells_quantity_cannot_be_changed(self):
        sell = Sell.objects.create(
            account=self.account, instrument=self.instrument, date=date(2024, 6, 1),
            quantity=Decimal('40'), unit_price=Money(12, 'AUD'),
            total_brokerage=Money(0, 'AUD'), strategy='FIFO',
        )
        self._refused(sell, quantity=Decimal('50'))

    def test_an_allocations_quantity_cannot_be_changed(self):
        sell = Sell.objects.create(
            account=self.account, instrument=self.instrument, date=date(2024, 6, 1),
            quantity=Decimal('40'), unit_price=Money(12, 'AUD'),
            total_brokerage=Money(0, 'AUD'), strategy='FIFO',
        )
        self._refused(sell.sale_allocation.get(), quantity=Decimal('30'))

    def test_a_splits_ratio_cannot_be_changed(self):
        split = ShareSplit.objects.create(
            account=self.account, instrument=self.instrument, date=date(2024, 3, 1),
            quantity_before=Decimal('1'), quantity_after=Decimal('2'),
        )
        self._refused(split, quantity_after=Decimal('3'))

    def test_an_adjustments_amount_cannot_be_changed(self):
        adjustment = CostBaseAdjustment.objects.create(
            account=self.account, instrument=self.instrument,
            financial_year_end_date=date(2024, 6, 30), cost_base_increase=Money(50, 'AUD'),
        )
        self._refused(adjustment, cost_base_increase=Money(80, 'AUD'))

    def test_saving_without_a_change_is_allowed(self):
        self.buy.notes = 'contract note filed'
        self.buy.full_clean()
        self.buy.save()

    def test_a_corrected_price_reaches_the_parcel_and_the_gain(self):
        sell = Sell.objects.create(
            account=self.account, instrument=self.instrument, date=date(2024, 6, 1),
            quantity=Decimal('40'), unit_price=Money(12, 'AUD'),
            total_brokerage=Money(0, 'AUD'), strategy='FIFO',
        )

        self.buy.unit_price = Money(11, 'AUD')
        self.buy.save()

        held = Parcel.objects.get(buy=self.buy, deactivation_date__isnull=True, sale_date__isnull=True)
        # 60 units at $11 plus 60% of the $10 brokerage.
        self.assertEqual(held.calculated_total_cost_base.amount, Decimal('666'))
        allocation = sell.sale_allocation.get()
        # 40 x $12 less (40 x $11 + 40% of $10).
        self.assertEqual(allocation.calculated_total_capital_gain.amount, Decimal('36'))


class OutOfOrderEntryTests(TransactionTestCase):
    """Entered by hand after a split, a trade dated before it would miss the split."""

    def setUp(self):
        self.account = create_account()
        self.instrument = create_instrument(account=self.account)
        Buy.objects.create(
            account=self.account, instrument=self.instrument, date=date(2020, 1, 10),
            quantity=Decimal('100'), unit_price=Money(10, 'AUD'), total_brokerage=Money(0, 'AUD'),
        )

    def _split(self, on=date(2021, 1, 1)):
        return ShareSplit.objects.create(
            account=self.account, instrument=self.instrument, date=on,
            quantity_before=Decimal('1'), quantity_after=Decimal('2'),
        )

    def _refused(self, record):
        with self.assertRaisesMessage(ValidationError, 'split'):
            record.full_clean()
        with self.assertRaisesMessage(ValueError, 'split'):
            record.save()

    def test_a_sale_dated_before_an_applied_split_is_refused(self):
        self._split()
        self._refused(Sell(
            account=self.account, instrument=self.instrument, date=date(2020, 6, 1),
            quantity=Decimal('50'), unit_price=Money(12, 'AUD'),
            total_brokerage=Money(0, 'AUD'), strategy='FIFO'))

    def test_a_buy_dated_before_an_applied_split_is_refused(self):
        self._split()
        self._refused(Buy(
            account=self.account, instrument=self.instrument, date=date(2020, 6, 1),
            quantity=Decimal('50'), unit_price=Money(12, 'AUD'), total_brokerage=Money(0, 'AUD')))

    def test_a_split_dated_before_an_allocated_sale_is_refused(self):
        Sell.objects.create(
            account=self.account, instrument=self.instrument, date=date(2022, 1, 1),
            quantity=Decimal('50'), unit_price=Money(12, 'AUD'),
            total_brokerage=Money(0, 'AUD'), strategy='FIFO',
        )
        self._refused(ShareSplit(
            account=self.account, instrument=self.instrument, date=date(2021, 1, 1),
            quantity_before=Decimal('1'), quantity_after=Decimal('2')))

    def test_trades_after_a_split_are_fine(self):
        self._split()
        Buy.objects.create(
            account=self.account, instrument=self.instrument, date=date(2021, 6, 1),
            quantity=Decimal('10'), unit_price=Money(6, 'AUD'), total_brokerage=Money(0, 'AUD'),
        )
        Sell.objects.create(
            account=self.account, instrument=self.instrument, date=date(2022, 1, 1),
            quantity=Decimal('50'), unit_price=Money(12, 'AUD'),
            total_brokerage=Money(0, 'AUD'), strategy='FIFO',
        )
        self.assertEqual(
            Instrument.objects.get(pk=self.instrument.pk).quantity_held, Decimal('160'))


class AdjustmentOutOfOrderTests(TransactionTestCase):
    """An adjustment is spread over the holding once, so a trade dated before it is refused."""

    def setUp(self):
        self.account = create_account()
        self.instrument = create_instrument(account=self.account)
        Buy.objects.create(
            account=self.account, instrument=self.instrument, date=date(2022, 7, 1),
            quantity=Decimal('100'), unit_price=Money(10, 'AUD'), total_brokerage=Money(0, 'AUD'),
        )

    def _adjustment(self, year_end=date(2024, 6, 30), amount=200):
        return CostBaseAdjustment.objects.create(
            account=self.account, instrument=self.instrument, financial_year_end_date=year_end,
            cost_base_increase=Money(amount, 'AUD'),
        )

    def _buy(self, on):
        return Buy(
            account=self.account, instrument=self.instrument, date=on,
            quantity=Decimal('100'), unit_price=Money(10, 'AUD'), total_brokerage=Money(0, 'AUD'))

    def _sell(self, on):
        return Sell(
            account=self.account, instrument=self.instrument, date=on,
            quantity=Decimal('50'), unit_price=Money(12, 'AUD'),
            total_brokerage=Money(0, 'AUD'), strategy='FIFO')

    def _refused(self, record):
        with self.assertRaisesMessage(ValidationError, 'cost base adjustment'):
            record.full_clean()
        with self.assertRaisesMessage(ValueError, 'cost base adjustment'):
            record.save()

    def test_a_buy_dated_before_an_applied_adjustments_year_end_is_refused(self):
        """It would get none of the adjustment, though held in its year."""
        self._adjustment()
        self._refused(self._buy(date(2022, 7, 1)))
        self._refused(self._buy(date(2024, 6, 30)))

    def test_a_buy_is_refused_even_where_the_adjustment_reached_no_parcel(self):
        self._adjustment(year_end=date(2022, 6, 30))
        self._refused(self._buy(date(2021, 7, 1)))

    def test_a_buy_after_the_year_end_is_fine(self):
        self._adjustment()
        self._buy(date(2024, 7, 1)).save()

    def test_a_sale_dated_before_an_applied_adjustments_year_end_is_refused(self):
        """The units sold would keep a share weighted as if held all year."""
        self._adjustment()
        self._refused(self._sell(date(2024, 6, 29)))

    def test_a_sale_on_the_year_end_is_fine(self):
        """The last day counts as held either way, so no weight changes."""
        self._adjustment()
        self._sell(date(2024, 6, 30)).save()
        self.assertEqual(
            Instrument.objects.get(pk=self.instrument.pk).quantity_held, Decimal('50'))

    def test_a_sale_is_fine_where_the_adjustment_reached_no_parcel(self):
        self._adjustment(year_end=date(2022, 6, 30))
        self._sell(date(2022, 8, 1)).save()

    def test_a_manual_adjustment_does_not_refuse_trades(self):
        CostBaseAdjustment.objects.create(
            account=self.account, instrument=self.instrument,
            financial_year_end_date=date(2024, 6, 30),
            cost_base_increase=Money(200, 'AUD'), allocation_method='MANUAL',
        )
        self._buy(date(2023, 1, 1)).save()

    def test_a_fresh_template_import_is_unaffected(self):
        """Buys load first, then sales and adjustments in date order."""
        account = create_account(
            owner=create_user(username='fresh'), description='Fresh',
            fy_type=create_fiscal_year_type(description='AU Tax Year fresh'))
        path = Path(tempfile.mkdtemp()) / 'fresh.xlsx'
        self.addCleanup(shutil.rmtree, path.parent)
        trade = {'instrument__name': 'BHP', 'unit_price_currency': 'AUD',
                 'total_brokerage': Decimal('0'), 'total_brokerage_currency': 'AUD'}
        data = {
            'Market': pd.DataFrame([{'code': 'ASX', 'suffix': 'AX'}]),
            'Instrument': pd.DataFrame([
                {'name': 'BHP', 'currency': 'AUD', 'market__code': 'ASX'}]),
            'Buy': pd.DataFrame([
                {**trade, 'legacy_id': 'B1', 'date': date(2022, 7, 1),
                 'quantity': Decimal('100'), 'unit_price': Decimal('10')},
                {**trade, 'legacy_id': 'B2', 'date': date(2024, 3, 1),
                 'quantity': Decimal('100'), 'unit_price': Decimal('11')},
                {**trade, 'legacy_id': 'B3', 'date': date(2025, 3, 1),
                 'quantity': Decimal('100'), 'unit_price': Decimal('12')},
            ]),
            'Sell': pd.DataFrame([
                {**trade, 'legacy_id': 'S1', 'date': date(2023, 12, 1), 'strategy': 'FIFO',
                 'quantity': Decimal('50'), 'unit_price': Decimal('12')},
                {**trade, 'legacy_id': 'S2', 'date': date(2025, 1, 1), 'strategy': 'FIFO',
                 'quantity': Decimal('50'), 'unit_price': Decimal('12')},
            ]),
            'CostBaseAdjustment': pd.DataFrame([
                {'legacy_id': 'A1', 'cost_base_increase': Decimal('200'),
                 'cost_base_increase_currency': 'AUD', 'instrument__name': 'BHP',
                 'financial_year_end_date': date(2024, 6, 30), 'allocation_method': 'QTY_HELD'},
            ]),
        }
        generator = excelinterface.ExcelGen(title='Fresh')
        for table_name, frame in data.items():
            generator.add_table(frame, table_name=table_name)
        generator.save(path)

        loading.DataLoader(account=account, input_file=path)

        self.assertEqual(Sell.objects.filter(account=account).count(), 2)
        adjustment = CostBaseAdjustment.objects.get(account=account)
        allocated = sum(
            allocation.cost_base_increase.amount
            for allocation in adjustment.cost_base_adjustment_allocation.filter(is_active=True))
        self.assertEqual(allocated, Decimal('200'))

    def test_an_adjustment_that_reached_no_parcel_is_reported(self):
        from share_dinkum_app import data_checks

        self.assertNotIn('empty_adjustments', {f.key for f in data_checks.run(self.account)})
        self._adjustment(year_end=date(2022, 6, 30))

        finding = {f.key: f for f in data_checks.run(self.account)}['empty_adjustments']
        self.assertEqual(finding.count, 1)
        self.assertTrue(finding.affects_gains)
        self.assertFalse(finding.repairable)


class AttachedDocumentTests(TransactionTestCase):
    """A document is deleted only once the change that replaced or removed it commits."""

    def setUp(self):
        from django.core.files.base import ContentFile
        from django.test import override_settings
        self.ContentFile = ContentFile
        media = override_settings(MEDIA_ROOT=tempfile.mkdtemp())
        media.enable()
        self.addCleanup(media.disable)

        account = create_account()
        self.buy = Buy.objects.create(
            account=account, instrument=create_instrument(account=account),
            date=date(2024, 1, 5), quantity=Decimal('100'),
            unit_price=Money(50, 'AUD'), total_brokerage=Money(0, 'AUD'),
        )
        self.buy.file.save('contract.pdf', ContentFile(b'original'))
        self.original = Path(self.buy.file.path)

    def test_a_rolled_back_replacement_keeps_the_original(self):
        with self.assertRaises(RuntimeError):
            with transaction.atomic():
                self.buy.file.save('replacement.pdf', self.ContentFile(b'new'))
                raise RuntimeError('the import failed after this row')

        self.assertTrue(self.original.exists())

    def test_a_committed_replacement_removes_the_original(self):
        self.buy.file.save('replacement.pdf', self.ContentFile(b'new'))
        self.assertFalse(self.original.exists())


class AdminDeletePermissionTests(TransactionTestCase):
    """Records the admin must not delete: parcels, and splits that can no longer be reversed."""

    def _request(self):
        from django.test import RequestFactory
        request = RequestFactory().get('/')
        request.user = AppUser.objects.create_superuser('root', 'root@example.com', 'x')
        return request

    def test_parcels_cannot_be_deleted(self):
        account = create_account()
        Buy.objects.create(
            account=account, instrument=create_instrument(account=account),
            date=date(2024, 1, 5), quantity=Decimal('100'),
            unit_price=Money(50, 'AUD'), total_brokerage=Money(0, 'AUD'),
        )
        parcel = Parcel.objects.get()
        self.assertFalse(admin.site._registry[Parcel].has_delete_permission(self._request(), parcel))

    def test_a_split_that_cannot_be_reversed_has_no_delete_button(self):
        account = create_account()
        instrument = create_instrument(account=account)
        Buy.objects.create(
            account=account, instrument=instrument, date=date(2024, 1, 5),
            quantity=Decimal('100'), unit_price=Money(50, 'AUD'), total_brokerage=Money(0, 'AUD'),
        )
        split = ShareSplit.objects.create(
            account=account, instrument=instrument, date=date(2024, 3, 1),
            quantity_before=Decimal('1'), quantity_after=Decimal('2'),
        )
        split_admin = admin.site._registry[ShareSplit]
        request = self._request()
        self.assertTrue(split_admin.has_delete_permission(request, split))

        Sell.objects.create(
            account=account, instrument=instrument, date=date(2024, 6, 1),
            quantity=Decimal('50'), unit_price=Money(55, 'AUD'),
            total_brokerage=Money(0, 'AUD'), strategy='FIFO',
        )
        self.assertFalse(split_admin.has_delete_permission(request, split))


# =============================================================================
# Models: CostBaseAdjustment
# =============================================================================


class CostBaseAdjustmentTests(TestCase):
    """Tests for CostBaseAdjustment model."""

    def test_get_description(self):
        acc = create_account()
        inst = create_instrument(account=acc)
        adj = CostBaseAdjustment(
            account=acc,
            instrument=inst,
            financial_year_end_date=date(2024, 6, 30),
            cost_base_increase=Money(100, 'AUD'),
            allocation_method='MANUAL',
        )
        adj.save()
        desc = adj.get_description()
        self.assertIn('2024-06-30', desc)
        self.assertIn(inst.name, desc)
        self.assertIn('100', desc)


# =============================================================================
# Models: LogEntry, BaseModel
# =============================================================================


class LogEntryTests(TestCase):
    """Tests for LogEntry and BaseModel.log_event."""

    def test_log_event_creates_entry(self):
        acc = create_account()
        inst = create_instrument(account=acc)
        inst.log_event('Test event')
        entries = LogEntry.objects.filter(account=acc, object_id=inst.id)
        self.assertEqual(entries.count(), 1)
        self.assertIn('Test event', entries.first().event)

    def test_log_entry_str_format(self):
        acc = create_account()
        inst = create_instrument(account=acc)
        inst.log_event('Something happened')
        entry = LogEntry.objects.filter(account=acc).first()
        self.assertIn('Something happened', str(entry))
        self.assertIn(str(entry.pk)[-4:], str(entry))


# =============================================================================
# Models: InstrumentPriceHistory
# =============================================================================


class InstrumentPriceHistoryTests(TestCase):
    """Tests for InstrumentPriceHistory (get_absolute_url)."""

    def test_get_absolute_url(self):
        acc = create_account()
        inst = create_instrument(account=acc)
        iph = InstrumentPriceHistory.objects.create(
            account=acc,
            instrument=inst,
            date=date(2024, 1, 15),
            open=Decimal('10'),
            high=Decimal('11'),
            low=Decimal('9'),
            close=Decimal('10.5'),
            volume=1000,
            stock_splits=Decimal('0'),
        )
        url = iph.get_absolute_url()
        self.assertIn('admin', url)
        self.assertIn(str(iph.id), url)


# =============================================================================
# Models: Dividend
# =============================================================================


class DividendTests(TestCase):
    """Tests for Dividend model calculated fields."""

    def test_total_franked_amount(self):
        acc = create_account()
        inst = create_instrument(account=acc)
        div = Dividend.objects.create(
            account=acc,
            instrument=inst,
            date=date(2024, 4, 1),
            quantity=Decimal('100'),
            franked_amount_per_share=Money(Decimal('0.50'), 'AUD'),
            unfranked_amount_per_share=Money(0, 'AUD'),
        )
        self.assertEqual(div.total_franked_amount.amount, Decimal('50'))
        self.assertEqual(str(div.total_franked_amount.currency), 'AUD')


# =============================================================================
# Reports
# =============================================================================


class RealisedCapitalGainReportTests(TestCase):
    """Tests for RealisedCapitalGainReport."""

    def test_generate_empty_account(self):
        acc = create_account()
        report = RealisedCapitalGainReport(account=acc)
        df = report.generate()
        self.assertTrue(df.empty)
        self.assertEqual(
            list(df.columns),
            [
                'sell_date', 'instrument', 'quantity_sold', 'buy_id', 'parcel_id', 'sell_id',
                'sell_allocation_id', 'buy_date', 'days_held', 'proceeds', 'cost_base',
                'capital_gain', 'fiscal_year',
            ],
        )

    def test_generate_with_one_sale_allocation(self):
        acc = create_account()
        inst = create_instrument(account=acc)
        buy = Buy.objects.create(
            account=acc,
            instrument=inst,
            date=date(2023, 1, 5),
            quantity=Decimal('100'),
            unit_price=Money(50, 'AUD'),
            total_brokerage=Money(0, 'AUD'),
        )
        sell = Sell.objects.create(
            account=acc,
            instrument=inst,
            date=date(2024, 2, 1),
            quantity=Decimal('50'),
            unit_price=Money(60, 'AUD'),
            total_brokerage=Money(0, 'AUD'),
            strategy='FIFO',
        )
        report = RealisedCapitalGainReport(account=acc)
        df = report.generate()
        self.assertEqual(len(df), 1)
        self.assertEqual(df.iloc[0]['instrument'], inst.name)
        self.assertEqual(df.iloc[0]['quantity_sold'], Decimal('50'))
        self.assertIn('capital_gain', df.columns)


# =============================================================================
# Reports: golden master
# =============================================================================


def create_golden_master_portfolio():
    """A fixed portfolio exercising every input to a capital gains figure.

    Two buy years, a split, an AMIT adjustment, a partial sell that bifurcates, a sell
    spanning two parcels, a gain and a loss. Figures asserted against it record current
    behaviour; if one moves, explain why before updating it.
    """
    account = create_account()
    instrument = create_instrument(account=account, name='GMT')

    buy_one = Buy.objects.create(
        account=account, instrument=instrument, date=date(2022, 8, 15),
        quantity=Decimal('1000'), unit_price=Money(Decimal('10.00'), 'AUD'),
        total_brokerage=Money(Decimal('19.95'), 'AUD'),
    )
    buy_two = Buy.objects.create(
        account=account, instrument=instrument, date=date(2023, 2, 20),
        quantity=Decimal('500'), unit_price=Money(Decimal('12.00'), 'AUD'),
        total_brokerage=Money(Decimal('9.50'), 'AUD'),
    )

    # One-for-two split: every parcel doubles, unit prices halve.
    share_split = ShareSplit.objects.create(
        account=account, instrument=instrument, date=date(2023, 9, 1),
        quantity_before=Decimal('1'), quantity_after=Decimal('2'),
    )

    # Bought part way through FY2024, and after the split so it is never doubled. This
    # parcel exercises the cost base allocation weighting: held 61 days of the year, it
    # gets that fraction of a full year's weight.
    buy_three = Buy.objects.create(
        account=account, instrument=instrument, date=date(2024, 5, 1),
        quantity=Decimal('200'), unit_price=Money(Decimal('7.00'), 'AUD'),
        total_brokerage=Money(Decimal('9.50'), 'AUD'),
    )

    # Partial sell at a gain; bifurcates the first parcel.
    sell_gain = Sell.objects.create(
        account=account, instrument=instrument, date=date(2024, 3, 10),
        quantity=Decimal('1500'), unit_price=Money(Decimal('8.00'), 'AUD'),
        total_brokerage=Money(Decimal('9.50'), 'AUD'), strategy='FIFO',
    )

    # AMIT cost base increase, spread across parcels by quantity x days held. Entered after
    # the sale before its year end, as an import's timeline does: the other way round is
    # refused, since the sold units' share would be weighted as if held all year.
    adjustment = CostBaseAdjustment.objects.create(
        account=account, instrument=instrument,
        financial_year_end_date=date(2024, 6, 30),
        cost_base_increase=Money(Decimal('150.00'), 'AUD'),
        allocation_method='QTY_HELD',
    )
    # Later sell at a loss, in the next fiscal year, spanning two parcels.
    sell_loss = Sell.objects.create(
        account=account, instrument=instrument, date=date(2024, 11, 5),
        quantity=Decimal('1000'), unit_price=Money(Decimal('4.00'), 'AUD'),
        total_brokerage=Money(Decimal('9.50'), 'AUD'), strategy='FIFO',
    )

    # So OpenParcelReport produces market values rather than nulls.
    instrument.current_unit_price = Decimal('5.5000')
    instrument.save()

    return {
        'account': account, 'instrument': instrument,
        'buy_one': buy_one, 'buy_two': buy_two, 'buy_three': buy_three,
        'share_split': share_split, 'adjustment': adjustment,
        'sell_gain': sell_gain, 'sell_loss': sell_loss,
    }


class CGTGoldenMasterTests(TransactionTestCase):
    """Pins every capital gains figure the app derives for the golden master portfolio.

    These record current behaviour, not independently derived tax answers.
    """

    def setUp(self):
        self.data = create_golden_master_portfolio()
        self.account = self.data['account']

    def _parcels(self):
        """Parcels with a cost base, oldest buy first, excluding deactivated ones."""
        return [
            p for p in Parcel.objects.filter(account=self.account).order_by('buy__date', 'id')
            if p.remaining_quantity or p.sale_date
        ]

    def test_parcel_cost_bases(self):
        parcels = self._parcels()
        self.assertEqual(len(parcels), 5)

        expected = [
            # buy date,        quantity, unit cost base,       total cost base,      sale date
            (date(2022, 8, 15), '1500', '5.050412266666666666666666667', '7575.6184', date(2024, 3, 10)),
            (date(2022, 8, 15), '500',  '5.0682428', '2534.1214', date(2024, 11, 5)),
            (date(2023, 2, 20), '500',  '6.067768',  '3033.8840', date(2024, 11, 5)),
            (date(2023, 2, 20), '500',  '6.0677678', '3033.8839', None),
            (date(2024, 5, 1),  '200',  '7.0572115', '1411.4423', None),
        ]
        for parcel, (buy_date, qty, unit_cb, total_cb, sale_date) in zip(parcels, expected):
            self.assertEqual(parcel.buy.date, buy_date)
            self.assertEqual(parcel.parcel_quantity, Decimal(qty))
            self.assertEqual(parcel.unit_cost_base.amount, Decimal(unit_cb))
            self.assertEqual(parcel.total_cost_base.amount, Decimal(total_cb))
            self.assertEqual(parcel.sale_date, sale_date)

    def test_cost_base_adjustment_allocation(self):
        """The $150 AMIT adjustment is weighted by quantity times days held in the year."""
        allocations = [p.total_adjustments.amount for p in self._parcels()]
        self.assertEqual(
            allocations,
            # The units sold on 10 March are weighted by the 254 days they were held.
            [Decimal('60.6559'), Decimal('29.1339'),
             Decimal('29.1340'), Decimal('29.1339'), Decimal('1.9423')],
        )
        # The whole adjustment is allocated, exactly: none lost to rounding, none
        # duplicated. The largest parcel absorbs the residual where the weights do not
        # divide evenly, and a bifurcated allocation gives its remainder the difference
        # rather than the complementary fraction.
        self.assertEqual(sum(allocations), Decimal('150.00'))

    def test_a_parcel_bought_late_in_the_year_gets_a_proportionate_share(self):
        """A parcel bought late in the year gets a proportionately smaller share."""
        parcels = {p.buy.date: p for p in self._parcels()}
        held_all_year = parcels[date(2023, 2, 20)]
        held_61_days = parcels[date(2024, 5, 1)]

        per_unit_full = held_all_year.total_adjustments.amount / held_all_year.parcel_quantity
        per_unit_part = held_61_days.total_adjustments.amount / held_61_days.parcel_quantity

        # 61 of 366 days, so a sixth of the per unit share of a parcel held throughout.
        # Not exact: each allocation is rounded to the cost base column's four decimal
        # places, which on a $1.65 share is a relative difference of a few parts in 100,000.
        self.assertAlmostEqual(
            float(per_unit_part / per_unit_full), 61 / 366, places=4)

    def test_realised_capital_gain_report(self):
        df = RealisedCapitalGainReport(account=self.account).generate()
        self.assertEqual(len(df), 3)

        expected = [
            # sell date,       qty,    days, proceeds,   cost base,              gain,                    fiscal year
            (date(2024, 3, 10), '1500', 573, '11990.50', '7575.6184', '4414.8816',  'FY2023/24'),
            (date(2024, 11, 5), '500',  813, '1995.25',  '2534.1214', '-538.8714',  'FY2024/25'),
            (date(2024, 11, 5), '500',  624, '1995.25',  '3033.8840', '-1038.6340', 'FY2024/25'),
        ]
        for (_, row), (sell_date, qty, days, proceeds, cost_base, gain, fy) in zip(df.iterrows(), expected):
            self.assertEqual(row['sell_date'], sell_date)
            self.assertEqual(row['quantity_sold'], Decimal(qty))
            self.assertEqual(row['days_held'], days)
            self.assertEqual(row['proceeds'].amount, Decimal(proceeds))
            self.assertEqual(row['cost_base'].amount, Decimal(cost_base))
            self.assertEqual(row['capital_gain'].amount, Decimal(gain))
            self.assertEqual(str(row['fiscal_year']), fy)

        # Both a gain and a loss are represented, so sign handling is exercised.
        gains = [row['capital_gain'].amount for _, row in df.iterrows()]
        self.assertTrue(any(g > 0 for g in gains))
        self.assertTrue(any(g < 0 for g in gains))

    def test_open_parcel_report(self):
        df = OpenParcelReport(account=self.account).generate()
        self.assertEqual(len(df), 2)

        expected = [
            # buy date,       qty,   unit cost base,        cost base,              market value, unrealised gain
            (date(2023, 2, 20), '500', '6.0677678', '3033.8839', '2750.00', '-283.8839'),
            (date(2024, 5, 1),  '200', '7.0572115', '1411.4423', '1100.00', '-311.4423'),
        ]
        for (_, row), (buy_date, qty, unit_cb, cost_base, value, gain) in zip(df.iterrows(), expected):
            self.assertEqual(row['buy_date'], buy_date)
            self.assertEqual(row['remaining_quantity'], Decimal(qty))
            self.assertEqual(row['unit_cost_base'].amount, Decimal(unit_cb))
            self.assertEqual(row['cost_base'].amount, Decimal(cost_base))
            self.assertEqual(row['current_value'].amount, Decimal(value))
            self.assertEqual(row['unrealised_gain'].amount, Decimal(gain))

        self.assertAlmostEqual(df.iloc[0]['unrealised_gain_pct'], -0.09357111522955773, places=12)
        self.assertAlmostEqual(df.iloc[1]['unrealised_gain_pct'], -0.22065535374701467, places=12)

    def test_quantities_reconcile(self):
        """Every unit bought is either held or allocated to a sale."""
        bought = Decimal('1000') * 2 + Decimal('500') * 2 + Decimal('200')   # post-split
        sold = sum(a.quantity for a in SellAllocation.objects.filter(account=self.account))
        held = sum(p.remaining_quantity for p in self._parcels())
        self.assertEqual(sold, Decimal('2500'))
        self.assertEqual(held, Decimal('700'))
        self.assertEqual(sold + held, bought)


class CGTEventTests(TransactionTestCase):
    """The cgt package reproduces the app's existing figures exactly."""

    def setUp(self):
        self.data = create_golden_master_portfolio()
        self.account = self.data['account']

    def test_events_match_the_realised_capital_gain_report(self):
        """Disposal events match RealisedCapitalGainReport row for row."""
        events = cgt.disposal_events(self.account)
        df = RealisedCapitalGainReport(account=self.account).generate()

        self.assertEqual(len(events), len(df))

        by_allocation = {str(row['sell_allocation_id']): row for _, row in df.iterrows()}
        for event in events:
            row = by_allocation[str(event.sell_allocation_id)]
            self.assertEqual(event.instrument, row['instrument'])
            self.assertEqual(event.quantity, row['quantity_sold'])
            self.assertEqual(event.days_held, row['days_held'])
            self.assertEqual(event.purchase_date, row['buy_date'])
            self.assertEqual(event.event_date, row['sell_date'])
            self.assertEqual(event.fiscal_year, row['fiscal_year'])
            self.assertEqual(event.cost_base.amount, row['cost_base'].amount)
            self.assertEqual(event.net_proceeds.amount, row['proceeds'].amount)
            self.assertEqual(event.capital_gain.amount, row['capital_gain'].amount)

    def test_cost_base_components_sum_to_the_cost_base(self):
        """Purchase, brokerage and adjustments sum to the cost base."""
        for event in cgt.disposal_events(self.account):
            components = (
                event.buy_consideration.amount
                + event.buy_brokerage.amount
                + event.cost_base_adjustments.amount
            )
            self.assertEqual(
                components, event.cost_base.amount,
                f'components do not rebuild the cost base for {event.sell_allocation_id}')

    def test_gross_gain_and_loss_split_the_signed_gain(self):
        for event in cgt.disposal_events(self.account):
            self.assertGreaterEqual(event.gross_gain.amount, 0)
            self.assertGreaterEqual(event.gross_loss.amount, 0)
            # Exactly one side carries the value.
            self.assertTrue(event.gross_gain.amount == 0 or event.gross_loss.amount == 0)
            self.assertEqual(
                event.gross_gain.amount - event.gross_loss.amount, event.capital_gain.amount)

    def test_fiscal_year_filter(self):
        all_events = cgt.disposal_events(self.account)
        self.assertEqual(len(all_events), 3)
        fy2024 = cgt.disposal_events(self.account, fiscal_year='FY2023/24')
        self.assertEqual(len(fy2024), 1)
        fy2025 = cgt.disposal_events(self.account, fiscal_year='FY2024/25')
        self.assertEqual(len(fy2025), 2)

    def test_characterisation_matches_todays_rules(self):
        for event in cgt.disposal_events(self.account):
            # Everything in the fixture is held well over a year, and sold before 2027.
            self.assertEqual(event.method, 'discount')
            self.assertEqual(event.discount_percentage, Decimal('0.5'))
            self.assertEqual(event.regime, 'pre_2027')

    def test_events_are_immutable(self):
        event = cgt.disposal_events(self.account)[0]
        with self.assertRaises(Exception):
            event.capital_gain = Money(Decimal('1'), 'AUD')

    def test_empty_account_yields_no_events(self):
        # Reuses the existing owner and fiscal year type: both carry unique constraints, so
        # create_account() cannot be called twice within one test.
        empty = Account.objects.create(
            owner=self.account.owner,
            description='Empty',
            currency=DEFAULT_CURRENCY,
            fiscal_year_type=self.account.fiscal_year_type,
        )
        self.assertEqual(cgt.disposal_events(empty), [])


class CGTClassificationTests(TransactionTestCase):
    """Deriving the CGT schedule category from an instrument's legal form and market."""

    def setUp(self):
        self.account = create_account()

    def _instrument(self, name, legal_form, market_code='ASX', suffix='AX',
                    country=None, listed=True):
        market, _ = Market.objects.get_or_create(
            account=self.account, code=market_code,
            defaults={'suffix': suffix},
        )
        if country is not None or not listed:
            market.country = country
            market.is_exchange_listed = listed
            market.save()
        instrument = Instrument.objects.create(
            account=self.account, market=market, name=name, currency=DEFAULT_CURRENCY)
        instrument.legal_form = legal_form
        instrument.legal_form_source = 'USER'
        instrument.save()
        return instrument

    def test_australian_listed_company_is_shares(self):
        instrument = self._instrument('AFI', 'COMPANY')
        self.assertEqual(instrument.cgt_asset_category, CGTAssetCategory.AU_LISTED_SHARES)

    def test_australian_listed_trust_is_units(self):
        instrument = self._instrument('VAS', 'UNIT_TRUST')
        self.assertEqual(instrument.cgt_asset_category, CGTAssetCategory.AU_LISTED_UNITS)

    def test_currency_and_holdings_do_not_decide_the_category(self):
        """VGS is Australian listed units, whatever its currency or holdings."""
        instrument = self._instrument('VGS', 'UNIT_TRUST')
        self.assertEqual(instrument.cgt_asset_category, CGTAssetCategory.AU_LISTED_UNITS)

    def test_foreign_listed_fund_is_other_units(self):
        instrument = self._instrument('VWRA', 'UNIT_TRUST', market_code='LSE', suffix='L')
        self.assertEqual(instrument.cgt_asset_category, CGTAssetCategory.OTHER_UNITS)

    def test_foreign_listed_company_is_other_shares(self):
        instrument = self._instrument('AAPL', 'COMPANY', market_code='NASDAQ', suffix='')
        self.assertEqual(instrument.cgt_asset_category, CGTAssetCategory.OTHER_SHARES)

    def test_stapled_securities_are_units(self):
        instrument = self._instrument('SCG', 'STAPLED')
        self.assertEqual(instrument.cgt_asset_category, CGTAssetCategory.AU_LISTED_UNITS)

    def test_unlisted_australian_company_is_other_shares(self):
        instrument = self._instrument(
            'PRIV', 'COMPANY', market_code='UNLISTED', suffix='', country='AU', listed=False)
        self.assertEqual(instrument.cgt_asset_category, CGTAssetCategory.OTHER_SHARES)

    def test_rights_and_debt_fall_to_other_assets(self):
        self.assertEqual(self._instrument('CTDR', 'RIGHT_OPTION').cgt_asset_category, CGTAssetCategory.OTHER_ASSETS)
        self.assertEqual(self._instrument('BOND', 'DEBT').cgt_asset_category, CGTAssetCategory.OTHER_ASSETS)

    def test_real_property_is_categorised_by_where_the_land_is(self):
        au = self._instrument('LAND', 'REAL_PROPERTY')
        self.assertEqual(au.cgt_asset_category, CGTAssetCategory.AU_REAL_ESTATE)
        overseas = self._instrument(
            'VILLA', 'REAL_PROPERTY', market_code='LSE', suffix='L')
        self.assertEqual(overseas.cgt_asset_category, CGTAssetCategory.OVERSEAS_REAL_ESTATE)
        self.assertTrue(cgt.is_real_property(au))

    def test_unclassified_is_never_guessed_into_a_category(self):
        market = Market.objects.create(account=self.account, code='ASX', suffix='AX')
        instrument = Instrument.objects.create(
            account=self.account, market=market, name='WHAT', currency=DEFAULT_CURRENCY)
        self.assertEqual(instrument.legal_form, 'UNKNOWN')
        self.assertEqual(instrument.cgt_asset_category, CGTAssetCategory.UNCLASSIFIED)
        self.assertFalse(instrument.is_classified)

    def test_an_override_wins_over_the_derived_category(self):
        instrument = self._instrument('ODD', 'COMPANY')
        instrument.cgt_asset_category_override = CGTAssetCategory.COLLECTABLES
        instrument.save()
        self.assertEqual(instrument.cgt_asset_category, CGTAssetCategory.COLLECTABLES)

    def test_a_suggestion_is_not_a_confirmation(self):
        market = Market.objects.create(account=self.account, code='ASX', suffix='AX')
        instrument = Instrument.objects.create(
            account=self.account, market=market, name='VAS', currency=DEFAULT_CURRENCY)
        # Seeded on creation, so the category resolves...
        self.assertEqual(instrument.legal_form, 'UNIT_TRUST')
        self.assertEqual(instrument.legal_form_source, 'SUGGESTED')
        # ...but it is not treated as settled.
        self.assertFalse(instrument.is_classified)


class InstrumentClassificationSignalTests(TransactionTestCase):
    """Suggestions on creation never overwrite a user's answer."""

    def setUp(self):
        self.account = create_account()
        self.market = Market.objects.create(account=self.account, code='ASX', suffix='AX')

    def test_market_country_is_suggested_from_the_suffix(self):
        self.assertEqual(self.market.country, 'AU')
        lse = Market.objects.create(account=self.account, code='LSE', suffix='L')
        self.assertEqual(lse.country, 'GB')

    def test_an_explicit_country_is_not_overwritten(self):
        market = Market.objects.create(
            account=self.account, code='CHIA', suffix='AX', country='NZ')
        self.assertEqual(market.country, 'NZ')

    def test_a_seeded_code_is_suggested(self):
        vas = Instrument.objects.create(
            account=self.account, market=self.market, name='VAS', currency=DEFAULT_CURRENCY)
        self.assertEqual(vas.legal_form, 'UNIT_TRUST')
        afi = Instrument.objects.create(
            account=self.account, market=self.market, name='AFI', currency=DEFAULT_CURRENCY)
        self.assertEqual(afi.legal_form, 'COMPANY')

    def test_an_unknown_code_stays_unknown(self):
        instrument = Instrument.objects.create(
            account=self.account, market=self.market, name='ZZZZ', currency=DEFAULT_CURRENCY)
        self.assertEqual(instrument.legal_form, 'UNKNOWN')
        self.assertEqual(instrument.legal_form_source, 'DEFAULT')

    def test_a_user_answer_survives_a_later_save(self):
        instrument = Instrument.objects.create(
            account=self.account, market=self.market, name='VAS', currency=DEFAULT_CURRENCY)
        instrument.legal_form = 'COMPANY'
        instrument.legal_form_source = 'USER'
        instrument.save()
        instrument.refresh_from_db()
        instrument.save()
        instrument.refresh_from_db()
        self.assertEqual(instrument.legal_form, 'COMPANY')
        self.assertEqual(instrument.legal_form_source, 'USER')

    def test_market_data_is_not_trusted_for_equities(self):
        """A yfinance EQUITY quoteType gives no suggestion; ETF suggests a unit trust."""
        instrument = Instrument.objects.create(
            account=self.account, market=self.market, name='ZZZZ', currency=DEFAULT_CURRENCY)
        self.assertIsNone(cgt.suggest_legal_form(instrument, {'quoteType': 'EQUITY'}))
        self.assertEqual(cgt.suggest_legal_form(instrument, {'quoteType': 'ETF'}), 'UNIT_TRUST')


class LegalFormFromActivityTests(TransactionTestCase):
    """Inferring legal form from dividend and distribution history."""

    def setUp(self):
        self.account = create_account()
        self.market = Market.objects.create(account=self.account, code='ASX', suffix='AX')

    def _instrument(self, name):
        return Instrument.objects.create(
            account=self.account, market=self.market, name=name, currency=DEFAULT_CURRENCY)

    def _dividend(self, instrument):
        Dividend.objects.create(
            account=self.account, instrument=instrument, date=date(2024, 3, 1),
            quantity=Decimal('100'),
            franked_amount_per_share=Money(Decimal('0.10'), 'AUD'),
            unfranked_amount_per_share=Money(0, 'AUD'),
        )

    def _distribution(self, instrument):
        Distribution.objects.create(
            account=self.account, instrument=instrument, date=date(2024, 3, 1),
            quantity=Decimal('100'),
            distribution_amount_per_share=Money(Decimal('0.10'), 'AUD'),
        )

    def test_dividends_alone_imply_a_company(self):
        instrument = self._instrument('WOW')
        self._dividend(instrument)
        self.assertEqual(cgt.suggest_legal_form_from_activity(instrument), 'COMPANY')

    def test_distributions_alone_imply_a_trust(self):
        instrument = self._instrument('ZZTRUST')
        self._distribution(instrument)
        self.assertEqual(cgt.suggest_legal_form_from_activity(instrument), 'UNIT_TRUST')

    def test_paying_both_implies_a_stapled_security(self):
        """An instrument paying dividends and distributions is suggested as stapled."""
        instrument = self._instrument('ZZSTAPLE')
        self._dividend(instrument)
        self._distribution(instrument)
        self.assertEqual(cgt.suggest_legal_form_from_activity(instrument), 'STAPLED')

    def test_no_income_history_gives_no_answer(self):
        """No income history gives no suggestion."""
        self.assertIsNone(cgt.suggest_legal_form_from_activity(self._instrument('ZZQUIET')))

    def test_trust_and_stapled_reach_the_same_schedule_category(self):
        """A unit trust and a stapled security share a schedule category; a company does not."""
        trust = self._instrument('ZZT')
        trust.legal_form, trust.legal_form_source = 'UNIT_TRUST', 'USER'
        trust.save()
        stapled = self._instrument('ZZS')
        stapled.legal_form, stapled.legal_form_source = 'STAPLED', 'USER'
        stapled.save()
        self.assertEqual(trust.cgt_asset_category, stapled.cgt_asset_category)

        company = self._instrument('ZZC')
        company.legal_form, company.legal_form_source = 'COMPANY', 'USER'
        company.save()
        self.assertNotEqual(company.cgt_asset_category, trust.cgt_asset_category)


class CGTDiscountTests(TestCase):
    """The discount eligibility and rate rules, in isolation."""

    def test_held_more_than_365_days_is_eligible(self):
        self.assertTrue(cgt.is_discount_eligible(date(2023, 1, 1), date(2024, 1, 2)))
        self.assertEqual(
            cgt.discount_percentage(date(2023, 1, 1), date(2024, 1, 2)), Decimal('0.5'))

    def test_held_exactly_365_days_is_not_eligible(self):
        self.assertFalse(cgt.is_discount_eligible(date(2023, 1, 1), date(2024, 1, 1)))
        self.assertEqual(
            cgt.discount_percentage(date(2023, 1, 1), date(2024, 1, 1)), Decimal('0'))

    def test_the_anniversary_itself_does_not_qualify(self):
        """A sale on the 12-month anniversary does not qualify (s115-25(1))."""
        self.assertFalse(cgt.is_discount_eligible(date(2023, 6, 15), date(2024, 6, 15)))
        self.assertTrue(cgt.is_discount_eligible(date(2023, 6, 15), date(2024, 6, 16)))

    def test_29_february_falls_back_to_28_february(self):
        """29 February's anniversary is 28 February."""
        from share_dinkum_app.cgt.discount import twelve_month_anniversary
        self.assertEqual(twelve_month_anniversary(date(2024, 2, 29)), date(2025, 2, 28))
        self.assertEqual(twelve_month_anniversary(date(2023, 3, 1)), date(2024, 3, 1))
        # Bought on the leap day, sold on 1 March the next year: past 28 February, eligible.
        self.assertTrue(cgt.is_discount_eligible(date(2024, 2, 29), date(2025, 3, 1)))
        self.assertFalse(cgt.is_discount_eligible(date(2024, 2, 29), date(2025, 2, 28)))

    def test_missing_dates_are_not_eligible(self):
        self.assertFalse(cgt.is_discount_eligible(None, date(2024, 1, 1)))
        self.assertFalse(cgt.is_discount_eligible(date(2024, 1, 1), None))

    def test_the_rate_is_exact_not_a_float(self):
        """The discount rate is an exact Decimal, not a float."""
        from share_dinkum_app.cgt.discount import FULL_DISCOUNT_RATE
        self.assertIsInstance(FULL_DISCOUNT_RATE, Decimal)
        self.assertEqual(FULL_DISCOUNT_RATE, Decimal('0.5'))

    def test_a_loss_is_never_discounted(self):
        from share_dinkum_app.cgt.discount import apply_discount
        loss = Money(Decimal('-100'), 'AUD')
        self.assertEqual(
            apply_discount(loss, date(2020, 1, 1), date(2024, 1, 1)), loss)


class AttributionStatementTests(TransactionTestCase):
    """Trust-attributed capital gains, using figures from a real Vanguard AMMA statement."""

    def setUp(self):
        self.account = create_account()
        self.instrument = create_instrument(account=self.account, name='VGS')
        self.instrument.legal_form = 'UNIT_TRUST'
        self.instrument.legal_form_source = 'USER'
        self.instrument.save()
        self.statement = AttributionStatement.objects.create(
            account=self.account, instrument=self.instrument,
            financial_year_end_date=date(2025, 6, 30),
        )

    def _component(self, component, amount):
        return AttributionComponent.objects.create(
            account=self.account, statement=self.statement,
            component=component, amount=Money(Decimal(amount), 'AUD'),
        )

    def _real_statement(self):
        """Record the FY2025 VGS statement: all NTAP, no other-method gains."""
        self._component('DISCOUNTED_NTAP', '4845.57')
        self._component('AMIT_GROSS_UP', '4845.57')
        self._component('NET_CAPITAL_GAIN', '4845.57')
        self._component('TOTAL_CY_CG', '9691.14')
        self._component('FOREIGN_SOURCE_INCOME', '9484.75')

    def test_the_grossed_up_total_is_derived_not_transcribed(self):
        self._real_statement()
        self.assertEqual(self.statement.discounted_capital_gain, Decimal('4845.57'))
        self.assertEqual(self.statement.other_method_capital_gain, Decimal('0'))
        # Twice the discounted amount, matching the statement's own stated total.
        self.assertEqual(
            self.statement.total_current_year_capital_gain, Decimal('9691.14'))

    def test_a_statement_that_reconciles_says_so(self):
        self._real_statement()
        self.assertIs(self.statement.reconciles, True)

    def test_a_misread_statement_is_detected(self):
        """A statement whose components do not match its stated total does not reconcile."""
        self._component('DISCOUNTED_NTAP', '4845.57')
        self._component('TOTAL_CY_CG', '5000.00')
        self.assertIs(self.statement.reconciles, False)

    def test_no_stated_total_means_nothing_to_check(self):
        self._component('DISCOUNTED_NTAP', '4845.57')
        self.assertIsNone(self.statement.reconciles)

    def test_an_attribution_becomes_a_grossed_up_cgt_event(self):
        self._real_statement()
        events = cgt.attribution_events(self.account)
        self.assertEqual(len(events), 1)

        event = events[0]
        self.assertEqual(event.source, 'trust_attribution')
        # The gross up matters: carrying the trust's halved figure would discount twice.
        self.assertEqual(event.capital_gain.amount, Decimal('9691.14'))
        self.assertEqual(event.method, 'discount')
        self.assertEqual(event.discount_percentage, Decimal('0.5'))
        self.assertEqual(event.tap_status, 'NTAP')
        self.assertEqual(event.event_date, date(2025, 6, 30))
        self.assertEqual(event.asset_category, CGTAssetCategory.AU_LISTED_UNITS)
        self.assertIs(event.source_reconciles, True)
        # The member never held what the trust sold, so there is no holding period.
        self.assertIsNone(event.purchase_date)
        self.assertIsNone(event.days_held)
        self.assertIsNone(event.cost_base)

    def test_discounted_and_other_method_gains_stay_separate(self):
        """Discounted and other-method gains become separate events."""
        self._component('DISCOUNTED_NTAP', '100.00')
        self._component('OTHER_NTAP', '40.00')
        self._component('TOTAL_CY_CG', '240.00')

        events = {event.method: event for event in cgt.attribution_events(self.account)}
        self.assertEqual(set(events), {'discount', 'other'})
        self.assertEqual(events['discount'].capital_gain.amount, Decimal('200.00'))
        self.assertEqual(events['discount'].discount_percentage, Decimal('0.5'))
        self.assertEqual(events['other'].capital_gain.amount, Decimal('40.00'))
        self.assertEqual(events['other'].discount_percentage, Decimal('0'))

    def test_the_tap_split_is_preserved(self):
        """TAP and NTAP discounted gains become one event each, as the statement states them."""
        self._component('DISCOUNTED_TAP', '30.00')
        self._component('DISCOUNTED_NTAP', '70.00')
        events = cgt.attribution_events(self.account)
        self.assertEqual(
            [(e.tap_status, e.capital_gain.amount) for e in events],
            [('TAP', Decimal('60.00')), ('NTAP', Decimal('140.00'))])

    def test_a_statement_with_no_capital_gains_produces_no_events(self):
        self._component('FOREIGN_SOURCE_INCOME', '1000.00')
        self.assertEqual(cgt.attribution_events(self.account), [])

    def test_attributions_join_disposals_in_the_combined_view(self):
        self._real_statement()
        Buy.objects.create(
            account=self.account, instrument=self.instrument, date=date(2020, 1, 6),
            quantity=Decimal('100'), unit_price=Money(Decimal('80.00'), 'AUD'),
            total_brokerage=Money(0, 'AUD'),
        )
        Sell.objects.create(
            account=self.account, instrument=self.instrument, date=date(2025, 2, 3),
            quantity=Decimal('40'), unit_price=Money(Decimal('110.00'), 'AUD'),
            total_brokerage=Money(0, 'AUD'), strategy='FIFO',
        )
        combined = cgt.all_events(self.account)
        self.assertEqual(len(combined), 2)
        self.assertEqual(
            [event.source for event in combined], ['disposal', 'trust_attribution'])

    def test_one_statement_per_instrument_per_year(self):
        with self.assertRaises(IntegrityError):
            AttributionStatement.objects.create(
                account=self.account, instrument=self.instrument,
                financial_year_end_date=date(2025, 6, 30),
            )

    def test_a_component_appears_once_per_statement(self):
        self._component('DISCOUNTED_NTAP', '10.00')
        with self.assertRaises(IntegrityError):
            self._component('DISCOUNTED_NTAP', '20.00')


class CostBaseAllocationReconciliationTests(TransactionTestCase):
    """Cost base adjustments still sum exactly after repeated parcel splits."""

    def setUp(self):
        self.account = create_account()
        self.instrument = create_instrument(account=self.account, name='RND')
        self.buy = Buy.objects.create(
            account=self.account, instrument=self.instrument, date=date(2022, 5, 3),
            quantity=Decimal('1000'), unit_price=Money(Decimal('3.37'), 'AUD'),
            total_brokerage=Money(Decimal('19.95'), 'AUD'),
        )
        # Chosen to land exactly on a rounding boundary when halved: 49.4505 / 2 is
        # 24.72525, which Django's DecimalField rounds half to even, giving 24.7252 twice
        # and losing a hundredth of a cent. An amount that divides cleanly proves nothing.
        self.adjustment_amount = Decimal('49.4505')
        self.adjustment = CostBaseAdjustment.objects.create(
            account=self.account, instrument=self.instrument,
            financial_year_end_date=date(2023, 6, 30),
            cost_base_increase=Money(self.adjustment_amount, 'AUD'),
            allocation_method='QTY_HELD',
        )

    def _live_allocation_total(self):
        return sum(
            allocation.cost_base_increase.amount
            for allocation in CostBaseAdjustmentAllocation.objects.filter(
                cost_base_adjustment=self.adjustment, deactivation_date__isnull=True)
        )

    def _live_parcel_quantity(self):
        return sum(
            parcel.parcel_quantity
            for parcel in Parcel.objects.filter(
                account=self.account, deactivation_date__isnull=True)
        )

    def test_allocation_survives_an_even_split_on_a_rounding_boundary(self):
        self.assertEqual(self._live_allocation_total(), self.adjustment_amount)
        # Exactly half the parcel, so each side lands on 24.72525.
        Sell.objects.create(
            account=self.account, instrument=self.instrument, date=date(2023, 9, 1),
            quantity=Decimal('500'), unit_price=Money(Decimal('4.11'), 'AUD'),
            total_brokerage=Money(Decimal('9.50'), 'AUD'), strategy='FIFO',
        )
        self.assertEqual(self._live_allocation_total(), self.adjustment_amount)

    def test_allocation_survives_being_split_seven_times(self):
        """An allocation split seven times still sums to the original."""
        for index, quantity in enumerate(['500', '250', '125', '62', '31', '15', '7']):
            Sell.objects.create(
                account=self.account, instrument=self.instrument,
                date=date(2023, 9, 1) + timedelta(days=index * 30),
                quantity=Decimal(quantity), unit_price=Money(Decimal('4.11'), 'AUD'),
                total_brokerage=Money(Decimal('9.50'), 'AUD'), strategy='FIFO',
            )
            self.assertEqual(
                self._live_allocation_total(), self.adjustment_amount,
                f'cost base went missing after split {index + 1}')

        # Halving repeatedly is the worst case for a per-split rounding loss.
        self.assertEqual(self._live_parcel_quantity(), Decimal('1000'))

    def test_every_unit_carries_a_share_after_splitting(self):
        """Every parcel carries a share of the adjustment after splitting."""
        Sell.objects.create(
            account=self.account, instrument=self.instrument, date=date(2023, 9, 1),
            quantity=Decimal('333'), unit_price=Money(Decimal('4.11'), 'AUD'),
            total_brokerage=Money(Decimal('9.50'), 'AUD'), strategy='FIFO',
        )
        parcels = Parcel.objects.filter(
            account=self.account, deactivation_date__isnull=True)
        for parcel in parcels:
            if parcel.parcel_quantity:
                self.assertGreater(
                    parcel.total_adjustments.amount, Decimal('0'),
                    f'parcel of {parcel.parcel_quantity} units has no adjustment share')


class MinCGTSelectionTests(TransactionTestCase):
    """Which parcels the MIN_CGT strategy picks, ranked by discounted gain per unit."""

    def _account_with_parcels(self):
        """Three parcels: two long held at different cost bases, one bought recently."""
        account = create_account()
        instrument = create_instrument(account=account, name='SEL')
        for buy_date, quantity, price, tag in [
            (date(2020, 1, 15), '100', '5.00', 'old_cheap'),
            (date(2020, 6, 15), '100', '9.00', 'old_dear'),
            (date(2024, 9, 1), '100', '7.00', 'recent'),
        ]:
            Buy.objects.create(
                account=account, instrument=instrument, date=buy_date,
                quantity=Decimal(quantity), unit_price=Money(Decimal(price), 'AUD'),
                total_brokerage=Money(0, 'AUD'),
            )
        return account, instrument

    def _allocated_buy_dates(self, sell):
        return [
            a.parcel.buy.date
            for a in sell.sale_allocation.filter(is_active=True).order_by('parcel__buy__date')
        ]

    def test_min_cgt_prefers_the_parcel_with_the_smallest_discounted_gain(self):
        account, instrument = self._account_with_parcels()
        sell = Sell.objects.create(
            account=account, instrument=instrument, date=date(2024, 10, 1),
            quantity=Decimal('100'), unit_price=Money(Decimal('10.00'), 'AUD'),
            total_brokerage=Money(0, 'AUD'), strategy='MIN_CGT',
        )
        # Per unit: old_cheap gain $5.00 halved to $2.50; old_dear $1.00 halved to $0.50;
        # recent $3.00 undiscounted because it is inside 365 days. old_dear ranks lowest.
        self.assertEqual(self._allocated_buy_dates(sell), [date(2020, 6, 15)])

    def test_the_discount_changes_the_choice(self):
        """The discount changes which parcels MIN_CGT picks."""
        account, instrument = self._account_with_parcels()
        sell = Sell.objects.create(
            account=account, instrument=instrument, date=date(2024, 10, 1),
            quantity=Decimal('250'), unit_price=Money(Decimal('10.00'), 'AUD'),
            total_brokerage=Money(0, 'AUD'), strategy='MIN_CGT',
        )
        self.assertEqual(
            self._allocated_buy_dates(sell),
            [date(2020, 1, 15), date(2020, 6, 15), date(2024, 9, 1)],
        )
        allocations = {
            a.parcel.buy.date: a.quantity
            for a in sell.sale_allocation.filter(is_active=True)
        }
        self.assertEqual(allocations[date(2020, 6, 15)], Decimal('100'))  # cheapest first
        self.assertEqual(allocations[date(2020, 1, 15)], Decimal('100'))  # then discounted
        self.assertEqual(allocations[date(2024, 9, 1)], Decimal('50'))    # remainder

    def test_holding_period_no_longer_depends_on_a_leap_day(self):
        """A leap day inside the holding no longer changes eligibility."""
        # Spans 29 February 2024: 366 days, and formerly eligible on that basis alone.
        across_leap_day = (date(2023, 3, 1), date(2024, 3, 1))
        self.assertEqual((across_leap_day[1] - across_leap_day[0]).days, 366)
        self.assertFalse(cgt.is_discount_eligible(*across_leap_day))

        # The same calendar holding a year later: 365 days, and never eligible.
        no_leap_day = (date(2022, 3, 1), date(2023, 3, 1))
        self.assertEqual((no_leap_day[1] - no_leap_day[0]).days, 365)
        self.assertFalse(cgt.is_discount_eligible(*no_leap_day))

        # A day past the anniversary qualifies in both cases.
        self.assertTrue(cgt.is_discount_eligible(date(2023, 3, 1), date(2024, 3, 2)))
        self.assertTrue(cgt.is_discount_eligible(date(2022, 3, 1), date(2023, 3, 2)))


# =============================================================================
# Reports: snapshots and basis change detection
# =============================================================================


class CGTReturnSnapshotTests(TransactionTestCase):
    """Capturing capital gains snapshots."""

    def setUp(self):
        self.data = create_golden_master_portfolio()
        self.account = self.data['account']
        self.fy2024 = FiscalYear.objects.get(
            fiscal_year_type=self.account.fiscal_year_type, start_year=2023)
        self.fy2025 = FiscalYear.objects.get(
            fiscal_year_type=self.account.fiscal_year_type, start_year=2024)

    def test_capture_records_only_the_requested_year(self):
        snapshot = CGTReturnSnapshot.capture(account=self.account, fiscal_year=self.fy2024)
        self.assertEqual(len(snapshot.rows), 1)
        self.assertEqual(snapshot.totals['row_count'], 1)
        self.assertEqual(Decimal(snapshot.totals['total_capital_gain']), Decimal('4414.8816'))

        later = CGTReturnSnapshot.capture(account=self.account, fiscal_year=self.fy2025)
        self.assertEqual(len(later.rows), 2)
        self.assertEqual(Decimal(later.totals['total_capital_gain']), Decimal('-1577.5054'))

    def test_captured_figures_keep_their_values(self):
        snapshot = CGTReturnSnapshot.capture(account=self.account, fiscal_year=self.fy2024)
        row = snapshot.rows[0]
        self.assertEqual(row['cost_base'].amount, Decimal('7575.6184'))
        self.assertEqual(row['capital_gain'].amount, Decimal('4414.8816'))
        self.assertEqual(row['proceeds'].amount, Decimal('11990.5000'))
        self.assertEqual(row['sell_date'], date(2024, 3, 10))
        # The fiscal year belongs to the snapshot, so it is not repeated on every row.
        self.assertEqual(snapshot.fiscal_year.name, 'FY2023/24')

    def test_each_disposal_is_its_own_row(self):
        """A snapshot stores one row per sell allocation."""
        snapshot = CGTReturnSnapshot.capture(account=self.account, fiscal_year=self.fy2024)
        self.assertEqual(snapshot.captured_rows.count(), 1)
        self.assertEqual(
            set(snapshot.rows[0].keys()), set(CGTReturnSnapshot.CAPTURED_FIELDS))
        self.assertFalse(hasattr(snapshot, 'payload'))

    def test_a_large_year_is_captured_rather_than_refused(self):
        """A snapshot can hold thousands of rows."""
        snapshot = CGTReturnSnapshot.capture(account=self.account, fiscal_year=self.fy2024)
        CGTReturnSnapshotRow.objects.bulk_create([
            CGTReturnSnapshotRow(
                account=self.account, snapshot=snapshot, instrument='BULK',
                sell_date=date(2024, 3, 10), quantity_sold=Decimal('1'),
                days_held=400, proceeds=Money(Decimal('10'), 'AUD'),
                cost_base=Money(Decimal('8'), 'AUD'),
                capital_gain=Money(Decimal('2'), 'AUD'))
            for _ in range(2000)
        ])
        self.assertEqual(snapshot.captured_rows.count(), 2001)
        self.assertEqual(snapshot.totals['row_count'], 2001)

    def test_the_allocation_reference_is_not_a_foreign_key(self):
        """A snapshot row survives deletion of its sell allocation."""
        snapshot = CGTReturnSnapshot.capture(account=self.account, fiscal_year=self.fy2024)
        row = snapshot.captured_rows.first()
        self.assertIsNotNone(row.sell_allocation_id)

        SellAllocation.objects.filter(id=row.sell_allocation_id).delete()

        row.refresh_from_db()
        self.assertIsNotNone(row.sell_allocation_id)
        self.assertEqual(row.capital_gain.amount, Decimal('4414.8816'))

    def test_recapturing_same_day_replaces_rather_than_duplicates(self):
        first = CGTReturnSnapshot.capture(
            account=self.account, fiscal_year=self.fy2024, taken_at=date(2026, 8, 31))
        second = CGTReturnSnapshot.capture(
            account=self.account, fiscal_year=self.fy2024, taken_at=date(2026, 8, 31))
        self.assertEqual(first.pk, second.pk)
        self.assertEqual(CGTReturnSnapshot.objects.filter(account=self.account).count(), 1)

    def test_capturing_a_later_day_adds_to_the_history(self):
        CGTReturnSnapshot.capture(
            account=self.account, fiscal_year=self.fy2024, taken_at=date(2026, 8, 30))
        CGTReturnSnapshot.capture(
            account=self.account, fiscal_year=self.fy2024, taken_at=date(2026, 8, 31))
        self.assertEqual(CGTReturnSnapshot.objects.filter(account=self.account).count(), 2)

    def test_capture_records_the_engine_version(self):
        snapshot = CGTReturnSnapshot.capture(account=self.account, fiscal_year=self.fy2024)
        self.assertEqual(snapshot.engine_version, version.__version__)

    def test_a_snapshot_exports_as_two_readable_sheets(self):
        """A snapshot exports as two readable sheets."""
        snapshot = CGTReturnSnapshot.capture(
            account=self.account, fiscal_year=self.fy2024, is_lodged=True)

        export = DataExport.objects.create(account=self.account, include_price_history=False)
        export.refresh_from_db()
        tables = excelinterface.get_all_tables_in_excel(Path(export.file.path))

        self.assertIn('CGTReturnSnapshot', tables)
        self.assertIn('CGTReturnSnapshotRow', tables)

        rows = tables['CGTReturnSnapshotRow']
        for column in ('instrument', 'sell_date', 'proceeds', 'cost_base', 'capital_gain'):
            self.assertIn(column, rows.columns)
        self.assertEqual(len(rows), snapshot.captured_rows.count())
        self.assertEqual(
            Decimal(str(rows.iloc[0]['capital_gain'])), Decimal('4414.8816'))

        self.assertEqual(
            Decimal(snapshot.totals['total_capital_gain']), Decimal('4414.8816'))


class CGTBasisChangeReportTests(TransactionTestCase):
    """The basis change report detects and attributes changed figures."""

    def setUp(self):
        self.data = create_golden_master_portfolio()
        self.account = self.data['account']
        self.fy2024 = FiscalYear.objects.get(
            fiscal_year_type=self.account.fiscal_year_type, start_year=2023)

    def test_the_current_basis_is_the_accounts_basis_today(self):
        """A snapshot on the legacy basis, compared after residency is declared."""
        snapshot = CGTReturnSnapshot.capture(account=self.account, fiscal_year=self.fy2024)
        CGTReturnSnapshotRow.objects.filter(snapshot=snapshot).update(
            capital_gain=Decimal('1'))
        declare(self.account, 'RESIDENT', date(2000, 1, 1))

        df = CGTBasisChangeReport(account=self.account).generate()

        self.assertEqual(set(df['snapshot_basis']), {cgt.BASIS_LEGACY})
        self.assertEqual(set(df['current_basis']), {cgt.BASIS_DIVISION_115})

    def test_no_change_produces_no_rows(self):
        CGTReturnSnapshot.capture(account=self.account, fiscal_year=self.fy2024)
        df = CGTBasisChangeReport(account=self.account).generate()
        self.assertTrue(df.empty)
        self.assertEqual(
            list(df.columns),
            ['fiscal_year', 'taken_at', 'snapshot_basis', 'current_basis',
             'snapshot_engine_version', 'sell_allocation_id', 'status', 'field',
             'snapshot_value', 'current_value', 'difference'],
        )

    def test_a_changed_cost_base_is_reported_with_its_difference(self):
        CGTReturnSnapshot.capture(account=self.account, fiscal_year=self.fy2024)

        # Stand in for any correction that moves a cost base: add a further adjustment.
        # What matters is that the report notices and quantifies it, not the cause.
        CostBaseAdjustment.objects.create(
            account=self.account, instrument=self.data['instrument'],
            financial_year_end_date=date(2023, 6, 30),
            cost_base_increase=Money(Decimal('80.00'), 'AUD'),
            allocation_method='QTY_HELD',
        )

        df = CGTBasisChangeReport(account=self.account).generate()
        self.assertFalse(df.empty)

        changed_fields = set(df['field'])
        self.assertIn('cost_base', changed_fields)
        self.assertIn('capital_gain', changed_fields)
        self.assertEqual(set(df['status']), {'CHANGED'})

        cost_base_row = df[df['field'] == 'cost_base'].iloc[0]
        self.assertEqual(cost_base_row['snapshot_value'], Decimal('7575.6184'))
        self.assertGreater(cost_base_row['current_value'], Decimal('7575.6184'))
        self.assertGreater(cost_base_row['difference'], Decimal('0'))

        # A higher cost base must show as a smaller gain, by the same amount.
        gain_row = df[df['field'] == 'capital_gain'].iloc[0]
        self.assertEqual(gain_row['difference'], -cost_base_row['difference'])

    def test_a_new_allocation_is_reported_as_added(self):
        CGTReturnSnapshot.capture(account=self.account, fiscal_year=self.fy2024)

        Sell.objects.create(
            account=self.account, instrument=self.data['instrument'],
            # On the adjustment's year end: an earlier sale would change its weights.
            date=date(2024, 6, 30), quantity=Decimal('100'),
            unit_price=Money(Decimal('9.00'), 'AUD'),
            total_brokerage=Money(Decimal('9.50'), 'AUD'), strategy='FIFO',
        )

        df = CGTBasisChangeReport(account=self.account).generate()
        added = df[df['status'] == 'ADDED']
        self.assertEqual(len(added), 1)
        self.assertIsNone(added.iloc[0]['snapshot_value'])
        self.assertIsNotNone(added.iloc[0]['current_value'])

    def test_lodged_only_filters_to_lodged_snapshots(self):
        CGTReturnSnapshot.capture(
            account=self.account, fiscal_year=self.fy2024,
            taken_at=date(2026, 8, 30), is_lodged=False)

        CostBaseAdjustment.objects.create(
            account=self.account, instrument=self.data['instrument'],
            financial_year_end_date=date(2023, 6, 30),
            cost_base_increase=Money(Decimal('80.00'), 'AUD'),
            allocation_method='QTY_HELD',
        )

        self.assertFalse(CGTBasisChangeReport(account=self.account).generate().empty)
        self.assertTrue(
            CGTBasisChangeReport(account=self.account, lodged_only=True).generate().empty)


# =============================================================================
# Loading: every model must be importable, not merely exportable
# =============================================================================


class DataLoaderModelCoverageTests(TestCase):
    """Every exported model must also be in the import load order.

    Export covers every model automatically; the load order is maintained by hand.
    """

    def test_every_model_is_in_the_load_order(self):
        from django.apps import apps

        app_models = {m.__name__ for m in apps.get_app_config('share_dinkum_app').get_models()}
        load_order = {m.__name__ for m in loading.DataLoader.get_model_load_order()}

        missing = app_models - load_order
        self.assertEqual(
            missing, set(),
            f"These models are exported but cannot be imported back: {sorted(missing)}. "
            f"Add them to DataLoader.get_model_load_order(), after the models they "
            f"reference by foreign key."
        )

    def test_load_order_contains_no_unknown_models(self):
        from django.apps import apps

        app_models = {m.__name__ for m in apps.get_app_config('share_dinkum_app').get_models()}
        load_order = {m.__name__ for m in loading.DataLoader.get_model_load_order()}

        unknown = load_order - app_models
        self.assertEqual(
            unknown, set(),
            f"The load order names models that no longer exist: {sorted(unknown)}."
        )

    def test_load_order_has_no_duplicates(self):
        names = [m.__name__ for m in loading.DataLoader.get_model_load_order()]
        self.assertEqual(len(names), len(set(names)))

    def test_required_foreign_keys_are_loaded_before_their_dependants(self):
        """No model loads before a model it requires by non-null foreign key.

        Nullable keys are skipped: Account.owner and AppUser.default_account form a cycle.
        """
        order = list(loading.DataLoader.get_model_load_order())
        position = {m.__name__: i for i, m in enumerate(order)}

        for model in order:
            for field in model._meta.get_fields():
                if not (field.many_to_one or field.one_to_one):
                    continue
                if not hasattr(field, 'related_model') or field.related_model is None:
                    continue
                if getattr(field, 'null', True):
                    continue  # optional: can be set after the row is created
                related = field.related_model
                if related is model or related.__name__ not in position:
                    continue  # self reference, or a model outside the load order
                self.assertLessEqual(
                    position[related.__name__], position[model.__name__],
                    f"{model.__name__} is loaded before {related.__name__}, which it "
                    f"requires via the non-nullable field '{field.name}'."
                )


# =============================================================================
# Signals: default account
# =============================================================================


class AssignDefaultAccountSignalTests(TransactionTestCase):
    """A user's first account becomes their default_account."""

    def test_first_account_set_as_default(self):
        user = create_user()
        self.assertIsNone(user.default_account_id)
        acc = create_account(owner=user)
        user.refresh_from_db()
        self.assertEqual(user.default_account_id, acc.id)


# =============================================================================
# Constants
# =============================================================================


class ConstantsTests(TestCase):
    """Sanity check for constants used in logic."""

    def test_cgt_constants(self):
        self.assertEqual(CGT_DISCOUNT_RATE, 0.5)
        self.assertEqual(CGT_DISCOUNT_THRESHOLD_DAYS, 365)

    def test_default_currency(self):
        self.assertEqual(DEFAULT_CURRENCY, 'AUD')


# =============================================================================
# Loading: importing files into more than one portfolio
# =============================================================================


class ImportWorkbookMixin:
    """Builds import workbooks, so tests use the real loading path."""

    def build_workbook(self, path, markets=None, instruments=None, buys=None, market_ids=None):
        if markets is None:
            markets = [{'code': 'ASX', 'description': 'Australian Securities Exchange', 'suffix': 'AX'}]
        if instruments is None:
            instruments = [{'name': 'BHP', 'description': 'BHP Group', 'currency': 'AUD', 'market__code': 'ASX'}]
        if buys is None:
            buys = [{
                'legacy_id': 'buy-1',
                'instrument__name': 'BHP',
                'date': date(2023, 7, 1),
                'quantity': Decimal('100'),
                'unit_price': Decimal('40'),
                'unit_price_currency': 'AUD',
                'total_brokerage': Decimal('10'),
                'total_brokerage_currency': 'AUD',
            }]

        if market_ids is not None:
            markets = [dict(market, id=market_id) for market, market_id in zip(markets, market_ids)]

        generator = excelinterface.ExcelGen(title='Test import')
        generator.add_table(pd.DataFrame(markets), table_name='Market')
        generator.add_table(pd.DataFrame(instruments), table_name='Instrument')
        generator.add_table(pd.DataFrame(buys), table_name='Buy')
        generator.save(path)
        return path


class DataLoaderMultiPortfolioTests(ImportWorkbookMixin, TransactionTestCase):
    """A file loads into exactly one portfolio and never changes another."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.workbook = Path(self.temp_dir.name) / 'import.xlsx'

        self.owner = create_user(username='owner')
        fy_type = create_fiscal_year_type()
        self.account_a = Account.objects.create(
            owner=self.owner, description='Portfolio A', currency='AUD', fiscal_year_type=fy_type)
        self.account_b = Account.objects.create(
            owner=self.owner, description='Portfolio B', currency='AUD', fiscal_year_type=fy_type)

    def test_each_portfolio_gets_its_own_records(self):
        self.build_workbook(self.workbook)

        loading.DataLoader(account=self.account_a, input_file=self.workbook)
        loading.DataLoader(account=self.account_b, input_file=self.workbook)

        for account in (self.account_a, self.account_b):
            self.assertEqual(Market.objects.filter(account=account).count(), 1)
            self.assertEqual(Instrument.objects.filter(account=account).count(), 1)
            self.assertEqual(Buy.objects.filter(account=account).count(), 1)
            self.assertEqual(Parcel.objects.filter(account=account).count(), 1)

    def test_loading_a_second_portfolio_leaves_the_first_alone(self):
        self.build_workbook(self.workbook)
        loading.DataLoader(account=self.account_a, input_file=self.workbook)

        buy_ids_before = set(Buy.objects.filter(account=self.account_a).values_list('id', flat=True))

        other_workbook = Path(self.temp_dir.name) / 'other.xlsx'
        self.build_workbook(
            other_workbook,
            instruments=[{'name': 'CBA', 'description': 'CBA', 'currency': 'AUD', 'market__code': 'ASX'}],
            buys=[{
                'legacy_id': 'buy-99',
                'instrument__name': 'CBA',
                'date': date(2024, 2, 1),
                'quantity': Decimal('5'),
                'unit_price': Decimal('100'),
                'unit_price_currency': 'AUD',
                'total_brokerage': Decimal('10'),
                'total_brokerage_currency': 'AUD',
            }],
        )
        loading.DataLoader(account=self.account_b, input_file=other_workbook)

        buy_ids_after = set(Buy.objects.filter(account=self.account_a).values_list('id', flat=True))
        self.assertEqual(buy_ids_before, buy_ids_after)
        self.assertEqual(Instrument.objects.filter(account=self.account_a, name='CBA').count(), 0)

    def test_loading_the_same_file_twice_is_not_doubled_up(self):
        self.build_workbook(self.workbook)

        loading.DataLoader(account=self.account_a, input_file=self.workbook)
        loading.DataLoader(account=self.account_a, input_file=self.workbook)

        # Reference tables are unique per portfolio, and the buy is matched on its legacy_id.
        self.assertEqual(Market.objects.filter(account=self.account_a).count(), 1)
        self.assertEqual(Instrument.objects.filter(account=self.account_a).count(), 1)
        self.assertEqual(Buy.objects.filter(account=self.account_a).count(), 1)
        self.assertEqual(Parcel.objects.filter(account=self.account_a).count(), 1)

    def test_a_row_with_no_legacy_id_is_refused_the_second_time(self):
        """It cannot be matched to the record it made, so it would be added again."""
        buy = {
            'instrument__name': 'BHP',
            'date': date(2023, 7, 1),
            'quantity': Decimal('100'),
            'unit_price': Decimal('40'),
            'unit_price_currency': 'AUD',
            'total_brokerage': Decimal('10'),
            'total_brokerage_currency': 'AUD',
        }
        self.build_workbook(self.workbook, buys=[buy])

        loading.DataLoader(account=self.account_a, input_file=self.workbook)
        with self.assertRaisesMessage(ValueError, 'legacy_id'):
            loading.DataLoader(account=self.account_a, input_file=self.workbook)

        self.assertEqual(Buy.objects.filter(account=self.account_a).count(), 1)
        # Another portfolio has none yet, so there it is simply added.
        loading.DataLoader(account=self.account_b, input_file=self.workbook)
        self.assertEqual(Buy.objects.filter(account=self.account_b).count(), 1)

    def test_records_are_not_moved_between_portfolios(self):
        self.build_workbook(self.workbook)
        loading.DataLoader(account=self.account_a, input_file=self.workbook)

        market = Market.objects.get(account=self.account_a)

        # An export carries the id of every row, so pointing one at another portfolio would
        # otherwise move the records rather than copy them.
        reparenting_workbook = Path(self.temp_dir.name) / 'export.xlsx'
        self.build_workbook(reparenting_workbook, market_ids=[str(market.id)])

        with self.assertRaises(ValueError) as raised:
            loading.DataLoader(account=self.account_b, input_file=reparenting_workbook)

        self.assertIn('Portfolio A', str(raised.exception))
        market.refresh_from_db()
        self.assertEqual(market.account_id, self.account_a.id)

    def test_a_file_that_fails_part_way_leaves_nothing_behind(self):
        # The buy names an instrument the file never lists, which fails after the reference tables
        # have already been written.
        self.build_workbook(
            self.workbook,
            buys=[{
                'legacy_id': 'buy-1',
                'instrument__name': 'NOT_LISTED',
                'date': date(2023, 7, 1),
                'quantity': Decimal('100'),
                'unit_price': Decimal('40'),
                'unit_price_currency': 'AUD',
                'total_brokerage': Decimal('10'),
                'total_brokerage_currency': 'AUD',
            }],
        )

        with self.assertRaises(Instrument.DoesNotExist):
            loading.DataLoader(account=self.account_a, input_file=self.workbook)

        self.assertEqual(Market.objects.filter(account=self.account_a).count(), 0)
        self.assertEqual(Instrument.objects.filter(account=self.account_a).count(), 0)
        self.assertEqual(Buy.objects.filter(account=self.account_a).count(), 0)


class AccountUniquenessTests(TestCase):
    """Portfolio names are unique per owner."""

    def test_one_owner_cannot_have_two_portfolios_of_the_same_name(self):
        owner = create_user(username='duplicate-owner')
        fy_type = create_fiscal_year_type()
        Account.objects.create(owner=owner, description='Shared name', fiscal_year_type=fy_type)

        with self.assertRaises(IntegrityError):
            Account.objects.create(owner=owner, description='Shared name', fiscal_year_type=fy_type)

    def test_different_owners_may_use_the_same_name(self):
        fy_type = create_fiscal_year_type()
        first = create_user(username='first-owner')
        second = create_user(username='second-owner')

        Account.objects.create(owner=first, description='My Portfolio', fiscal_year_type=fy_type)
        Account.objects.create(owner=second, description='My Portfolio', fiscal_year_type=fy_type)

        self.assertEqual(Account.objects.filter(description='My Portfolio').count(), 2)


# =============================================================================
# Version and update check
# =============================================================================


class VersionTests(TestCase):

    def test_parse_version_accepts_a_tag(self):
        self.assertEqual(version.parse_version('v1.2.3'), (1, 2, 3))
        self.assertEqual(version.parse_version('1.2.3'), (1, 2, 3))

    def test_parse_version_ignores_anything_else(self):
        self.assertIsNone(version.parse_version(None))
        self.assertIsNone(version.parse_version(''))
        self.assertIsNone(version.parse_version('not-a-version'))

    def _temporary_cache(self):
        folder = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, folder, ignore_errors=True)
        cache = patch.object(version, 'get_cache_path', return_value=folder / 'check.json')
        cache.start()
        self.addCleanup(cache.stop)

    def test_an_unreachable_github_reports_no_update(self):
        self._temporary_cache()
        with patch.object(version.requests, 'get', side_effect=OSError('no network')):
            result = version.check_for_update()

        self.assertFalse(result['update_available'])
        self.assertIsNone(result['latest_version'])
        self.assertEqual(result['current_version'], version.__version__)

    def test_an_unreachable_github_is_not_asked_again_on_the_next_page(self):
        """Offline, every dashboard load waited on the request, since a failure was not kept."""
        self._temporary_cache()
        with patch.object(version.requests, 'get', side_effect=OSError('no network')) as get:
            version.check_for_update()
            result = version.check_for_update()

        self.assertEqual(get.call_count, 1)
        self.assertFalse(result['update_available'])

    def test_an_unreachable_github_is_asked_again_after_an_hour(self):
        self._temporary_cache()
        version.write_cache(latest_version=None, release_url=None, reached=False)
        stale = json.loads(version.get_cache_path().read_text(encoding='utf-8'))
        stale['checked_at'] = (timezone.now() - timedelta(hours=2)).isoformat()
        version.get_cache_path().write_text(json.dumps(stale), encoding='utf-8')

        self.assertIsNone(version.read_cache())

    def test_no_published_release_reports_no_update(self):
        response = MagicMock(status_code=404)
        with patch.object(version.requests, 'get', return_value=response):
            with patch.object(version, 'read_cache', return_value=None):
                with patch.object(version, 'write_cache', side_effect=lambda latest_version, release_url: {
                        'latest_version': latest_version, 'release_url': release_url}):
                    result = version.check_for_update()

        self.assertFalse(result['update_available'])
        self.assertIsNone(result['latest_version'])

    def test_a_newer_release_is_reported(self):
        cached = {'latest_version': '99.0.0', 'release_url': 'https://example.invalid/releases/99.0.0'}
        with patch.object(version, 'read_cache', return_value=cached):
            result = version.check_for_update()

        self.assertTrue(result['update_available'])
        self.assertEqual(result['latest_version'], '99.0.0')
        self.assertEqual(result['release_url'], 'https://example.invalid/releases/99.0.0')

    def test_the_installed_release_is_not_reported_as_an_update(self):
        cached = {'latest_version': version.__version__, 'release_url': 'https://example.invalid/'}
        with patch.object(version, 'read_cache', return_value=cached):
            result = version.check_for_update()

        self.assertFalse(result['update_available'])


class ImportTemplateCommandTests(TestCase):
    """The generated import template loads without error."""

    def test_the_template_round_trips_as_an_empty_file(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / 'blank.xlsx'
            call_command('make_import_template', output=str(output_path))

            tables = excelinterface.get_all_tables_in_excel(output_path)

        for model in make_import_template.TEMPLATE_MODELS:
            # The loader finds each table by the model's own name.
            self.assertIn(model.__name__, tables)
            self.assertEqual(len(tables[model.__name__]), 0)

    def test_the_template_offers_no_column_the_app_owns(self):
        for model in make_import_template.TEMPLATE_MODELS:
            columns = make_import_template.get_template_columns(model)
            for owned in ('id', 'account', 'account_id', 'created_at', 'updated_at', 'current_unit_price'):
                self.assertNotIn(owned, columns, f'{model.__name__} should not offer {owned}')
            self.assertFalse([column for column in columns if column.startswith('calculated_')])


class OptionalTemplateTablesTests(TransactionTestCase):
    """The optional template tables: marked as optional, and loadable from a filled-in template."""

    OPTIONAL = ('ResidencyPeriod', 'InstrumentValuation', 'CapitalLossCarryForward',
                'AttributionStatement', 'AttributionComponent')

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.path = Path(self.temp_dir.name) / 'optional.xlsx'
        self.account = create_account()

    def _template(self, **overrides):
        """A template with a buy, an adjustment, and one row in each optional table."""
        data = {
            'Market': pd.DataFrame([{'code': 'ASX', 'suffix': 'AX'}]),
            'Instrument': pd.DataFrame([{'name': 'BHP', 'currency': 'AUD', 'market__code': 'ASX'}]),
            'Buy': pd.DataFrame([{
                'legacy_id': 'B001', 'instrument__name': 'BHP', 'date': date(2022, 7, 1),
                'quantity': Decimal('100'), 'unit_price': Decimal('10'), 'unit_price_currency': 'AUD',
                'total_brokerage': Decimal('0'), 'total_brokerage_currency': 'AUD'}]),
            'CostBaseAdjustment': pd.DataFrame([{
                'legacy_id': 'A001', 'cost_base_increase': Decimal('50'), 'cost_base_increase_currency': 'AUD',
                'instrument__name': 'BHP', 'financial_year_end_date': date(2024, 6, 30),
                'allocation_method': 'QTY_HELD'}]),
            'ResidencyPeriod': pd.DataFrame([{
                'legacy_id': 'R001', 'status': 'RESIDENT', 'start_date': date(2022, 1, 1)}]),
            'InstrumentValuation': pd.DataFrame([{
                'legacy_id': 'V001', 'instrument__name': 'BHP', 'valuation_date': date(2027, 6, 30),
                'unit_value': Decimal('45'), 'unit_value_currency': 'AUD', 'purpose': 'CUTOVER_2027'}]),
            'CapitalLossCarryForward': pd.DataFrame([{
                'legacy_id': 'L001', 'financial_year_end_date': date(2024, 6, 30),
                'amount': Decimal('1500'), 'amount_currency': 'AUD'}]),
            'AttributionStatement': pd.DataFrame([{
                'legacy_id': 'AS001', 'instrument__name': 'BHP',
                'financial_year_end_date': date(2024, 6, 30), 'lookup_legacy_adjustment': 'A001'}]),
            'AttributionComponent': pd.DataFrame([{
                'legacy_id': 'AC001', 'component': 'COSTBASE_INCREASE', 'amount': Decimal('50'),
                'amount_currency': 'AUD', 'lookup_legacy_statement': 'AS001'}]),
        }
        data.update(overrides)
        generator = excelinterface.ExcelGen(title='Optional tables')
        for table_name, frame in data.items():
            generator.add_table(frame, table_name=table_name)
        generator.save(self.path)
        return self.path

    def test_the_template_marks_exactly_the_optional_tables(self):
        call_command('make_import_template', output=str(self.path))

        workbook = openpyxl.load_workbook(self.path)
        descriptions = {
            row[1].value: row[2].value for row in workbook['Index'].iter_rows(min_row=2)}
        tab_colours = {
            list(sheet.tables)[0]: sheet.sheet_properties.tabColor for sheet in workbook.worksheets
            if sheet.tables and sheet.title != 'Index'}

        for model in make_import_template.TEMPLATE_MODELS:
            name = model.__name__
            optional = name in self.OPTIONAL
            self.assertEqual(str(descriptions[name]).startswith('Optional:'), optional, name)
            self.assertEqual(tab_colours[name] is not None, optional, name)

    def test_the_template_offers_the_links_and_hides_the_foreign_keys(self):
        columns = make_import_template.get_template_columns
        self.assertIn('lookup_legacy_adjustment', columns(AttributionStatement))
        self.assertIn('lookup_legacy_statement', columns(AttributionComponent))
        loss_columns = columns(CapitalLossCarryForward)
        self.assertIn('financial_year_end_date', loss_columns)
        self.assertNotIn('fiscal_year__name', loss_columns)

    def test_every_optional_table_loads_and_links_up(self):
        loading.DataLoader(account=self.account, input_file=self._template())

        residency = ResidencyPeriod.objects.get(account=self.account)
        self.assertEqual(residency.status, 'RESIDENT')

        valuation = InstrumentValuation.objects.get(account=self.account)
        self.assertEqual(valuation.source, 'USER')  # blank cell: the default, not a NOT NULL failure

        loss = CapitalLossCarryForward.objects.get(account=self.account)
        self.assertEqual(loss.fiscal_year.start_year, 2023)  # the year ending 30 June 2024
        self.assertFalse(loss.is_opening_balance)  # blank cell: the default

        statement = AttributionStatement.objects.get(account=self.account)
        self.assertEqual(statement.cost_base_adjustment.legacy_id, 'A001')
        component = AttributionComponent.objects.get(account=self.account)
        self.assertEqual(component.statement, statement)

    def test_loading_the_same_template_again_adds_nothing(self):
        loading.DataLoader(account=self.account, input_file=self._template())
        loading.DataLoader(account=self.account, input_file=self.path)

        for model in (ResidencyPeriod, InstrumentValuation, CapitalLossCarryForward,
                      AttributionStatement, AttributionComponent):
            self.assertEqual(model.objects.filter(account=self.account).count(), 1, model.__name__)

    def test_an_empty_optional_table_loads_nothing(self):
        call_command('make_import_template', output=str(self.path))
        loading.DataLoader(account=self.account, input_file=self.path)

        for model in (ResidencyPeriod, InstrumentValuation, CapitalLossCarryForward,
                      AttributionStatement, AttributionComponent):
            self.assertEqual(model.objects.filter(account=self.account).count(), 0, model.__name__)

    def test_a_statement_naming_an_unknown_adjustment_says_so(self):
        template = self._template(AttributionStatement=pd.DataFrame([{
            'legacy_id': 'AS001', 'instrument__name': 'BHP',
            'financial_year_end_date': date(2024, 6, 30), 'lookup_legacy_adjustment': 'NOPE'}]))

        with self.assertRaisesMessage(ValueError, 'no CostBaseAdjustment with that legacy_id'):
            loading.DataLoader(account=self.account, input_file=template)
        self.assertEqual(Buy.objects.filter(account=self.account).count(), 0)  # nothing loaded

    def test_a_loss_with_no_year_says_so(self):
        template = self._template(CapitalLossCarryForward=pd.DataFrame([{
            'legacy_id': 'L001', 'amount': Decimal('1500'), 'amount_currency': 'AUD',
            'financial_year_end_date': None}]))

        with self.assertRaisesMessage(ValueError, 'has no financial_year_end_date'):
            loading.DataLoader(account=self.account, input_file=template)

    def test_a_fiscal_year_name_is_looked_up_within_the_portfolios_type(self):
        other = create_fiscal_year_type(description='Another July year')
        FiscalYear.objects.create(fiscal_year_type=other, start_year=2023)
        mine, _ = self.account.fiscal_year_type.classify_date(date(2024, 6, 30))

        loader = loading.DataLoader(account=self.account)
        found = loader.get_related_obj_by_name(
            related_model=FiscalYear, account=self.account, filters={'name': mine.name})

        self.assertEqual(found, mine)


class MakeFakeDataCommandTests(TestCase):
    """The fake portfolio generator, run on synthetic market data so it needs no network."""

    AS_OF = '2026-10-02'

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.folder = Path(self.temp_dir.name)

        days = pd.bdate_range('2016-06-01', self.AS_OF)
        market = {}
        for holding in make_fake_data.HOLDINGS.values():
            frame = pd.DataFrame({
                'Close': [50 + 0.01 * i for i in range(len(days))],
                'Dividends': [0.0] * len(days),
                'Stock Splits': [0.0] * len(days),
            }, index=days)
            # A payment on the first trading day of each quarter's first month.
            for day in days:
                if day.month in (1, 4, 7, 10) and day.day <= 3 and day.weekday() == 0:
                    frame.loc[day, 'Dividends'] = 0.5
            market[holding.ticker] = frame
        market['AAPL'].loc[pd.Timestamp('2020-08-31'), 'Stock Splits'] = 4.0
        market['AUDUSD=X'] = pd.DataFrame({'Close': [0.7] * len(days)}, index=days)
        self.cache = self.folder / 'market.pkl'
        with open(self.cache, 'wb') as f:
            pickle.dump(market, f)

    def _run(self, name, **options):
        output = self.folder / name
        call_command('make_import_template', output=str(output))
        options.setdefault('force', True)
        call_command('make_fake_data', output=str(output), as_of=self.AS_OF, cache=str(self.cache), **options)
        return output, excelinterface.get_all_tables_in_excel(output)

    def test_it_will_not_overwrite_without_force(self):
        output = self.folder / 'template.xlsx'
        call_command('make_import_template', output=str(output))

        with self.assertRaisesMessage(CommandError, '--force'):
            call_command('make_fake_data', output=str(output), as_of=self.AS_OF, cache=str(self.cache))

    def test_it_needs_a_file_to_fill_in(self):
        with self.assertRaisesMessage(CommandError, 'does not exist'):
            call_command('make_fake_data', output=str(self.folder / 'missing.xlsx'), force=True)

    def test_every_populated_table_is_marked_as_fake_on_its_first_row(self):
        _, tables = self._run('fake.xlsx')

        for name in ('Buy', 'Sell', 'SellAllocation', 'ShareSplit', 'CostBaseAdjustment', 'Dividend',
                     'Distribution', 'ResidencyPeriod'):
            self.assertTrue(len(tables[name]) > 0, name)
            self.assertIn(make_fake_data.FAKE_NOTE, tables[name].iloc[0]['notes'], name)

    def test_the_optional_tables_the_portfolio_does_not_use_stay_empty(self):
        _, tables = self._run('fake.xlsx')

        for name in ('InstrumentValuation', 'CapitalLossCarryForward', 'AttributionStatement',
                     'AttributionComponent'):
            self.assertEqual(len(tables[name]), 0, name)

    def test_the_same_seed_gives_the_same_file_and_another_seed_a_different_one(self):
        _, first = self._run('one.xlsx', seed=7)
        _, again = self._run('two.xlsx', seed=7)
        _, other = self._run('three.xlsx', seed=8)

        pd.testing.assert_frame_equal(first['Buy'], again['Buy'])
        self.assertFalse(first['Buy'].equals(other['Buy']))

    def test_the_portfolio_holds_what_it_should(self):
        _, tables = self._run('fake.xlsx')
        buys, sells = tables['Buy'], tables['Sell']

        self.assertTrue({'VGS', 'A200', 'MSFT', 'AMZN', 'AAPL', 'CBA', 'BHP'} <= set(buys['instrument__name']))
        # Fully sold: every unit bought was sold, and nothing is bought afterwards.
        for name in make_fake_data.FULLY_SOLD:
            bought = buys[buys['instrument__name'] == name]
            sold = sells[sells['instrument__name'] == name]
            self.assertEqual(bought['quantity'].sum(), sold['quantity'].sum(), name)
            self.assertLess(bought['date'].max(), sold['date'].max(), name)
        # The USD satellites are priced in USD.
        self.assertEqual(set(buys[buys['instrument__name'].isin(make_fake_data.SATELLITES)]['unit_price_currency']),
                         {'USD'})
        # The split comes from the market data.
        self.assertEqual(list(tables['ShareSplit']['instrument__name']), ['AAPL'])

    def test_a_manual_sale_has_allocations_that_add_up(self):
        _, tables = self._run('fake.xlsx')
        sells, allocations = tables['Sell'], tables['SellAllocation']

        manual = sells[sells['strategy'] == 'MANUAL']
        self.assertTrue(len(manual) > 0)
        for _, sale in manual.iterrows():
            taken = allocations[allocations['lookup_legacy_sell'] == sale['legacy_id']]['quantity'].sum()
            self.assertEqual(taken, sale['quantity'], sale['legacy_id'])
        self.assertIn('SellAllocation', str(manual.iloc[0]['notes']))


class ColumnHelpTests(TransactionTestCase):
    """Every column says what it is, on the model and as a note on its Excel header."""

    # Fields the app does not declare: the currency beside each money amount is built by django-money,
    # and the user model's own fields come from Django. column_help describes them instead.
    NOT_DECLARED_HERE = ('password', 'last_login', 'first_name', 'last_name', 'email', 'date_joined')

    def _models(self):
        return apps.get_app_config('share_dinkum_app').get_models()

    def _declared_fields(self):
        from djmoney.models.fields import CurrencyField
        for model in self._models():
            for field in model._meta.fields:
                if isinstance(field, CurrencyField) and field.name != 'currency' and getattr(field, 'price_field', None):
                    continue
                if model is AppUser and field.name in self.NOT_DECLARED_HERE:
                    continue
                yield model, field

    def test_every_field_has_help_text(self):
        missing = sorted(f'{model.__name__}.{field.name}' for model, field in self._declared_fields()
                         if not field.help_text)
        self.assertEqual(missing, [], 'These fields need a help_text, which becomes the note on their Excel header.')

    def test_every_field_can_be_described(self):
        missing = sorted(f'{model.__name__}.{field.name}' for model in self._models() for field in model._meta.fields
                         if column_help.field_help(field) is None)
        self.assertEqual(missing, [])

    def test_every_template_header_has_a_note(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'blank.xlsx'
            call_command('make_import_template', output=str(path))
            workbook = openpyxl.load_workbook(path)

        for sheet in workbook.worksheets:
            for cell in sheet[1]:
                if cell.value:
                    self.assertIsNotNone(cell.comment, f'{sheet.title}.{cell.value} has no note')

    def test_a_template_note_says_what_is_required_and_what_a_blank_means(self):
        text = column_help.describe_columns(Sell, ['strategy', 'instrument__name', 'unit_price_currency'], template=True)

        self.assertIn('One of: FIFO', text['strategy'])
        self.assertIn('Blank = MIN_CGT', text['strategy'])
        self.assertIn('Required.', text['instrument__name'])
        self.assertIn("Blank = the instrument's currency", text['unit_price_currency'])

    def test_the_extra_template_columns_are_described(self):
        for model in make_import_template.TEMPLATE_MODELS:
            columns = make_import_template.get_template_columns(model)
            described = column_help.describe_columns(model, columns, template=True)
            self.assertEqual(sorted(set(columns) - set(described)), [], model.__name__)

    def test_a_column_nothing_is_known_about_is_left_out(self):
        self.assertEqual(column_help.describe_columns(Sell, ['no_such_column']), {})

    def test_an_export_puts_a_note_on_every_header(self):
        data = create_golden_master_portfolio()
        export = DataExport.objects.create(account=data['account'])
        export.refresh_from_db()
        workbook = openpyxl.load_workbook(export.file.path)

        model_names = {model.__name__ for model in self._models()}
        checked = 0
        for sheet in workbook.worksheets:
            tables = list(sheet.tables)
            if not tables or tables[0] not in model_names:
                continue
            for cell in sheet[1]:
                if cell.value:
                    checked += 1
                    self.assertIsNotNone(cell.comment, f'{tables[0]}.{cell.value} has no note')
        self.assertGreater(checked, 100)

    def test_add_table_attaches_the_description_to_its_header(self):
        generator = excelinterface.ExcelGen(title='Notes')
        generator.add_table(
            pd.DataFrame([{'a': 1, 'b': 2}]), table_name='Demo',
            column_descriptions={'a': 'The first column.'})
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'notes.xlsx'
            generator.save(path)
            sheet = openpyxl.load_workbook(path)['01']

        self.assertEqual(sheet['A1'].comment.text, 'The first column.')
        self.assertIsNone(sheet['B1'].comment)


class ExcelReaderTests(TestCase):
    """What the reader makes of a hand-edited sheet."""

    def _read(self, rows, headers=('name', 'market__code'), table='Instrument', ref=None):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(list(headers))
        for row in rows:
            sheet.append(row)
        from openpyxl.worksheet.table import Table
        sheet.add_table(Table(displayName=table, ref=ref or f'A1:{chr(64 + len(headers))}{len(rows) + 1}'))
        stream = io.BytesIO()
        workbook.save(stream)
        stream.seek(0)
        return excelinterface.get_all_tables_in_excel(stream)

    def test_surrounding_spaces_are_stripped(self):
        frame = self._read([[' BHP ', 'ASX '], ['VGS', 'ASX']])['Instrument']

        self.assertEqual(list(frame['name']), ['BHP', 'VGS'])
        self.assertEqual(list(frame['market__code']), ['ASX', 'ASX'])

    def test_a_cell_of_only_spaces_is_blank(self):
        frame = self._read([['BHP', '   '], ['VGS', 'ASX']])['Instrument']

        self.assertIsNone(frame['market__code'].iloc[0])

    def test_a_name_column_is_read_as_text_even_when_excel_made_it_a_number(self):
        frame = self._read([[4013, 'HKEX'], [700.0, 'HKEX'], ['BHP', 'ASX']])['Instrument']

        self.assertEqual(list(frame['name']), ['4013', '700', 'BHP'])

    def test_other_columns_keep_their_numbers(self):
        frame = self._read([['BHP', 12.5]], headers=('name', 'unit_price'), table='Buy')['Buy']

        self.assertEqual(frame['unit_price'].iloc[0], 12.5)

    def test_a_table_claiming_an_enormous_range_is_refused(self):
        with self.assertRaisesMessage(ValueError, 'limit'):
            self._read([['BHP', 'ASX']], ref='A1:B1048576')

class DropdownTests(TestCase):
    """Choice columns get a dropdown, backed by a list on the value assistance sheet."""

    def _validations(self, workbook):
        """Each sheet's list validations as {column letter: [allowed values]}."""
        assistance = workbook[excelinterface.VALUE_ASSISTANCE_SHEET]
        found = {}
        for sheet in workbook.worksheets:
            for validation in sheet.data_validations.dataValidation:
                range_ref = validation.formula1.split('!')[1].replace('$', '')
                values = [cell.value for row in assistance[range_ref] for cell in row]
                for cell_range in validation.sqref.ranges:
                    found[(sheet.title, cell_range.coord)] = values
        return found

    def _template(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'blank.xlsx'
            call_command('make_import_template', output=str(path))
            return openpyxl.load_workbook(path), excelinterface.get_all_tables_in_excel(path)

    def test_add_table_validates_a_column_against_its_list(self):
        generator = excelinterface.ExcelGen(title='Dropdowns')
        generator.add_table(pd.DataFrame([{'colour': 'red', 'size': 1}]), table_name='Demo',
                            dropdowns={'colour': ['red', 'green', 'blue']})
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'dropdowns.xlsx'
            generator.save(path)
            workbook = openpyxl.load_workbook(path)

        found = self._validations(workbook)

        # Rows 2 to the one row of data plus the spare rows beneath it.
        self.assertEqual(found, {('01', f'A2:A{1 + 1 + excelinterface.DROPDOWN_SPARE_ROWS}'): ['red', 'green', 'blue']})

    def test_a_list_shared_by_name_is_written_once(self):
        generator = excelinterface.ExcelGen(title='Shared')
        for table in ('One', 'Two'):
            generator.add_table(pd.DataFrame([{'c': 'AUD'}]), table_name=table,
                                dropdowns={'c': ('currency', ['AUD', 'USD'])})
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'shared.xlsx'
            generator.save(path)
            workbook = openpyxl.load_workbook(path)

        self.assertEqual(workbook[excelinterface.VALUE_ASSISTANCE_SHEET].max_column, 1)
        self.assertEqual(len(self._validations(workbook)), 2)

    def test_a_dropdown_for_a_missing_column_is_skipped(self):
        generator = excelinterface.ExcelGen(title='Missing')
        generator.add_table(pd.DataFrame([{'a': 1}]), table_name='Demo', dropdowns={'nope': ['x']})

        with tempfile.TemporaryDirectory() as folder:
            generator.save(Path(folder) / 'missing.xlsx')  # no error

    def test_every_choice_column_in_the_template_has_a_dropdown(self):
        workbook, _ = self._template()
        validated = {(sheet, coord[0]) for sheet, coord in self._validations(workbook)}
        titles = {list(sheet.tables)[0]: sheet.title for sheet in workbook.worksheets if sheet.tables}

        for model in make_import_template.TEMPLATE_MODELS:
            columns = make_import_template.get_template_columns(model)
            for column in make_import_template.get_dropdowns(model, columns):
                letter = openpyxl.utils.get_column_letter(columns.index(column) + 1)
                self.assertIn((titles[model.__name__], letter), validated, f'{model.__name__}.{column}')

    def test_the_lists_hold_the_keys_the_loader_accepts(self):
        workbook, _ = self._template()
        lists = {name: [cell.value for cell in column[1:] if cell.value]
                 for name, column in ((column[0].value, column) for column in workbook[excelinterface.VALUE_ASSISTANCE_SHEET].columns)}

        self.assertEqual(lists['Sell.strategy'], ['FIFO', 'LIFO', 'MIN_CGT', 'MANUAL'])
        self.assertEqual(lists['Dividend.dividend_type'], ['LOCAL', 'FOREIGN'])
        self.assertEqual(lists['ResidencyPeriod.status'], ['RESIDENT', 'FOREIGN', 'TEMPORARY'])
        self.assertIn('AUD', lists['currency'])
        self.assertIn('USD', lists['currency'])

    def test_a_lookup_column_gets_no_dropdown(self):
        """An instrument loaded in an earlier file is a legitimate answer, so lookups stay free text."""
        columns = make_import_template.get_template_columns(Buy)

        self.assertNotIn('instrument__name', make_import_template.get_dropdowns(Buy, columns))

    def test_the_reader_does_not_take_the_value_sheet_for_a_table(self):
        _, tables = self._template()

        self.assertNotIn(excelinterface.VALUE_ASSISTANCE_SHEET, tables)
        self.assertIn('Buy', tables)


# --- Phase 4: residency, apportionment and the foreign resident disregard ---

def declare(account, status, start, end=None, i1=None):
    """Create a residency period for `account`."""
    return ResidencyPeriod.objects.create(
        account=account, status=status, start_date=start, end_date=end,
        i1_election_made=i1,
    )


class ResidencyPeriodValidationTests(TransactionTestCase):
    """Inconsistent residency histories are refused."""

    def setUp(self):
        self.account = create_account()
        self.instrument = create_instrument(account=self.account)

    def test_end_before_start_is_refused(self):
        period = ResidencyPeriod(
            account=self.account, status='RESIDENT',
            start_date=date(2020, 1, 1), end_date=date(2019, 1, 1))
        with self.assertRaises(ValidationError):
            period.full_clean()

    def test_overlapping_periods_are_refused(self):
        declare(self.account, 'RESIDENT', date(2015, 1, 1), date(2021, 6, 30))
        clash = ResidencyPeriod(
            account=self.account, status='FOREIGN',
            start_date=date(2020, 1, 1), end_date=date(2022, 1, 1))
        with self.assertRaises(ValidationError):
            clash.full_clean()

    def test_overlap_is_refused_even_without_a_form(self):
        """save() refuses overlapping periods, so imports are checked too."""
        declare(self.account, 'RESIDENT', date(2015, 1, 1), date(2021, 6, 30))
        with self.assertRaises(ValidationError):
            ResidencyPeriod.objects.create(
                account=self.account, status='FOREIGN',
                start_date=date(2020, 1, 1), end_date=date(2022, 1, 1))

    def test_a_gap_is_refused(self):
        declare(self.account, 'RESIDENT', date(2015, 1, 1), date(2021, 6, 30))
        later = ResidencyPeriod(
            account=self.account, status='FOREIGN', start_date=date(2021, 8, 1))
        with self.assertRaises(ValidationError):
            later.full_clean()

    def test_contiguous_periods_are_accepted(self):
        declare(self.account, 'RESIDENT', date(2015, 1, 1), date(2021, 6, 30))
        later = ResidencyPeriod(
            account=self.account, status='FOREIGN', start_date=date(2021, 7, 1))
        later.full_clean()
        later.save()
        self.assertEqual(ResidencyPeriod.objects.filter(account=self.account).count(), 2)

    def test_history_must_reach_back_to_the_earliest_purchase(self):
        """A residency history starting after the earliest purchase is refused."""
        Buy.objects.create(
            account=self.account, instrument=self.instrument, date=date(2010, 5, 1),
            quantity=Decimal('100'), unit_price=Money(Decimal('10'), 'AUD'),
            total_brokerage=Money(Decimal('9.95'), 'AUD'))
        late = ResidencyPeriod(
            account=self.account, status='RESIDENT', start_date=date(2015, 1, 1))
        with self.assertRaises(ValidationError):
            late.full_clean()

    def test_coverage_problems_reports_what_clean_would_have_refused(self):
        """coverage_problems reports histories that clean() would refuse."""
        declare(self.account, 'RESIDENT', date(2015, 1, 1), date(2021, 6, 30))
        ResidencyPeriod.objects.create(
            account=self.account, status='FOREIGN', start_date=date(2021, 8, 1))
        problems = cgt.residency.coverage_problems(self.account)
        self.assertEqual(len(problems), 1)
        self.assertIn('not declared between', problems[0])

    def test_a_sound_history_reports_no_problems(self):
        declare(self.account, 'RESIDENT', date(2015, 1, 1), date(2021, 6, 30))
        declare(self.account, 'FOREIGN', date(2021, 7, 1))
        self.assertEqual(cgt.residency.coverage_problems(self.account), [])


class ResidencyDayCountingTests(TransactionTestCase):
    """Residency day counting: inclusive endpoints, undeclared days counted separately."""

    def setUp(self):
        self.account = create_account()

    def test_both_endpoints_are_counted(self):
        declare(self.account, 'RESIDENT', date(2020, 1, 1))
        self.assertEqual(
            cgt.resident_days(self.account, date(2020, 1, 1), date(2020, 1, 1)), 1)
        self.assertEqual(
            cgt.resident_days(self.account, date(2020, 1, 1), date(2020, 1, 31)), 31)

    def test_days_are_split_between_statuses(self):
        declare(self.account, 'RESIDENT', date(2020, 1, 1), date(2020, 1, 10))
        declare(self.account, 'FOREIGN', date(2020, 1, 11))
        counts = cgt.residency.days_by_status(
            self.account, date(2020, 1, 1), date(2020, 1, 31))
        self.assertEqual(counts, {'RESIDENT': 10, 'FOREIGN': 21})

    def test_undeclared_days_are_counted_separately_not_ignored(self):
        """Days no period covers are counted under None."""
        declare(self.account, 'RESIDENT', date(2020, 1, 11), date(2020, 1, 20))
        counts = cgt.residency.days_by_status(
            self.account, date(2020, 1, 1), date(2020, 1, 31))
        self.assertEqual(counts, {'RESIDENT': 10, None: 21})

    def test_temporary_residents_count_with_foreign_residents(self):
        declare(self.account, 'TEMPORARY', date(2020, 1, 1), date(2020, 1, 10))
        declare(self.account, 'FOREIGN', date(2020, 1, 11), date(2020, 1, 20))
        declare(self.account, 'RESIDENT', date(2020, 1, 21))
        self.assertEqual(
            cgt.residency.non_resident_days(
                self.account, date(2020, 1, 1), date(2020, 1, 31)), 20)


class ResidencyInertForResidentsTests(TransactionTestCase):
    """Declaring unbroken Australian residency changes no figure."""

    def setUp(self):
        self.data = create_golden_master_portfolio()
        self.account = self.data['account']

    def test_declaring_residency_changes_no_figure(self):
        before = cgt.disposal_events(self.account)
        declare(self.account, 'RESIDENT', date(2000, 1, 1))
        after = cgt.disposal_events(self.account)

        self.assertEqual(len(before), 3)
        self.assertEqual(len(after), len(before))

        # Everything except how the answer was arrived at must be identical.
        annotations = {'residency_status', 'discount_basis', 'tap_status', 'is_disregarded',
                       'disregard_reason'}
        for old_event, new_event in zip(before, after):
            for field in cgt.event_fields():
                if field in annotations:
                    continue
                self.assertEqual(
                    getattr(old_event, field), getattr(new_event, field),
                    f'{field} moved when residency was declared')

    def test_the_discount_is_still_a_flat_half(self):
        declare(self.account, 'RESIDENT', date(2000, 1, 1))
        eligible = [e for e in cgt.disposal_events(self.account)
                    if e.method == cgt.events.METHOD_DISCOUNT]
        self.assertTrue(eligible)
        for event in eligible:
            self.assertEqual(event.discount_percentage, Decimal('0.5'))

    def test_the_basis_changes_from_assumed_to_declared(self):
        events = cgt.disposal_events(self.account)
        self.assertTrue(all(e.discount_basis == cgt.BASIS_LEGACY for e in events))
        declare(self.account, 'RESIDENT', date(2000, 1, 1))
        events = cgt.disposal_events(self.account)
        self.assertTrue(all(e.discount_basis == cgt.BASIS_DIVISION_115 for e in events))
        self.assertTrue(all(e.residency_status == 'RESIDENT' for e in events))

    def test_a_resident_disregards_nothing(self):
        declare(self.account, 'RESIDENT', date(2000, 1, 1))
        events = cgt.disposal_events(self.account)
        self.assertTrue(all(not e.is_disregarded for e in events))
        # Still not taxable Australian property. It is residency, not the asset, that makes
        # the gain assessable here.
        self.assertTrue(all(e.tap_status == cgt.tap.NTAP for e in events))


class UndeclaredDaysTests(TransactionTestCase):
    """A day no residency period covers counts as not resident, in every apportionment case.

    It used to depend on which formula the purchase date chose, and on whether some other
    day happened to be foreign.
    """

    def setUp(self):
        self.account = create_account()
        self.bought, self.sold = date(2010, 1, 1), date(2020, 1, 1)
        self.total = (self.sold - self.bought).days + 1
        self.gap = (date(2014, 12, 31) - date(2014, 1, 1)).days + 1

    def _fraction(self, counted):
        return (Decimal(counted) / Decimal(self.total)).quantize(Decimal('0.00000001'))

    def test_a_gap_counts_against_a_parcel_held_on_8_may_2012(self):
        declare(self.account, 'RESIDENT', date(2005, 1, 1), date(2013, 12, 31))
        declare(self.account, 'FOREIGN', date(2015, 1, 1), date(2015, 12, 31))
        declare(self.account, 'RESIDENT', date(2016, 1, 1))
        foreign = (date(2015, 12, 31) - date(2015, 1, 1)).days + 1

        self.assertEqual(
            cgt.apportionment_fraction(self.account, self.bought, self.sold),
            self._fraction(self.total - self.gap - foreign))

    def test_a_gap_alone_apportions_the_discount(self):
        declare(self.account, 'RESIDENT', date(2005, 1, 1), date(2013, 12, 31))
        declare(self.account, 'RESIDENT', date(2015, 1, 1))

        self.assertEqual(
            cgt.discount_percentage(self.bought, self.sold, account=self.account),
            Decimal('0.5') * self._fraction(self.total - self.gap))

    def test_a_complete_history_is_unaffected(self):
        declare(self.account, 'RESIDENT', date(2005, 1, 1))
        self.assertEqual(
            cgt.discount_percentage(self.bought, self.sold, account=self.account),
            Decimal('0.5'))


class ResidencyCoverageTests(TransactionTestCase):
    """What the residency history leaves out is reported rather than assumed."""

    def setUp(self):
        self.account = create_account()
        Buy.objects.create(
            account=self.account, instrument=create_instrument(account=self.account),
            date=date(2015, 1, 5), quantity=Decimal('100'),
            unit_price=Money(10, 'AUD'), total_brokerage=Money(0, 'AUD'),
        )

    def test_a_history_ending_before_today_is_reported(self):
        declare(self.account, 'RESIDENT', date(2010, 1, 1), date(2020, 12, 31))
        problems = cgt.residency.coverage_problems(self.account)
        self.assertTrue(any('only up to 2020-12-31' in p for p in problems), problems)

    def test_an_open_ended_history_is_not(self):
        declare(self.account, 'RESIDENT', date(2010, 1, 1))
        self.assertEqual(cgt.residency.coverage_problems(self.account), [])


class DiscountApportionmentTests(TransactionTestCase):
    """s115-115 apportionment in each of its three cases."""

    def setUp(self):
        self.account = create_account()

    def test_undeclared_residency_keeps_the_flat_half(self):
        self.assertEqual(
            cgt.discount_percentage(
                date(2015, 1, 1), date(2020, 1, 1), account=self.account),
            Decimal('0.5'))

    def test_acquired_after_8_may_2012_apportions_over_the_whole_holding(self):
        """s115-115(2): bought after 8 May 2012, every day of the holding is apportioned."""
        declare(self.account, 'RESIDENT', date(2015, 1, 1), date(2017, 6, 30))
        declare(self.account, 'FOREIGN', date(2017, 7, 1))
        self.assertEqual(
            cgt.apportionment_fraction(self.account, date(2015, 1, 1), date(2020, 1, 1)),
            Decimal('0.49917898'))
        self.assertEqual(
            cgt.discount_percentage(
                date(2015, 1, 1), date(2020, 1, 1), account=self.account),
            Decimal('0.249589490'))

    def test_resident_on_8_may_2012_keeps_the_earlier_years_whatever_happened(self):
        """s115-115(3): resident on 8 May 2012, earlier absences do not reduce the discount."""
        declare(self.account, 'RESIDENT', date(2005, 6, 1), date(2007, 12, 31))
        declare(self.account, 'FOREIGN', date(2008, 1, 1), date(2010, 12, 31))
        declare(self.account, 'RESIDENT', date(2011, 1, 1), date(2017, 6, 30))
        declare(self.account, 'FOREIGN', date(2017, 7, 1))

        self.assertEqual(
            cgt.apportionment_fraction(self.account, date(2005, 6, 1), date(2020, 1, 1)),
            Decimal('0.82826577'))
        self.assertEqual(
            cgt.discount_percentage(
                date(2005, 6, 1), date(2020, 1, 1), account=self.account),
            Decimal('0.414132885'))

    def test_already_abroad_on_8_may_2012_counts_only_later_resident_days(self):
        """s115-115(6): abroad on 8 May 2012, only later resident days count."""
        declare(self.account, 'RESIDENT', date(2010, 1, 1), date(2011, 12, 31))
        declare(self.account, 'FOREIGN', date(2012, 1, 1), date(2014, 12, 31))
        declare(self.account, 'RESIDENT', date(2015, 1, 1))
        self.assertEqual(
            cgt.apportionment_fraction(self.account, date(2010, 1, 1), date(2020, 1, 1)),
            Decimal('0.50013687'))

    def test_absence_entirely_before_8_may_2012_does_not_apportion_at_all(self):
        """Absence wholly before 8 May 2012 leaves the full discount."""
        declare(self.account, 'RESIDENT', date(2005, 1, 1), date(2007, 12, 31))
        declare(self.account, 'FOREIGN', date(2008, 1, 1), date(2010, 12, 31))
        declare(self.account, 'RESIDENT', date(2011, 1, 1))
        self.assertEqual(
            cgt.discount_percentage(
                date(2005, 1, 1), date(2020, 1, 1), account=self.account),
            Decimal('0.5'))

    def test_a_wholly_foreign_holding_gets_no_discount(self):
        declare(self.account, 'FOREIGN', date(2014, 1, 1))
        self.assertEqual(
            cgt.discount_percentage(
                date(2015, 1, 1), date(2020, 1, 1), account=self.account),
            Decimal('0'))

    def test_apportionment_never_rescues_a_holding_under_twelve_months(self):
        declare(self.account, 'RESIDENT', date(2015, 1, 1))
        self.assertEqual(
            cgt.discount_percentage(
                date(2019, 6, 1), date(2019, 12, 1), account=self.account),
            Decimal('0'))


class TaxpayerTypeDiscountTests(TransactionTestCase):
    """s115-10 and s115-100: the discount rate by taxpayer type."""

    def setUp(self):
        self.account = create_account()
        self.purchase = date(2015, 1, 1)
        self.sale = date(2020, 1, 1)

    def _rate(self):
        return cgt.discount_percentage(self.purchase, self.sale, account=self.account)

    def test_undeclared_keeps_the_behaviour_the_app_has_always_had(self):
        self.assertEqual(self.account.taxpayer_type, 'UNDECLARED')
        self.assertEqual(self._rate(), Decimal('0.5'))

    def test_an_individual_discounts_a_half(self):
        self.account.taxpayer_type = 'INDIVIDUAL'
        self.assertEqual(self._rate(), Decimal('0.5'))

    def test_a_complying_superannuation_fund_discounts_a_third(self):
        self.account.taxpayer_type = 'SMSF'
        self.assertEqual(self._rate(), Decimal(1) / Decimal(3))

    def test_a_company_gets_no_discount(self):
        """A company gets no discount (s115-10)."""
        self.account.taxpayer_type = 'COMPANY'
        self.assertEqual(self._rate(), Decimal('0'))

    def test_a_superannuation_funds_rate_is_not_apportioned_by_residency(self):
        """A super fund's rate is not apportioned by residency (s115-105)."""
        self.account.taxpayer_type = 'SMSF'
        declare(self.account, 'RESIDENT', date(2015, 1, 1), date(2017, 6, 30))
        declare(self.account, 'FOREIGN', date(2017, 7, 1))
        self.assertEqual(self._rate(), Decimal(1) / Decimal(3))

    def test_an_individuals_rate_is_apportioned(self):
        self.account.taxpayer_type = 'INDIVIDUAL'
        declare(self.account, 'RESIDENT', date(2015, 1, 1), date(2017, 6, 30))
        declare(self.account, 'FOREIGN', date(2017, 7, 1))
        self.assertEqual(self._rate(), Decimal('0.249589490'))


class ForeignResidentDisregardTests(TransactionTestCase):
    """s855-10, s768-915 and the s104-165(3) deeming."""

    def setUp(self):
        self.account = create_account()
        self.instrument = create_instrument(account=self.account)

    def _status(self, acquired, sold):
        return cgt.parcel_tap_status(self.account, self.instrument, acquired, sold)

    def test_listed_shares_are_not_taxable_australian_property(self):
        declare(self.account, 'RESIDENT', date(2010, 1, 1))
        self.assertEqual(self._status(date(2015, 1, 1), date(2020, 1, 1)), cgt.tap.NTAP)

    def test_undeclared_residency_leaves_the_question_unanswered(self):
        """Undeclared residency gives an unknown TAP status and nothing disregarded."""
        self.assertIsNone(self._status(date(2015, 1, 1), date(2020, 1, 1)))
        disregarded, reason = cgt.disregard(self.account, None, date(2020, 1, 1))
        self.assertFalse(disregarded)
        self.assertIsNone(reason)

    def test_a_foreign_resident_disregards_a_gain_on_listed_shares(self):
        declare(self.account, 'RESIDENT', date(2010, 1, 1), date(2017, 6, 30))
        declare(self.account, 'FOREIGN', date(2017, 7, 1))
        disregarded, reason = cgt.disregard(self.account, cgt.tap.NTAP, date(2020, 1, 1))
        self.assertTrue(disregarded)
        self.assertIn('s855-10', reason)

    def test_a_temporary_resident_disregards_it_too(self):
        declare(self.account, 'TEMPORARY', date(2010, 1, 1))
        disregarded, reason = cgt.disregard(self.account, cgt.tap.NTAP, date(2020, 1, 1))
        self.assertTrue(disregarded)
        self.assertIn('s768-915', reason)

    def test_a_resident_disregards_nothing(self):
        declare(self.account, 'RESIDENT', date(2010, 1, 1))
        disregarded, _ = cgt.disregard(self.account, cgt.tap.NTAP, date(2020, 1, 1))
        self.assertFalse(disregarded)

    def test_real_property_is_taxable_australian_property(self):
        declare(self.account, 'FOREIGN', date(2010, 1, 1))
        self.instrument.legal_form = 'REAL_PROPERTY'
        self.assertEqual(self._status(date(2015, 1, 1), date(2020, 1, 1)), cgt.tap.TAP)

    def test_the_instrument_override_wins(self):
        declare(self.account, 'FOREIGN', date(2010, 1, 1))
        self.instrument.is_taxable_australian_property_override = True
        self.assertEqual(self._status(date(2015, 1, 1), date(2020, 1, 1)), cgt.tap.TAP)
        disregarded, _ = cgt.disregard(self.account, cgt.tap.TAP, date(2020, 1, 1))
        self.assertFalse(disregarded)

    def test_a_false_override_suppresses_the_departure_deeming(self):
        """A "no" TAP override suppresses the departure deeming."""
        declare(self.account, 'RESIDENT', date(2000, 1, 1), end=date(2021, 6, 30))
        declare(self.account, 'FOREIGN', date(2021, 7, 1), i1=True)

        held_at_departure = (date(2015, 1, 1), date(2024, 1, 1))
        self.assertEqual(self._status(*held_at_departure), cgt.tap.TAP)

        self.instrument.is_taxable_australian_property_override = False
        self.assertEqual(self._status(*held_at_departure), cgt.tap.NTAP)
        disregarded, _ = cgt.disregard(self.account, cgt.tap.NTAP, date(2024, 1, 1))
        self.assertTrue(disregarded)


class I1ElectionDeemingTests(TransactionTestCase):
    """s104-165(3): an I1 election keeps parcels held on departure inside the CGT net."""

    def setUp(self):
        self.account = create_account()
        self.instrument = create_instrument(account=self.account)
        declare(self.account, 'RESIDENT', date(2010, 1, 1), date(2017, 6, 30))
        declare(self.account, 'FOREIGN', date(2017, 7, 1), i1=True)

    def _status(self, acquired, sold=date(2020, 1, 1)):
        return cgt.parcel_tap_status(self.account, self.instrument, acquired, sold)

    def test_a_parcel_held_at_departure_is_deemed_taxable_australian_property(self):
        self.assertEqual(self._status(date(2015, 1, 1)), cgt.tap.TAP)

    def test_a_parcel_bought_after_departure_is_not(self):
        self.assertEqual(self._status(date(2018, 1, 1)), cgt.tap.NTAP)

    def test_the_two_coexist_for_one_instrument(self):
        """Two parcels of one instrument can get different TAP statuses."""
        self.assertEqual(self._status(date(2015, 1, 1)), cgt.tap.TAP)
        self.assertEqual(self._status(date(2018, 1, 1)), cgt.tap.NTAP)

    def test_without_the_election_there_is_no_deeming(self):
        """Without an I1 election there is no deeming."""
        ResidencyPeriod.objects.filter(account=self.account, status='FOREIGN').update(
            i1_election_made=False)
        self.assertEqual(self._status(date(2015, 1, 1)), cgt.tap.NTAP)

    def test_the_deeming_lapses_on_returning_to_australia(self):
        """The deeming lapses on resuming residency (s104-165(3))."""
        ResidencyPeriod.objects.filter(account=self.account, status='FOREIGN').update(
            end_date=date(2019, 6, 30))
        declare(self.account, 'RESIDENT', date(2019, 7, 1))
        self.assertEqual(self._status(date(2015, 1, 1), date(2020, 1, 1)), cgt.tap.NTAP)
        # Sold while still abroad, it would have been caught.
        self.assertEqual(self._status(date(2015, 1, 1), date(2018, 1, 1)), cgt.tap.TAP)

    def _suppressed(self, acquired, sold=date(2020, 1, 1)):
        return cgt.tap.override_suppresses_deeming(
            self.account, self.instrument, acquired, sold)

    def test_a_no_override_on_a_deemed_parcel_is_reported_as_suppressing_it(self):
        """A "no" override on a deemed parcel is reported as suppressing the deeming."""
        self.instrument.is_taxable_australian_property_override = False
        self.assertTrue(self._suppressed(date(2015, 1, 1)))

    def test_an_override_agreeing_with_the_derivation_is_not_reported(self):
        """An override that matches the derivation is not reported."""
        self.instrument.is_taxable_australian_property_override = False
        self.assertFalse(self._suppressed(date(2018, 1, 1)))

    def test_an_unset_override_is_not_reported(self):
        self.assertFalse(self._suppressed(date(2015, 1, 1)))

    def test_a_yes_override_is_not_reported(self):
        """A "yes" override is not reported."""
        self.instrument.is_taxable_australian_property_override = True
        self.assertFalse(self._suppressed(date(2015, 1, 1)))


class AttributionDisregardTests(TransactionTestCase):
    """s855-40(2), s276-55: a foreign resident member's non-TAP attributions are disregarded."""

    def setUp(self):
        self.account = create_account()
        self.instrument = create_instrument(account=self.account, name='VGS')
        self.statement = AttributionStatement.objects.create(
            account=self.account, instrument=self.instrument,
            financial_year_end_date=date(2025, 6, 30))

    def _component(self, component, amount):
        AttributionComponent.objects.create(
            account=self.account, statement=self.statement,
            component=component, amount=Money(Decimal(amount), 'AUD'))

    def test_a_foreign_resident_disregards_a_non_tap_attribution(self):
        self._component('DISCOUNTED_NTAP', '4845.57')
        declare(self.account, 'FOREIGN', date(2010, 1, 1))
        events = cgt.attribution_events(self.account)
        self.assertEqual(len(events), 1)
        event = events[0]
        self.assertEqual(event.tap_status, cgt.tap.NTAP)
        self.assertTrue(event.is_disregarded)
        self.assertIn('s855-40(2)', event.disregard_reason)
        # Still grossed up and still reported, so the figure is explainable rather than
        # simply missing from the schedule.
        self.assertEqual(event.capital_gain, Money(Decimal('9691.14'), 'AUD'))

    def test_a_tap_attribution_stays_assessable(self):
        self._component('DISCOUNTED_TAP', '100.00')
        declare(self.account, 'FOREIGN', date(2010, 1, 1))
        event = cgt.attribution_events(self.account)[0]
        self.assertEqual(event.tap_status, cgt.tap.TAP)
        self.assertFalse(event.is_disregarded)

    def test_a_resident_disregards_nothing(self):
        self._component('DISCOUNTED_NTAP', '4845.57')
        declare(self.account, 'RESIDENT', date(2010, 1, 1))
        self.assertFalse(cgt.attribution_events(self.account)[0].is_disregarded)

    def test_undeclared_residency_disregards_nothing(self):
        """Undeclared residency disregards nothing."""
        self._component('DISCOUNTED_NTAP', '4845.57')
        self.assertFalse(cgt.attribution_events(self.account)[0].is_disregarded)

    def test_the_two_halves_of_a_mixed_statement_are_treated_separately(self):
        """s855-40(2): a foreign resident's NTAP part is disregarded, the TAP part is not.

        The statement states each amount, so nothing is apportioned on a guessed ratio.
        """
        self._component('DISCOUNTED_TAP', '100.00')
        self._component('DISCOUNTED_NTAP', '900.00')
        self._component('OTHER_NTAP', '50.00')
        declare(self.account, 'FOREIGN', date(2010, 1, 1))
        events = cgt.attribution_events(self.account)
        self.assertEqual(
            [(e.method, e.tap_status, e.capital_gain.amount, e.is_disregarded) for e in events],
            [('discount', cgt.tap.TAP, Decimal('200.00'), False),
             ('discount', cgt.tap.NTAP, Decimal('1800.00'), True),
             ('other', cgt.tap.NTAP, Decimal('50.00'), True)])


class TaxSettingsBannerTests(TransactionTestCase):
    """The dashboard's tax settings warning."""

    def setUp(self):
        from share_dinkum_app.dashboard import _tax_settings_warning
        self._warning = _tax_settings_warning

    def test_an_account_with_no_sales_is_left_alone(self):
        account = create_account()
        self.assertIsNone(self._warning(account))

    def test_an_account_with_sales_and_no_declaration_is_warned(self):
        account = create_golden_master_portfolio()['account']
        warning = self._warning(account)
        self.assertIsNotNone(warning)
        self.assertIn('flat 50%', warning)

    def test_declaring_everything_silences_it(self):
        account = create_golden_master_portfolio()['account']
        account.taxpayer_type = 'INDIVIDUAL'
        account.save()
        declare(account, 'RESIDENT', date(2000, 1, 1))
        self.assertIsNone(self._warning(account))

    def test_reviewing_the_settings_silences_it_whatever_was_chosen(self):
        """Setting tax_settings_reviewed_at silences the warning, whatever was chosen."""
        from django.utils import timezone
        account = create_golden_master_portfolio()['account']
        account.tax_settings_reviewed_at = timezone.now()
        account.save()
        self.assertIsNone(self._warning(account))


# --- Phase 5: the 2027 regime ---

def load_test_cpi(quarters=None):
    """Load CPI rising in round steps from the first indexable quarter."""
    quarters = quarters or {
        date(2027, 7, 1): '100.0',
        date(2027, 10, 1): '102.5',
        date(2028, 1, 1): '105.0',
        date(2028, 4, 1): '107.5',
        date(2028, 7, 1): '110.0',
        date(2029, 1, 1): '115.0',
        date(2030, 1, 1): '125.0',
    }
    for quarter_start, index_number in quarters.items():
        CPIIndex.objects.update_or_create(
            quarter_start_date=quarter_start,
            defaults={'index_number': Decimal(index_number), 'source': 'test'})


def enable_2027_regime(test_case, account=None):
    """Turn on `model_2027_regime` for the test's account."""
    account = account or getattr(test_case, 'account', None)
    if account is None:
        raise AssertionError('enable_2027_regime needs an account to turn the setting on')
    account.model_2027_regime = True
    account.save()
    return account


class IndexationFactorTests(TestCase):
    """Subdivision 960-M indexation factors, and missing CPI quarters."""

    def setUp(self):
        load_test_cpi()

    def test_the_factor_is_the_ratio_of_two_quarters(self):
        # Bought at the cutover, sold in the September 2028 quarter: 110.0 over 100.0.
        self.assertEqual(
            cgt.indexation_factor(date(2027, 7, 1), date(2028, 8, 20)),
            Decimal('1.100'))

    def test_indexation_never_reaches_back_before_the_cutover(self):
        """Indexation runs only from the quarter starting 1 July 2027 (s960-275(1B))."""
        self.assertEqual(
            cgt.indexation_factor(date(2010, 1, 1), date(2028, 8, 20)),
            Decimal('1.100'))

    def test_the_factor_is_three_decimal_places(self):
        factor = cgt.indexation_factor(date(2027, 7, 1), date(2027, 11, 1))
        self.assertEqual(factor, Decimal('1.025'))
        self.assertEqual(factor.as_tuple().exponent, -3)

    def test_a_falling_index_does_not_shrink_a_cost_base(self):
        CPIIndex.objects.create(
            quarter_start_date=date(2031, 1, 1), index_number=Decimal('90.0'))
        self.assertEqual(
            cgt.indexation_factor(date(2027, 7, 1), date(2031, 2, 1)),
            Decimal('1.000'))

    def test_a_missing_quarter_raises_rather_than_defaulting(self):
        """A missing CPI quarter raises rather than defaulting."""
        with self.assertRaises(cgt.IndexationDataUnavailable) as caught:
            cgt.indexation_factor(date(2027, 7, 1), date(2033, 5, 1))
        self.assertIn('2033-04-01', str(caught.exception))

    def test_an_empty_table_raises_too(self):
        CPIIndex.objects.all().delete()
        with self.assertRaises(cgt.IndexationDataUnavailable):
            cgt.indexation_factor(date(2027, 7, 1), date(2028, 8, 20))


class IndexationEligibilityTests(TransactionTestCase):
    """s114-25: the testing period runs from the later of the cutover and acquisition."""

    def setUp(self):
        self.account = create_account()
        load_test_cpi()

    def test_past_non_residency_does_not_disqualify(self):
        """Non-residency before 1 July 2027 does not affect indexation."""
        declare(self.account, 'RESIDENT', date(2010, 1, 1), date(2020, 12, 31))
        declare(self.account, 'FOREIGN', date(2021, 1, 1), date(2026, 6, 30))
        declare(self.account, 'RESIDENT', date(2026, 7, 1))
        self.assertTrue(cgt.is_indexation_eligible(
            self.account, date(2020, 1, 15), date(2030, 1, 15)))

    def test_one_week_abroad_inside_the_testing_period_denies_it_entirely(self):
        """One week abroad in the testing period denies indexation entirely."""
        declare(self.account, 'RESIDENT', date(2010, 1, 1), date(2020, 12, 31))
        declare(self.account, 'FOREIGN', date(2021, 1, 1), date(2026, 6, 30))
        declare(self.account, 'RESIDENT', date(2026, 7, 1), date(2028, 2, 29))
        declare(self.account, 'FOREIGN', date(2028, 3, 1), date(2028, 3, 7))
        declare(self.account, 'RESIDENT', date(2028, 3, 8))
        self.assertFalse(cgt.is_indexation_eligible(
            self.account, date(2020, 1, 15), date(2030, 1, 15)))

    def test_a_disposal_before_the_cutover_is_never_indexed(self):
        declare(self.account, 'RESIDENT', date(2010, 1, 1))
        self.assertFalse(cgt.is_indexation_eligible(
            self.account, date(2020, 1, 15), date(2026, 1, 15)))

    def test_undeclared_residency_does_not_get_indexation_on_an_assumption(self):
        """An account with undeclared residency is not eligible for indexation."""
        self.assertFalse(cgt.is_indexation_eligible(
            self.account, date(2020, 1, 15), date(2030, 1, 15)))

    def test_a_company_or_super_fund_is_not_indexed(self):
        """s110-36(1A) indexes only for Australian resident individuals and trusts."""
        declare(self.account, 'RESIDENT', date(2010, 1, 1))
        for taxpayer_type in ('COMPANY', 'SMSF'):
            self.account.taxpayer_type = taxpayer_type
            self.account.save()
            self.assertFalse(cgt.is_indexation_eligible(
                self.account, date(2020, 1, 15), date(2030, 1, 15)), taxpayer_type)

    def test_a_trust_is_indexed(self):
        declare(self.account, 'RESIDENT', date(2010, 1, 1))
        self.account.taxpayer_type = 'TRUST'
        self.account.save()
        self.assertTrue(cgt.is_indexation_eligible(
            self.account, date(2020, 1, 15), date(2030, 1, 15)))

    def test_a_super_fund_keeps_its_third_after_the_cutover(self):
        """s115-100(b) is unchanged, so a complying fund's gain is discounted, not indexed."""
        data = create_cutover_portfolio(suffix='smsf')
        declare(data['account'], 'RESIDENT', date(2010, 1, 1))
        data['account'].taxpayer_type = 'SMSF'
        enable_2027_regime(self, data['account'])

        event = cgt.disposal_events(data['account'])[0]

        self.assertEqual(event.discount_percentage, Decimal(1) / Decimal(3))
        self.assertEqual(event.method, cgt.events.METHOD_DISCOUNT)


def create_cutover_portfolio(sell_date=date(2028, 8, 20), unit_value='15.00', suffix='',
                             adjustments=()):
    """A parcel bought well before the cutover and sold well after it.

    Cost base 10,019.95, cutover value 15,000, net proceeds 19,990.05. `suffix` allows a
    second, independent portfolio in the same test. `adjustments` are
    `(financial_year_end_date, amount)` pairs, entered before the sale.
    """
    account = create_account(
        owner=create_user(username=f'cutover{suffix}'),
        description=f'Cutover Account {suffix}'.strip(),
        fy_type=create_fiscal_year_type(description=f'AU Tax Year {suffix}'.strip()),
    )
    instrument = create_instrument(account=account, name='CUT')

    buy = Buy.objects.create(
        account=account, instrument=instrument, date=date(2020, 1, 15),
        quantity=Decimal('1000'), unit_price=Money(Decimal('10.00'), 'AUD'),
        total_brokerage=Money(Decimal('19.95'), 'AUD'),
    )
    InstrumentValuation.objects.create(
        account=account, instrument=instrument,
        valuation_date=date(2027, 6, 30),
        unit_value=Money(Decimal(unit_value), 'AUD'),
        purpose='CUTOVER_2027', source='USER',
    )
    for year_end, amount in adjustments:
        CostBaseAdjustment.objects.create(
            account=account, instrument=instrument, financial_year_end_date=year_end,
            cost_base_increase=Money(Decimal(amount), 'AUD'),
        )
    sell = Sell.objects.create(
        account=account, instrument=instrument, date=sell_date,
        quantity=Decimal('1000'), unit_price=Money(Decimal('20.00'), 'AUD'),
        total_brokerage=Money(Decimal('9.95'), 'AUD'), strategy='FIFO',
    )
    load_test_cpi()
    return {'account': account, 'instrument': instrument, 'buy': buy, 'sell': sell}


class DeemedSaleSplitTests(TransactionTestCase):
    """s112-155: a straddling disposal splits into deferred and post-cutover gains."""

    def setUp(self):
        self.data = create_cutover_portfolio()
        self.account = self.data['account']
        declare(self.account, 'RESIDENT', date(2010, 1, 1))
        enable_2027_regime(self)

    def _events(self):
        return cgt.disposal_events(self.account)

    def test_one_disposal_becomes_two_rows(self):
        events = self._events()
        self.assertEqual(len(events), 2)
        deferred, post = events
        self.assertEqual(deferred.slice, cgt.events.SLICE_PRE_CUTOVER)
        self.assertEqual(post.slice, cgt.events.SLICE_POST_CUTOVER)
        # One sale, so a report can still present them together.
        self.assertEqual(deferred.sell_allocation_id, post.sell_allocation_id)

    def test_the_deferred_gain_is_measured_to_the_cutover_valuation(self):
        deferred = self._events()[0]
        self.assertEqual(deferred.net_proceeds, Money(Decimal('15000.00'), 'AUD'))
        self.assertEqual(deferred.cost_base, Money(Decimal('10019.95'), 'AUD'))
        self.assertEqual(deferred.capital_gain, Money(Decimal('4980.05'), 'AUD'))

    def test_the_deferred_gain_keeps_the_discount(self):
        deferred = self._events()[0]
        self.assertEqual(deferred.method, cgt.events.METHOD_DISCOUNT)
        self.assertEqual(deferred.discount_percentage, Decimal('0.5'))
        self.assertEqual(
            deferred.gain_category, constants.CGT_GAIN_DEFERRED_NON_RESIDENTIAL)

    def test_the_post_cutover_gain_is_indexed_and_undiscounted(self):
        """The post-cutover slice is indexed and gets no discount (s110-36(1A), s115-20)."""
        post = self._events()[1]
        self.assertEqual(post.indexation_factor, Decimal('1.100'))
        # 15,000 reacquisition cost lifted by 10% inflation.
        self.assertEqual(post.cost_base, Money(Decimal('16500.000'), 'AUD'))
        self.assertEqual(post.capital_gain, Money(Decimal('3490.050'), 'AUD'))
        self.assertEqual(post.discount_percentage, Decimal('0'))
        self.assertEqual(post.method, cgt.events.METHOD_OTHER)
        self.assertEqual(post.gain_category, constants.CGT_GAIN_NON_RESIDENTIAL)

    def test_the_two_slices_sum_to_the_whole_gain_before_indexation_relief(self):
        """The two cutover slices sum to the whole gain, less the indexation relief."""
        deferred, post = self._events()
        whole_gain = Decimal('19990.05') - Decimal('10019.95')
        relief = Decimal('16500.000') - Decimal('15000.00')
        self.assertEqual(
            deferred.capital_gain.amount + post.capital_gain.amount,
            whole_gain - relief)

    def test_the_twelve_month_rule_ignores_the_deemed_reacquisition(self):
        """s114-10(9): the deferred slice's 12-month test runs from the original purchase."""
        data = create_cutover_portfolio(sell_date=date(2027, 8, 20), suffix='b')
        declare(data['account'], 'RESIDENT', date(2010, 1, 1))
        deferred = cgt.disposal_events(data['account'])[0]
        self.assertEqual(deferred.method, cgt.events.METHOD_DISCOUNT)
        self.assertEqual(deferred.discount_percentage, Decimal('0.5'))

    @patch('share_dinkum_app.models.yfinanceinterface.get_exchange_rate',
           return_value=Decimal('1.5'))
    def test_a_foreign_valuation_is_converted_at_the_cutover_rate(self, mock_get_rate):
        """A valuation in another currency is converted, not subtracted from an AUD cost base."""
        InstrumentValuation.objects.filter(account=self.account).update(
            unit_value=Decimal('10.00'), unit_value_currency='USD')

        deferred = self._events()[0]

        # 1,000 units at US$10, converted at 1.5.
        self.assertEqual(deferred.net_proceeds, Money(Decimal('15000.00'), 'AUD'))

    def test_a_split_after_the_sale_does_not_scale_the_cutover_value(self):
        """The parcel was sold before the split, so its units never changed."""
        ShareSplit.objects.create(
            account=self.account, instrument=self.data['instrument'], date=date(2029, 1, 1),
            quantity_before=Decimal('1'), quantity_after=Decimal('2'),
        )

        deferred = self._events()[0]

        self.assertEqual(deferred.net_proceeds, Money(Decimal('15000.00'), 'AUD'))

    def test_without_a_valuation_the_disposal_is_not_split(self):
        InstrumentValuation.objects.all().delete()
        events = self._events()
        self.assertEqual(len(events), 1)
        self.assertEqual(events[0].slice, cgt.events.SLICE_WHOLE)
        self.assertIn('No market value recorded', events[0].pending_reason)

    def test_a_missing_cpi_quarter_leaves_the_row_pending(self):
        CPIIndex.objects.all().delete()
        post = self._events()[1]
        self.assertIsNone(post.indexation_factor)
        self.assertIn('No CPI index number', post.pending_reason)
        # And no relief is given: the cost base is the unindexed reacquisition cost.
        self.assertEqual(post.cost_base, Money(Decimal('15000.00'), 'AUD'))

    def test_the_rollout_gate_reverts_to_a_single_unsplit_row(self):
        self.account.model_2027_regime = False
        self.account.save()
        events = self._events()
        self.assertEqual(len(events), 1)
        self.assertEqual(events[0].slice, cgt.events.SLICE_WHOLE)
        self.assertEqual(events[0].capital_gain, Money(Decimal('9970.10'), 'AUD'))


class CutoverAdjustmentSliceTests(TransactionTestCase):
    """Adjustments for years after the cutover belong to the post-cutover slice."""

    def _events(self, year_end, amount):
        data = create_cutover_portfolio(adjustments=[(year_end, amount)])
        declare(data['account'], 'RESIDENT', date(2010, 1, 1))
        enable_2027_regime(self, data['account'])
        return cgt.disposal_events(data['account'])

    def test_a_later_increase_moves_to_the_post_cutover_slice(self):
        deferred, post = self._events(date(2028, 6, 30), '1000')

        # The deferred slice is as if there were no adjustment.
        self.assertEqual(deferred.cost_base, Money(Decimal('10019.95'), 'AUD'))
        self.assertEqual(deferred.cost_base_adjustments, Money(Decimal('0'), 'AUD'))
        self.assertEqual(deferred.capital_gain, Money(Decimal('4980.05'), 'AUD'))
        # 15,000 lifted by 10% inflation from July 2027, plus 1,000 lifted from the quarter
        # holding 30 June 2028 (110.0 / 107.5 = 1.023).
        self.assertEqual(post.buy_consideration, Money(Decimal('15000.00'), 'AUD'))
        self.assertEqual(post.cost_base_adjustments, Money(Decimal('1000.00'), 'AUD'))
        self.assertEqual(post.cost_base, Money(Decimal('17523.000'), 'AUD'))
        self.assertEqual(post.capital_gain, Money(Decimal('2467.050'), 'AUD'))
        # The factor reported is the market value's.
        self.assertEqual(post.indexation_factor, Decimal('1.100'))

    def test_a_later_decrease_moves_to_the_post_cutover_slice(self):
        """A decrease, common for property trusts, raises the undiscounted gain."""
        deferred, post = self._events(date(2028, 6, 30), '-500')

        self.assertEqual(deferred.cost_base, Money(Decimal('10019.95'), 'AUD'))
        self.assertEqual(deferred.capital_gain, Money(Decimal('4980.05'), 'AUD'))
        self.assertEqual(post.cost_base_adjustments, Money(Decimal('-500.00'), 'AUD'))
        # s114-15(3): the decrease takes off only the indexation from its own quarter.
        # 16,500 less 500 x 1.023.
        self.assertEqual(post.cost_base, Money(Decimal('15988.500'), 'AUD'))
        self.assertEqual(post.capital_gain, Money(Decimal('4001.550'), 'AUD'))

    def test_without_indexing_increases_an_increase_is_at_face_value(self):
        """Interpretation 2: s114-15(2) does not reach a total cost base increase."""
        with patch.object(constants, 'CGT_INDEX_COST_BASE_INCREASES', False):
            _deferred, post = self._events(date(2028, 6, 30), '1000')

        self.assertEqual(post.cost_base, Money(Decimal('17500.000'), 'AUD'))
        self.assertEqual(post.capital_gain, Money(Decimal('2490.050'), 'AUD'))

    def test_without_indexing_increases_a_decrease_is_still_indexed(self):
        """s114-15(3) expressly covers a reduction of the total cost base."""
        with patch.object(constants, 'CGT_INDEX_COST_BASE_INCREASES', False):
            _deferred, post = self._events(date(2028, 6, 30), '-500')

        self.assertEqual(post.cost_base, Money(Decimal('15988.500'), 'AUD'))

    def test_an_adjustment_for_the_year_of_the_sale_is_made_at_the_sale(self):
        """s104-107B(4)(b): sold before the year ends, so the adjustment is not indexed."""
        data = create_cutover_portfolio()
        CostBaseAdjustment.objects.create(
            account=data['account'], instrument=data['instrument'],
            financial_year_end_date=date(2029, 6, 30),
            cost_base_increase=Money(Decimal('1000'), 'AUD'),
        )
        declare(data['account'], 'RESIDENT', date(2010, 1, 1))
        enable_2027_regime(self, data['account'])

        _deferred, post = cgt.disposal_events(data['account'])

        # 16,500 plus the parcel's share of the adjustment at a factor of 1.000.
        adjustment = post.cost_base_adjustments.amount
        self.assertGreater(adjustment, Decimal('0'))
        self.assertEqual(
            post.cost_base, Money(Decimal('16500.000') + adjustment, 'AUD'))

    def test_an_adjustment_for_the_year_ending_at_the_cutover_stays_deferred(self):
        """s104-107B applies it at the end of that year, the day of the deemed sale."""
        deferred, post = self._events(date(2027, 6, 30), '1000')

        self.assertEqual(deferred.cost_base_adjustments, Money(Decimal('1000.00'), 'AUD'))
        self.assertEqual(deferred.cost_base, Money(Decimal('11019.95'), 'AUD'))
        self.assertEqual(deferred.capital_gain, Money(Decimal('3980.05'), 'AUD'))
        self.assertEqual(post.cost_base_adjustments, Money(Decimal('0'), 'AUD'))
        self.assertEqual(post.cost_base, Money(Decimal('16500.000'), 'AUD'))
        self.assertEqual(post.capital_gain, Money(Decimal('3490.050'), 'AUD'))


class ReturnedExpatIndexationTests(TransactionTestCase):
    """A returned expatriate gets no deemed sale and no discount, only indexation from 2027.

    s112-155(1)(d) denies the split; s114-25 then allows indexation, and s115-20 removes
    the discount.
    """

    def setUp(self):
        self.data = create_cutover_portfolio()
        self.account = self.data['account']
        declare(self.account, 'RESIDENT', date(2010, 1, 1), date(2014, 12, 31))
        declare(self.account, 'FOREIGN', date(2015, 1, 1), date(2020, 12, 31))
        declare(self.account, 'RESIDENT', date(2021, 1, 1))
        enable_2027_regime(self)

    def test_there_is_no_deemed_sale(self):
        applies, reason = cgt.deemed_sale_applies(
            self.account, date(2020, 1, 15), date(2028, 8, 20))
        self.assertFalse(applies)
        self.assertIn('s112-155(1)(d)', reason)

    def test_the_gain_stays_a_single_row(self):
        events = cgt.disposal_events(self.account)
        self.assertEqual(len(events), 1)
        self.assertEqual(events[0].slice, cgt.events.SLICE_WHOLE)

    def test_indexation_applies_and_takes_the_discount_with_it(self):
        event = cgt.disposal_events(self.account)[0]
        self.assertEqual(event.indexation_factor, Decimal('1.100'))
        # The whole original cost base indexed, not a 2027 market value.
        self.assertEqual(event.cost_base, Money(Decimal('11021.945'), 'AUD'))
        self.assertEqual(event.discount_percentage, Decimal('0'))
        self.assertEqual(event.method, cgt.events.METHOD_OTHER)

    def test_each_adjustment_is_indexed_from_its_own_quarter(self):
        """One before the cutover is indexed from July 2027, one after from its year end."""
        data = create_cutover_portfolio(suffix='adj', adjustments=[
            (date(2026, 6, 30), '200'), (date(2028, 6, 30), '-500')])
        declare(data['account'], 'RESIDENT', date(2010, 1, 1), date(2014, 12, 31))
        declare(data['account'], 'FOREIGN', date(2015, 1, 1), date(2020, 12, 31))
        declare(data['account'], 'RESIDENT', date(2021, 1, 1))
        enable_2027_regime(self, data['account'])

        event = cgt.disposal_events(data['account'])[0]

        self.assertEqual(event.indexation_factor, Decimal('1.100'))
        # (10,019.95 + 200) x 1.100 - 500 x 1.023.
        self.assertEqual(event.cost_base, Money(Decimal('10730.445'), 'AUD'))

    def test_the_report_explains_it_rather_than_leaving_it_to_be_discovered(self):
        event = cgt.disposal_events(self.account)[0]
        self.assertIn('s115-105 applies', event.pending_reason)

    def test_the_explanation_says_what_was_done(self):
        """Indexed with no discount, not the apportioned discount the other case gets."""
        reason = cgt.disposal_events(self.account)[0].pending_reason
        self.assertIn('indexed from then', reason)
        self.assertNotIn('no indexation', reason)

    def test_returning_after_the_cutover_keeps_an_apportioned_discount_instead(self):
        """Not resident every day since 1 July 2027, so s114-25 denies indexation."""
        data = create_cutover_portfolio(suffix='late')
        declare(data['account'], 'RESIDENT', date(2010, 1, 1), date(2014, 12, 31))
        declare(data['account'], 'FOREIGN', date(2015, 1, 1), date(2027, 12, 31))
        declare(data['account'], 'RESIDENT', date(2028, 1, 1))
        enable_2027_regime(self, data['account'])

        event = cgt.disposal_events(data['account'])[0]

        self.assertGreater(event.discount_percentage, Decimal('0'))
        self.assertIn('no indexation', event.pending_reason)

    def test_they_are_worse_off_than_if_they_had_never_left(self):
        """A returned expatriate is worse off than if they had never left."""
        caught = cgt.disposal_events(self.account)[0]

        never_left = create_cutover_portfolio(suffix='b')
        declare(never_left['account'], 'RESIDENT', date(2010, 1, 1))
        # The setting is per portfolio now, so a second one built here needs it too. The
        # module constant it replaced covered every account at once.
        enable_2027_regime(self, never_left['account'])
        deferred, post = cgt.disposal_events(never_left['account'])

        taxed_if_caught = caught.capital_gain.amount
        taxed_if_not = (
            deferred.capital_gain.amount * (Decimal('1') - deferred.discount_percentage)
            + post.capital_gain.amount)
        self.assertGreater(taxed_if_caught, taxed_if_not)


class CutoverValuationTests(TransactionTestCase):
    """Cutover valuations are per unit and survive share splits."""

    def setUp(self):
        self.data = create_cutover_portfolio()
        self.account = self.data['account']

    def test_a_parcel_is_valued_from_the_per_unit_figure(self):
        parcel = Parcel.objects.filter(account=self.account).first()
        value, source = parcel.market_value_at(date(2027, 6, 30))
        self.assertEqual(value, Money(Decimal('15000.00'), 'AUD'))
        self.assertEqual(source, 'USER')

    def test_a_later_split_does_not_double_the_valuation(self):
        """A later share split does not change a parcel's cutover valuation."""
        held = create_instrument(
            account=self.account, market=self.data['instrument'].market, name='HELD')
        Buy.objects.create(
            account=self.account, instrument=held, date=date(2020, 1, 15),
            quantity=Decimal('1000'), unit_price=Money(Decimal('10.00'), 'AUD'),
            total_brokerage=Money(Decimal('0'), 'AUD'),
        )
        InstrumentValuation.objects.create(
            account=self.account, instrument=held, valuation_date=date(2027, 6, 30),
            unit_value=Money(Decimal('15.00'), 'AUD'), purpose='CUTOVER_2027',
            source='USER',
        )
        ShareSplit.objects.create(
            account=self.account, instrument=held, date=date(2027, 9, 1),
            quantity_before=Decimal('1'), quantity_after=Decimal('2'),
        )

        parcel = [p for p in Parcel.objects.filter(account=self.account, buy__instrument=held)
                  if p.remaining_quantity][0]
        self.assertEqual(parcel.parcel_quantity, Decimal('2000'))

        # 2,000 units at a value recorded when a unit was worth twice as much.
        value, _source = parcel.market_value_at(date(2027, 6, 30))
        self.assertEqual(value, Money(Decimal('15000.00'), 'AUD'))

    def test_a_price_history_close_is_used_where_no_valuation_was_recorded(self):
        InstrumentValuation.objects.all().delete()
        InstrumentPriceHistory.objects.create(
            account=self.account, instrument=self.data['instrument'],
            date=date(2027, 6, 30), open=Decimal('14'), high=Decimal('16'),
            low=Decimal('14'), close=Decimal('15.5'), volume=1000,
            stock_splits=Decimal('0'),
        )
        parcel = Parcel.objects.filter(account=self.account).first()
        value, source = parcel.market_value_at(date(2027, 6, 30))
        self.assertEqual(value, Money(Decimal('15500.0'), 'AUD'))
        self.assertEqual(source, 'PRICE_HISTORY')

    def test_a_recorded_valuation_beats_a_closing_price(self):
        """A recorded valuation takes precedence over a closing price."""
        InstrumentPriceHistory.objects.create(
            account=self.account, instrument=self.data['instrument'],
            date=date(2027, 6, 30), open=Decimal('14'), high=Decimal('16'),
            low=Decimal('14'), close=Decimal('15.5'), volume=1000,
            stock_splits=Decimal('0'),
        )
        parcel = Parcel.objects.filter(account=self.account).first()
        value, source = parcel.market_value_at(date(2027, 6, 30))
        self.assertEqual(value, Money(Decimal('15000.00'), 'AUD'))
        self.assertEqual(source, 'USER')

    def test_nothing_is_invented_where_there_is_no_price(self):
        InstrumentValuation.objects.all().delete()
        parcel = Parcel.objects.filter(account=self.account).first()
        self.assertEqual(parcel.market_value_at(date(2027, 6, 30)), (None, None))


class DeemedResetDateTests(TransactionTestCase):
    """The dates on which a cost base is deemed reset to market value."""

    def setUp(self):
        self.account = create_account()

    def test_the_cutover_is_always_a_reset_date(self):
        self.assertEqual(
            cgt.deemed_reset_dates(self.account),
            [(date(2027, 7, 1), 'CUTOVER_2027')])

    def test_leaving_australia_without_the_election_adds_one(self):
        """A departure without an I1 election adds a reset date (s104-165)."""
        declare(self.account, 'RESIDENT', date(2010, 1, 1), date(2017, 6, 30))
        declare(self.account, 'FOREIGN', date(2017, 7, 1), i1=False)
        self.assertIn((date(2017, 7, 1), 'DEPARTURE'), cgt.deemed_reset_dates(self.account))

    def test_the_i1_election_means_there_is_no_reset(self):
        """A departure with an I1 election adds no reset date."""
        declare(self.account, 'RESIDENT', date(2010, 1, 1), date(2017, 6, 30))
        declare(self.account, 'FOREIGN', date(2017, 7, 1), i1=True)
        self.assertEqual(
            cgt.deemed_reset_dates(self.account),
            [(date(2027, 7, 1), 'CUTOVER_2027')])

    def test_arriving_in_australia_adds_one(self):
        """Becoming a resident adds a reset date (s855-45)."""
        declare(self.account, 'FOREIGN', date(2010, 1, 1), date(2016, 12, 31))
        declare(self.account, 'RESIDENT', date(2017, 1, 1))
        self.assertIn((date(2017, 1, 1), 'ARRIVAL'), cgt.deemed_reset_dates(self.account))


class CapitalGainScheduleTests(TransactionTestCase):
    """s102-5: netting, loss ordering, and discounting after losses."""

    def setUp(self):
        self.data = create_golden_master_portfolio()
        self.account = self.data['account']

    def test_losses_reduce_gains_rather_than_each_disposal_being_floored(self):
        """Losses net against gains across the year, not per disposal."""
        schedule = cgt.build_schedule(self.account, 'FY2024/25')
        self.assertEqual(schedule.gross_gains.amount, Decimal('0'))
        self.assertEqual(schedule.gross_losses.amount, Decimal('1577.5054'))
        self.assertEqual(schedule.net_capital_gain.amount, Decimal('0'))
        # Nothing to absorb them, so the whole amount is carried forward.
        self.assertEqual(schedule.losses_carried_forward.amount, Decimal('1577.5054'))

    def test_the_discount_is_applied_after_losses_not_before(self):
        """Prior-year losses are applied before the discount."""
        schedule = cgt.build_schedule(
            self.account, 'FY2023/24', prior_year_losses=Money(Decimal('1000'), 'AUD'))
        self.assertEqual(schedule.prior_year_losses_applied.amount, Decimal('1000'))
        self.assertEqual(schedule.net_capital_gain.amount, Decimal('1707.44080'))

    def test_a_year_with_a_gain_and_no_losses_is_simply_discounted(self):
        schedule = cgt.build_schedule(self.account, 'FY2023/24')
        self.assertEqual(schedule.gross_gains.amount, Decimal('4414.8816'))
        self.assertEqual(schedule.net_capital_gain.amount, Decimal('2207.44080'))

    def test_prior_year_losses_come_from_the_carry_forward_model(self):
        fiscal_year = FiscalYear.objects.first()
        CapitalLossCarryForward.objects.create(
            account=self.account, fiscal_year=fiscal_year,
            amount=Money(Decimal('500'), 'AUD'), is_opening_balance=True,
        )
        schedule = cgt.build_schedule(self.account, 'FY2023/24')
        self.assertEqual(schedule.prior_year_losses_applied.amount, Decimal('500'))

    def test_an_opening_balance_is_the_only_way_to_bring_in_earlier_losses(self):
        """A carried-forward loss must be recorded as a positive amount."""
        loss = CapitalLossCarryForward(
            account=self.account,
            fiscal_year=FiscalYear.objects.first(),
            amount=Money(Decimal('-500'), 'AUD'),
        )
        with self.assertRaises(ValidationError):
            loss.full_clean()

    def test_an_undeclared_account_produces_a_draft_not_a_schedule(self):
        schedule = cgt.build_schedule(self.account, 'FY2023/24')
        self.assertTrue(schedule.is_draft)
        joined = ' '.join(schedule.warnings)
        self.assertIn('Residency has not been declared', joined)
        self.assertIn('does not say who owns this portfolio', joined)

    def test_declaring_everything_clears_the_draft_warnings(self):
        self.account.taxpayer_type = 'INDIVIDUAL'
        self.account.save()
        declare(self.account, 'RESIDENT', date(2000, 1, 1))
        self.data['instrument'].legal_form = 'COMPANY'
        self.data['instrument'].legal_form_source = 'USER'
        self.data['instrument'].save()
        self.data['instrument'].market.country = 'AU'
        self.data['instrument'].market.save()

        schedule = cgt.build_schedule(self.account, 'FY2023/24')
        self.assertEqual(schedule.warnings, [])
        self.assertFalse(schedule.is_draft)

    def test_a_statement_disagreeing_on_cost_base_makes_the_schedule_a_draft(self):
        """A cost base disagreement makes the schedule a draft, even with no attributed gain."""
        self.account.taxpayer_type = 'INDIVIDUAL'
        self.account.save()
        declare(self.account, 'RESIDENT', date(2000, 1, 1))

        instrument = self.data['instrument']
        adjustment = CostBaseAdjustment.objects.create(
            account=self.account, instrument=instrument,
            financial_year_end_date=date(2024, 6, 30),
            cost_base_increase=Money(Decimal('100.00'), 'AUD'))
        statement = AttributionStatement.objects.create(
            account=self.account, instrument=instrument,
            financial_year_end_date=date(2024, 6, 30),
            cost_base_adjustment=adjustment)
        AttributionComponent.objects.create(
            account=self.account, statement=statement,
            component='COSTBASE_INCREASE', amount=Money(Decimal('251.84'), 'AUD'))

        self.assertEqual(
            [e for e in cgt.attribution_events(self.account) if e.instrument == instrument.name],
            [], 'this statement should attribute no capital gain at all')

        schedule = cgt.build_schedule(self.account, 'FY2023/24')
        joined = ' '.join(schedule.warnings)
        self.assertTrue(schedule.is_draft)
        self.assertIn('disagree with the cost base adjustment', joined)
        self.assertIn('251.84', joined)
        self.assertIn('100.00', joined)

        AttributionComponent.objects.filter(statement=statement).update(
            amount=Money(Decimal('100.00'), 'AUD'))
        cleared = ' '.join(cgt.build_schedule(self.account, 'FY2023/24').warnings)
        self.assertNotIn('disagree with the cost base adjustment', cleared)

    def test_an_override_suppressing_the_deeming_makes_the_schedule_a_draft(self):
        """A TAP override suppressing the deeming makes the schedule a draft."""
        self.account.taxpayer_type = 'INDIVIDUAL'
        self.account.save()
        # Departure after both buys (2022-08-15, 2023-02-20) and before both sells, so the
        # parcels were held at the I1 moment and the deeming is what would catch them.
        declare(self.account, 'RESIDENT', date(2000, 1, 1), date(2023, 6, 30))
        declare(self.account, 'FOREIGN', date(2023, 7, 1), i1=True)

        instrument = self.data['instrument']
        instrument.legal_form = 'COMPANY'
        instrument.legal_form_source = 'USER'
        instrument.is_taxable_australian_property_override = False
        instrument.save()
        instrument.market.country = 'AU'
        instrument.market.save()

        schedule = cgt.build_schedule(self.account, 'FY2023/24')
        self.assertTrue(schedule.is_draft)
        joined = ' '.join(schedule.warnings)
        self.assertIn('the only reason their gains are disregarded', joined)
        self.assertIn(instrument.name, joined)

        instrument.is_taxable_australian_property_override = None
        instrument.save()
        cleared = ' '.join(cgt.build_schedule(self.account, 'FY2023/24').warnings)
        self.assertNotIn('the only reason their gains are disregarded', cleared)


class UnsettledForeignResidentDiscountTests(TransactionTestCase):
    """The Act leaves open whether a discount apportioned under s115-115 survives 2027.

    s115-100(aa) ends the 50% for events from 1 July 2027 and new (f) sets 0% where nothing
    else applies, but s115-105 and s115-115 are unamended. The schedule applies the
    apportioned discount and says it is unsettled.
    """

    def _warnings(self, *periods, regime=False):
        data = create_cutover_portfolio()  # bought 15 January 2020, sold 20 August 2028
        for status, start, end, i1 in periods:
            declare(data['account'], status, start, end, i1=i1)
        if regime:
            enable_2027_regime(self, data['account'])
        return cgt.build_schedule(data['account'], 'FY2028/29').warnings

    def _flagged(self, warnings):
        return [w for w in warnings if 's115-100(f)' in w]

    def test_a_foreign_residents_sale_after_the_cutover_is_flagged(self):
        flagged = self._flagged(self._warnings(
            ('RESIDENT', date(2010, 1, 1), date(2021, 6, 30), None),
            ('FOREIGN', date(2021, 7, 1), None, True)))
        self.assertEqual(len(flagged), 1)
        self.assertIn('CUT', flagged[0])

    def test_it_is_flagged_whether_or_not_the_2027_regime_is_modelled(self):
        self.assertTrue(self._flagged(self._warnings(
            ('RESIDENT', date(2010, 1, 1), date(2021, 6, 30), None),
            ('FOREIGN', date(2021, 7, 1), None, True), regime=True)))

    def test_a_resident_throughout_is_not_flagged(self):
        self.assertFalse(self._flagged(self._warnings(
            ('RESIDENT', date(2010, 1, 1), None, None))))

    def test_a_disregarded_gain_is_not_flagged(self):
        """No election, so the parcel is not TAP and its gain is disregarded (s855-10)."""
        self.assertFalse(self._flagged(self._warnings(
            ('RESIDENT', date(2010, 1, 1), date(2021, 6, 30), None),
            ('FOREIGN', date(2021, 7, 1), None, None))))


class StatutoryLossOrderingTests(TransactionTestCase):
    """s102-5 Step 1: losses go against deferred (discounted) gains first."""

    def setUp(self):
        self.data = create_cutover_portfolio()
        self.account = self.data['account']
        declare(self.account, 'RESIDENT', date(2010, 1, 1))
        enable_2027_regime(self)

        # A second holding, bought and sold entirely after the cutover, at a loss, in the
        # same fiscal year as the first sale.
        losing = create_instrument(
            account=self.account, market=self.data['instrument'].market, name='LOSS')
        Buy.objects.create(
            account=self.account, instrument=losing, date=date(2028, 9, 1),
            quantity=Decimal('1000'), unit_price=Money(Decimal('10.00'), 'AUD'),
            total_brokerage=Money(Decimal('0'), 'AUD'),
        )
        Sell.objects.create(
            account=self.account, instrument=losing, date=date(2029, 1, 15),
            quantity=Decimal('1000'), unit_price=Money(Decimal('8.00'), 'AUD'),
            total_brokerage=Money(Decimal('0'), 'AUD'), strategy='FIFO',
        )

    def _schedule(self):
        year = cgt.disposal_events(self.account)[0].fiscal_year
        return cgt.build_schedule(self.account, year)

    def test_both_categories_are_present(self):
        schedule = self._schedule()
        categories = [line.category for line in schedule.lines]
        self.assertEqual(categories, [
            constants.CGT_GAIN_DEFERRED_NON_RESIDENTIAL,
            constants.CGT_GAIN_NON_RESIDENTIAL,
        ])

    def test_the_loss_hits_the_deferred_gain_first(self):
        """A loss is applied to the deferred gain first (s102-5 Step 1(a))."""
        schedule = self._schedule()
        deferred, non_residential = schedule.lines
        self.assertEqual(
            deferred.current_year_losses_applied.amount, Decimal('2000.00'))
        self.assertEqual(
            non_residential.current_year_losses_applied.amount, Decimal('0'))

    def test_spending_the_loss_the_other_way_would_have_been_worth_more(self):
        """The statutory order costs more tax than the reverse would."""
        schedule = self._schedule()
        as_required = schedule.net_capital_gain.amount

        deferred, non_residential = schedule.lines
        # The same loss taken off the undiscounted gain instead.
        if_chosen = (
            (deferred.gross_gains.amount) * Decimal('0.5')
            + (non_residential.gross_gains.amount - Decimal('2000.00')))
        self.assertGreater(as_required, if_chosen)


class CGTReportTests(TransactionTestCase):
    """The CGT event and schedule reports, and the unchanged realised gain report."""

    def setUp(self):
        self.data = create_golden_master_portfolio()
        self.account = self.data['account']

    def test_the_event_report_emits_every_field_of_the_fact_table(self):
        df = CGTEventReport(account=self.account).generate()
        self.assertEqual(list(df.columns), cgt.event_fields())
        self.assertEqual(len(df), 3)

    def test_the_realised_gain_report_did_not_grow_a_column(self):
        """RealisedCapitalGainReport's columns are unchanged."""
        df = RealisedCapitalGainReport(account=self.account).generate()
        self.assertEqual(list(df.columns), [
            'sell_date', 'instrument', 'quantity_sold', 'buy_id', 'parcel_id', 'sell_id',
            'sell_allocation_id', 'buy_date', 'days_held', 'proceeds', 'cost_base',
            'capital_gain', 'fiscal_year',
        ])

    def test_the_schedule_report_refuses_to_call_itself_final(self):
        report = CGTScheduleReport(account=self.account, fiscal_year='FY2023/24')
        self.assertTrue(report.is_draft)
        self.assertTrue(report.warnings())
        self.assertTrue(report.summary()['is_draft'])

    def test_the_schedule_report_carries_the_figures_a_return_asks_for(self):
        summary = CGTScheduleReport(
            account=self.account, fiscal_year='FY2023/24').summary()
        self.assertEqual(
            summary['total_current_year_capital_gains'].amount, Decimal('4414.8816'))
        self.assertEqual(summary['net_capital_gain'].amount, Decimal('2207.44080'))
        self.assertEqual(
            summary['minimum_tax_capital_gain_base'].amount, Decimal('2207.44080'))


class LoadCPICommandTests(TestCase):
    """The load_cpi command."""

    def _write(self, rows):
        handle = tempfile.NamedTemporaryFile(
            mode='w', suffix='.csv', delete=False, newline='', encoding='utf-8')
        handle.write(rows)
        handle.close()
        self.addCleanup(lambda: Path(handle.name).unlink(missing_ok=True))
        return handle.name

    def test_rows_are_loaded_and_normalised_to_the_quarter_start(self):
        """Loaded CPI rows are normalised to the start of their quarter."""
        path = self._write('date,index\n2027-09-30,100.0\n2027-12-31,102.5\n')
        call_command('load_cpi', path)
        self.assertEqual(
            list(CPIIndex.objects.values_list('quarter_start_date', flat=True)),
            [date(2027, 7, 1), date(2027, 10, 1)])

    def test_loading_twice_updates_rather_than_duplicating(self):
        path = self._write('2027-07-01,100.0\n')
        call_command('load_cpi', path)
        call_command('load_cpi', self._write('2027-07-01,101.0\n'))
        self.assertEqual(CPIIndex.objects.count(), 1)
        self.assertEqual(
            CPIIndex.objects.first().index_number, Decimal('101.0000'))

    def test_a_dry_run_writes_nothing(self):
        path = self._write('2027-07-01,100.0\n')
        call_command('load_cpi', path, '--dry-run')
        self.assertEqual(CPIIndex.objects.count(), 0)

    def test_a_file_with_no_usable_rows_is_an_error(self):
        path = self._write('nothing,useful\n')
        with self.assertRaises(CommandError):
            call_command('load_cpi', path)

    def test_a_quarter_must_start_a_quarter(self):
        entry = CPIIndex(quarter_start_date=date(2027, 8, 1), index_number=Decimal('100'))
        with self.assertRaises(ValidationError):
            entry.full_clean()


class CaptureCutoverValuationsCommandTests(TransactionTestCase):
    """The capture_cutover_valuations command."""

    def setUp(self):
        self.account = create_account()
        self.instrument = create_instrument(account=self.account, name='CUT')
        Buy.objects.create(
            account=self.account, instrument=self.instrument, date=date(2020, 1, 15),
            quantity=Decimal('1000'), unit_price=Money(Decimal('10.00'), 'AUD'),
            total_brokerage=Money(Decimal('19.95'), 'AUD'),
        )
        InstrumentPriceHistory.objects.create(
            account=self.account, instrument=self.instrument, date=date(2027, 6, 30),
            open=Decimal('14'), high=Decimal('16'), low=Decimal('14'),
            close=Decimal('15.5'), volume=1000, stock_splits=Decimal('0'),
        )

    def test_it_records_the_close_for_the_day_before_the_cutover(self):
        """The command values 30 June 2027, the day before the cutover."""
        call_command('capture_cutover_valuations', '--account', self.account.description)
        valuation = InstrumentValuation.objects.get()
        self.assertEqual(valuation.valuation_date, date(2027, 6, 30))
        self.assertEqual(valuation.unit_value, Money(Decimal('15.500000'), 'AUD'))
        self.assertEqual(valuation.source, 'PRICE_HISTORY')

    def test_it_does_not_overwrite_a_value_the_user_sourced(self):
        InstrumentValuation.objects.create(
            account=self.account, instrument=self.instrument,
            valuation_date=date(2027, 6, 30),
            unit_value=Money(Decimal('14.00'), 'AUD'),
            purpose='CUTOVER_2027', source='USER',
        )
        call_command('capture_cutover_valuations', '--account', self.account.description)
        self.assertEqual(
            InstrumentValuation.objects.get().unit_value, Money(Decimal('14.00'), 'AUD'))

    def test_overwrite_replaces_it(self):
        InstrumentValuation.objects.create(
            account=self.account, instrument=self.instrument,
            valuation_date=date(2027, 6, 30),
            unit_value=Money(Decimal('14.00'), 'AUD'),
            purpose='CUTOVER_2027', source='USER',
        )
        call_command(
            'capture_cutover_valuations', '--account', self.account.description,
            '--overwrite')
        self.assertEqual(
            InstrumentValuation.objects.get().unit_value,
            Money(Decimal('15.500000'), 'AUD'))

    def test_a_dry_run_writes_nothing(self):
        call_command(
            'capture_cutover_valuations', '--account', self.account.description,
            '--dry-run')
        self.assertEqual(InstrumentValuation.objects.count(), 0)

    def test_an_unknown_portfolio_is_an_error(self):
        with self.assertRaises(CommandError):
            call_command('capture_cutover_valuations', '--account', 'no such portfolio')


class IndexationNeverDeepensALossTests(TransactionTestCase):
    """s110-55: indexation never creates or deepens a loss."""

    def setUp(self):
        self.account = create_account(
            owner=create_user(username='indexloss'),
            description='Indexation Loss',
            fy_type=create_fiscal_year_type(description='AU Tax Year IL'),
        )
        self.market = create_market(account=self.account)
        declare(self.account, 'RESIDENT', date(2010, 1, 1))
        load_test_cpi()
        enable_2027_regime(self)

    def _sell_at(self, name, unit_price):
        """Buy after the cutover for 10,000 and sell when CPI is up 4.5% (indexed: 10,450)."""
        instrument = create_instrument(
            account=self.account, market=self.market, name=name)
        Buy.objects.create(
            account=self.account, instrument=instrument, date=date(2028, 9, 1),
            quantity=Decimal('1000'), unit_price=Money(Decimal('10.00'), 'AUD'),
            total_brokerage=Money(Decimal('0'), 'AUD'),
        )
        Sell.objects.create(
            account=self.account, instrument=instrument, date=date(2029, 1, 15),
            quantity=Decimal('1000'), unit_price=Money(Decimal(unit_price), 'AUD'),
            total_brokerage=Money(Decimal('0'), 'AUD'), strategy='FIFO',
        )
        return [e for e in cgt.disposal_events(self.account)
                if e.instrument == name][0]

    def test_the_indexation_factor_is_what_this_rests_on(self):
        self.assertEqual(
            cgt.indexation_factor(date(2028, 9, 1), date(2029, 1, 15)),
            Decimal('1.045'))

    def test_a_loss_is_measured_against_the_unindexed_cost_base(self):
        event = self._sell_at('DOWN', '8.00')
        self.assertEqual(event.capital_gain, Money(Decimal('-2000.00'), 'AUD'))
        self.assertEqual(event.cost_base, Money(Decimal('10000.00'), 'AUD'))
        # Indexing it would have reported 2,450: a 450 loss the holder never made.
        self.assertEqual(event.gross_loss, Money(Decimal('2000.00'), 'AUD'))

    def test_proceeds_between_the_two_cost_bases_are_neither_a_gain_nor_a_loss(self):
        """Proceeds between the plain and indexed cost bases give neither a gain nor a loss."""
        event = self._sell_at('FLAT', '10.20')
        self.assertEqual(event.capital_gain, Money(Decimal('0'), 'AUD'))
        self.assertEqual(event.gross_gain, Money(Decimal('0'), 'AUD'))
        self.assertEqual(event.gross_loss, Money(Decimal('0'), 'AUD'))

    def test_a_gain_is_measured_against_the_indexed_cost_base(self):
        event = self._sell_at('UP', '11.00')
        self.assertEqual(event.cost_base, Money(Decimal('10450.000'), 'AUD'))
        self.assertEqual(event.capital_gain, Money(Decimal('550.000'), 'AUD'))
        self.assertEqual(event.discount_percentage, Decimal('0'))


class ExportRoundTripTests(TransactionTestCase):
    """An export loads back into the portfolio it came from, or into an empty database."""

    def _export(self, account):
        export = DataExport.objects.create(account=account)
        export.refresh_from_db()
        self.assertTrue(export.file, 'the export produced no file')
        return Path(export.file.path)

    def _detached_export(self, account):
        """An export copied out of the media folder, as a kept backup would be."""
        source = self._export(account)
        destination = Path(tempfile.mkdtemp()) / source.name
        shutil.copy2(source, destination)
        return destination

    def _wipe(self):
        """Empty every table, as `DataLoader.clear_all_data` does.

        Not by deleting each model: deleting a share split reverses it, and one whose
        parcels were since sold refuses.
        """
        call_command('flush', interactive=False, verbosity=0)
        self.assertEqual(Account.objects.count(), 0)

    def test_an_export_restores_into_an_empty_database(self):
        """An export restores into an empty database, bringing its own user and portfolio."""
        data = create_golden_master_portfolio()
        account = data['account']
        declare(account, 'RESIDENT', date(2000, 1, 1))
        account_id, username = account.id, account.owner.username
        path = self._detached_export(account)

        before = {model.__name__: model.objects.count()
                  for model in (Buy, Sell, Parcel, SellAllocation, Instrument, Market,
                                ResidencyPeriod, AppUser, Account)}

        # Everything goes, exactly as it would in the disaster this is for. The export file
        # lives outside the tables, which is the whole point of it.
        self._wipe()

        loader = loading.DataLoader(input_file=path)

        self.assertEqual(loader.account.id, account_id,
                         'the portfolio must keep the id every other row refers to')
        self.assertEqual(AppUser.objects.get().username, username)
        for name, count in before.items():
            self.assertEqual(
                apps.get_model('share_dinkum_app', name).objects.count(), count,
                f'{name} did not come back with the same number of rows')

    def test_the_optional_tables_restore_with_their_links(self):
        """Losses, valuations and attribution statements come back, still linked, from an export."""
        data = create_golden_master_portfolio()
        account, instrument = data['account'], data['instrument']
        year, _ = account.fiscal_year_type.classify_date(date(2024, 6, 30))
        CapitalLossCarryForward.objects.create(
            account=account, fiscal_year=year, amount=Money(Decimal('1500'), 'AUD'), is_opening_balance=True)
        InstrumentValuation.objects.create(
            account=account, instrument=instrument, valuation_date=date(2027, 6, 30),
            unit_value=Money(Decimal('45'), 'AUD'))
        statement = AttributionStatement.objects.create(
            account=account, instrument=instrument, financial_year_end_date=date(2024, 6, 30),
            cost_base_adjustment=data['adjustment'])
        AttributionComponent.objects.create(
            account=account, statement=statement, component='COSTBASE_INCREASE',
            amount=Money(Decimal('50'), 'AUD'))
        path = self._detached_export(account)

        self._wipe()
        loading.DataLoader(input_file=path)

        loss = CapitalLossCarryForward.objects.get()
        self.assertEqual(loss.fiscal_year.start_year, 2023)
        self.assertTrue(loss.is_opening_balance)
        self.assertEqual(InstrumentValuation.objects.get().unit_value, Money(Decimal('45'), 'AUD'))
        restored = AttributionStatement.objects.get()
        self.assertEqual(restored.cost_base_adjustment_id, data['adjustment'].id)
        self.assertEqual(AttributionComponent.objects.get().statement, restored)

    def test_restoring_does_not_derive_what_the_file_already_holds(self):
        """A restore keeps the file's parcels and does not derive a second set by signal."""
        data = create_golden_master_portfolio()
        account = data['account']
        declare(account, 'RESIDENT', date(2000, 1, 1))
        path = self._detached_export(account)
        parcels_before = Parcel.objects.count()
        self.assertGreater(parcels_before, 0)

        self._wipe()

        loading.DataLoader(input_file=path)

        self.assertEqual(Parcel.objects.count(), parcels_before)
        self.assertEqual(
            Parcel.objects.filter(buy__isnull=False).count(), parcels_before,
            'every parcel should be one the file supplied, not one a signal invented')

    def test_a_column_the_model_no_longer_has_is_ignored(self):
        """A column the model no longer has is ignored on import."""
        account = create_account()
        instrument = create_instrument(account=account)
        df = pd.DataFrame([{
            'id': instrument.id,
            'name': instrument.name,
            'is_taxable_australian_property': True,   # renamed away
            'a_field_that_never_existed': 'x',
        }])

        loader = loading.DataLoader(account=account)
        loader.load_table_to_model(model=Instrument, df=df)

        instrument.refresh_from_db()
        self.assertEqual(instrument.name, instrument.name)

    def test_a_blank_file_cell_loads_as_no_file(self):
        """A blank (NaN) file cell loads as no file."""
        account = create_account()
        instrument = create_instrument(account=account)
        df = pd.DataFrame([{
            'legacy_id': 'B-NAN',
            'instrument__name': instrument.name,
            'date': date(2024, 1, 10),
            'quantity': Decimal('10'),
            'unit_price': Decimal('50'),
            'total_brokerage': Decimal('10'),
            'file': float('nan'),
        }])

        loader = loading.DataLoader(account=account)
        loader.load_table_to_model(model=Buy, df=df)

        buy = Buy.objects.get(legacy_id='B-NAN')
        self.assertFalse(buy.file)

    def test_an_export_reloads_into_its_own_portfolio(self):
        data = create_golden_master_portfolio()
        account = data['account']
        declare(account, 'RESIDENT', date(2000, 1, 1))

        path = self._export(account)
        loading.DataLoader(account=account, input_file=path)

        # Nothing duplicated: every row matched the one it came from.
        self.assertEqual(Buy.objects.filter(account=account).count(), 3)
        self.assertEqual(Sell.objects.filter(account=account).count(), 2)
        self.assertEqual(ResidencyPeriod.objects.filter(account=account).count(), 1)
        self.assertEqual(AppUser.objects.count(), 1)

    def test_reloading_an_export_keeps_the_stored_export(self):
        """The file lists its own DataExport row, without a file; that row is not loaded."""
        account = create_golden_master_portfolio()['account']
        path = self._export(account)

        loading.DataLoader(account=account, input_file=path)

        self.assertTrue(path.exists())
        self.assertEqual(DataExport.objects.filter(account=account).count(), 1)

    def test_an_export_holds_only_its_own_portfolio(self):
        """Another portfolio's Account row would stop the file being restored on its own."""
        account = create_golden_master_portfolio()['account']
        create_account(owner=create_user('other'), description='Other',
                       fy_type=account.fiscal_year_type)

        tables = excelinterface.get_all_tables_in_excel(self._export(account))

        self.assertEqual(list(tables['Account']['id']), [str(account.id)])

    def test_a_restore_recalculates_the_stored_figures(self):
        """Saved as their rows loaded, before the parcels they count were there."""
        data = create_golden_master_portfolio()
        path = self._detached_export(data['account'])
        self._wipe()

        loading.DataLoader(input_file=path)

        instrument = Instrument.objects.get(name='GMT')
        self.assertEqual(instrument.calculated_quantity_held, instrument.quantity_held)
        self.assertGreater(instrument.calculated_quantity_held, 0)
        for parcel in Parcel.objects.filter(deactivation_date__isnull=True):
            self.assertEqual(parcel.calculated_is_sold, parcel.is_sold)

    def test_loading_an_export_into_another_portfolio_points_to_a_template(self):
        """Removing the id column, as once advised, pointed copied rows at the original's."""
        account = create_golden_master_portfolio()['account']
        other = create_account(owner=create_user('other'), description='Other',
                               fy_type=account.fiscal_year_type)

        with self.assertRaisesMessage(ValueError, 'make_import_template'):
            loading.DataLoader(account=other, input_file=self._export(account))

    def test_a_blank_text_column_loads_as_empty_rather_than_failing(self):
        """A blank email loads as an empty string, not NULL."""
        account = create_account()
        generator = excelinterface.ExcelGen(title='Blank email')
        generator.add_table(
            pd.DataFrame([{
                'id': str(account.owner.id),
                'username': account.owner.username,
                'email': None,
                'first_name': None,
                'is_active': True,
            }]),
            table_name='AppUser',
        )
        path = Path(tempfile.mkdtemp()) / 'blank.xlsx'
        generator.save(path)

        loading.DataLoader(account=account, input_file=path)

        account.owner.refresh_from_db()
        self.assertEqual(account.owner.email, '')
        self.assertEqual(account.owner.first_name, '')

    def test_a_genuinely_missing_required_value_still_fails(self):
        """A missing required value, such as a date, still fails."""
        account = create_account()
        market = create_market(account=account)
        create_instrument(account=account, market=market, name='BHP')

        generator = excelinterface.ExcelGen(title='Missing date')
        generator.add_table(
            pd.DataFrame([{
                'legacy_id': 'broken-1',
                'instrument__name': 'BHP',
                'date': None,
                'quantity': Decimal('100'),
                'unit_price': Decimal('40'),
                'unit_price_currency': 'AUD',
                'total_brokerage': Decimal('10'),
                'total_brokerage_currency': 'AUD',
            }]),
            table_name='Buy',
        )
        path = Path(tempfile.mkdtemp()) / 'broken.xlsx'
        generator.save(path)

        with self.assertRaises(Exception):
            loading.DataLoader(account=account, input_file=path)
        self.assertEqual(Buy.objects.filter(account=account).count(), 0)


class UpgradeIsInertUntilDataIsTouchedTests(TransactionTestCase):
    """Reading reports never rewrites a stored figure, so upgrading moves no number."""

    def setUp(self):
        self.data = create_golden_master_portfolio()
        self.account = self.data['account']

    def _allocation_state(self):
        return sorted(
            (str(a.id), a.cost_base_increase.amount)
            for a in CostBaseAdjustmentAllocation.objects.filter(account=self.account)
        )

    def test_generating_every_report_leaves_the_allocations_untouched(self):
        before = self._allocation_state()
        self.assertTrue(before)

        RealisedCapitalGainReport(account=self.account).generate()
        OpenParcelReport(account=self.account).generate()
        CGTEventReport(account=self.account).generate()
        CGTScheduleReport(account=self.account, fiscal_year='FY2023/24').generate()
        cgt.all_events(self.account)

        self.assertEqual(self._allocation_state(), before)

    def test_declaring_residency_does_not_rewrite_anything_either(self):
        before = self._allocation_state()
        declare(self.account, 'RESIDENT', date(2000, 1, 1))
        self.account.taxpayer_type = 'INDIVIDUAL'
        self.account.save()

        CGTEventReport(account=self.account).generate()
        self.assertEqual(self._allocation_state(), before)

    def test_re_saving_an_adjustment_does_not_re_spread_it(self):
        """Saving an adjustment again does not re-allocate it."""
        adjustment = self.data['adjustment']
        CostBaseAdjustmentAllocation.objects.filter(
            account=self.account, cost_base_adjustment=adjustment).delete()

        adjustment.save()

        self.assertEqual(self._allocation_state(), [])


class FreshImportPicksUpTheCorrectedWeightingTests(ImportWorkbookMixin, TransactionTestCase):
    """Loading a file into a new portfolio applies the corrected cost base weighting.

    Adjustments are allocated once, on creation, so upgrading or re-importing into the
    same portfolio does not re-spread them.
    """

    def _workbook(self, path):
        """A workbook with one parcel held all year and one bought two months before year end."""
        generator = excelinterface.ExcelGen(title='Weighting')
        generator.add_table(
            pd.DataFrame([{'code': 'ASX', 'suffix': 'AX'}]), table_name='Market')
        generator.add_table(
            pd.DataFrame([{'name': 'AMT', 'currency': 'AUD', 'market__code': 'ASX'}]),
            table_name='Instrument')
        generator.add_table(
            pd.DataFrame([
                {'legacy_id': 'all-year', 'instrument__name': 'AMT',
                 'date': date(2023, 7, 1), 'quantity': Decimal('1000'),
                 'unit_price': Decimal('10'), 'unit_price_currency': 'AUD',
                 'total_brokerage': Decimal('0'), 'total_brokerage_currency': 'AUD'},
                {'legacy_id': 'late', 'instrument__name': 'AMT',
                 'date': date(2024, 5, 1), 'quantity': Decimal('1000'),
                 'unit_price': Decimal('10'), 'unit_price_currency': 'AUD',
                 'total_brokerage': Decimal('0'), 'total_brokerage_currency': 'AUD'},
            ]),
            table_name='Buy')
        generator.add_table(
            pd.DataFrame([{
                'legacy_id': 'amit-2024', 'instrument__name': 'AMT',
                'financial_year_end_date': date(2024, 6, 30),
                'cost_base_increase': Decimal('100'),
                'cost_base_increase_currency': 'AUD',
                'allocation_method': 'QTY_HELD',
            }]),
            table_name='CostBaseAdjustment')
        generator.save(path)
        return path

    def test_a_fresh_import_weights_by_the_days_actually_held(self):
        account = create_account()
        path = self._workbook(Path(tempfile.mkdtemp()) / 'weighting.xlsx')
        loading.DataLoader(account=account, input_file=path)

        by_legacy = {
            p.buy.legacy_id: p.total_adjustments.amount
            for p in Parcel.objects.filter(account=account)
            if p.remaining_quantity or p.sale_date
        }
        self.assertEqual(len(by_legacy), 2)

        # 366 days against 61, on equal quantities, so roughly six to one -- not one to one,
        # which is what the old weighting gave.
        self.assertGreater(by_legacy['all-year'], by_legacy['late'] * 5)
        self.assertEqual(sum(by_legacy.values()), Decimal('100.00'))

    def test_reloading_the_same_file_does_not_re_spread_it(self):
        """Reloading the same file does not re-allocate adjustments."""
        account = create_account()
        path = self._workbook(Path(tempfile.mkdtemp()) / 'weighting.xlsx')
        loading.DataLoader(account=account, input_file=path)

        before = sorted(
            a.cost_base_increase.amount
            for a in CostBaseAdjustmentAllocation.objects.filter(account=account))

        loading.DataLoader(account=account, input_file=path)

        after = sorted(
            a.cost_base_increase.amount
            for a in CostBaseAdjustmentAllocation.objects.filter(account=account))
        self.assertEqual(after, before)
        self.assertEqual(CostBaseAdjustment.objects.filter(account=account).count(), 1)


class RefreshPricesButtonTests(TransactionTestCase):
    """The dashboard's Refresh prices button."""

    def setUp(self):
        self.data = create_golden_master_portfolio()
        self.account = self.data['account']
        self.user = self.account.owner
        self.user.is_staff = True
        self.user.is_superuser = True
        self.user.save()
        self.client.force_login(self.user)
        self.url = reverse('admin:dashboard_refresh_prices')

    def test_the_button_is_on_the_dashboard(self):
        response = self.client.get(reverse('admin:dashboard'))
        self.assertEqual(response.status_code, 200)
        body = response.content.decode()
        self.assertIn('Refresh prices', body)
        self.assertIn(self.url, body)
        # A state changing action has to be a form, not a link.
        self.assertIn('csrfmiddlewaretoken', body)

    def test_it_triggers_the_same_refresh_the_account_field_did(self):
        with patch.object(
                Account, 'update_all_price_history') as prices, patch.object(
                Account, 'update_all_exchange_rate_history') as rates:
            response = self.client.post(self.url)

        self.assertEqual(response.status_code, 302)
        self.assertEqual(response['Location'], reverse('admin:dashboard'))
        prices.assert_called_once()
        rates.assert_called_once()

    def test_exchange_rates_are_refreshed_before_prices(self):
        """Exchange rates are refreshed before prices."""
        calls = []
        with patch.object(Account, 'update_all_price_history',
                          side_effect=lambda: calls.append('prices')), \
             patch.object(Account, 'update_all_exchange_rate_history',
                          side_effect=lambda: calls.append('rates')):
            self.client.post(self.url)

        self.assertEqual(calls, ['rates', 'prices'])

    def test_the_flag_is_left_clear_afterwards(self):
        """update_price_history is cleared after the refresh."""
        with patch.object(Account, 'update_all_price_history'), \
             patch.object(Account, 'update_all_exchange_rate_history'):
            self.client.post(self.url)

        self.account.refresh_from_db()
        self.assertFalse(self.account.update_price_history)

    def test_a_get_does_not_refresh_anything(self):
        """A GET does not refresh anything."""
        with patch.object(Account, 'update_all_price_history') as prices:
            response = self.client.get(self.url)
        self.assertEqual(response.status_code, 405)
        prices.assert_not_called()

    def test_a_user_who_is_not_staff_is_refused(self):
        """A signed-in user without staff rights is refused."""
        outsider = AppUser.objects.create_user(
            username='outsider', password='x', is_staff=False)
        self.client.force_login(outsider)

        with patch.object(Account, 'update_all_price_history') as prices:
            response = self.client.post(self.url)

        self.assertEqual(response.status_code, 302)
        self.assertIn('login', response['Location'])
        prices.assert_not_called()

    def test_a_provider_failure_says_so_and_changes_nothing(self):
        with patch.object(Account, 'update_all_price_history',
                          side_effect=RuntimeError('yfinance is down')), \
             patch.object(Account, 'update_all_exchange_rate_history'):
            response = self.client.post(self.url, follow=True)

        body = response.content.decode()
        self.assertIn('Could not refresh prices', body)
        self.assertIn('yfinance is down', body)
        self.assertIn('existing figures are unchanged', body.replace('Your e', 'e'))

    def test_the_page_says_how_current_the_prices_are(self):
        InstrumentPriceHistory.objects.create(
            account=self.account, instrument=self.data['instrument'],
            date=date(2026, 8, 28), open=Decimal('5'), high=Decimal('6'),
            low=Decimal('5'), close=Decimal('5.5'), volume=100,
            stock_splits=Decimal('0'),
        )
        response = self.client.get(reverse('admin:dashboard'))
        self.assertContains(response, 'Latest close held')
        self.assertContains(response, 'Aug. 28, 2026')

    def test_a_portfolio_with_no_prices_says_that_instead(self):
        response = self.client.get(reverse('admin:dashboard'))
        self.assertContains(response, 'No prices held yet')


class CaptureCGTSnapshotCommandTests(TransactionTestCase):
    """The capture_cgt_snapshot management command."""

    def setUp(self):
        self.data = create_golden_master_portfolio()
        self.account = self.data['account']

    def test_it_captures_every_year_that_has_a_sale(self):
        call_command('capture_cgt_snapshot', '--account', self.account.description)

        years = sorted(
            s.fiscal_year.name
            for s in CGTReturnSnapshot.objects.filter(account=self.account))
        self.assertEqual(years, ['FY2023/24', 'FY2024/25'])

    def test_the_figures_match_the_report(self):
        call_command(
            'capture_cgt_snapshot', '--account', self.account.description,
            '--fiscal-year', 'FY2023/24')

        snapshot = CGTReturnSnapshot.objects.get(account=self.account)
        self.assertEqual(len(snapshot.rows), 1)
        self.assertEqual(
            Decimal(snapshot.totals['total_capital_gain']), Decimal('4414.8816'))

    def test_it_records_which_basis_produced_the_figures(self):
        """A snapshot records the residency basis of its figures."""
        call_command('capture_cgt_snapshot', '--account', self.account.description)
        self.assertEqual(
            CGTReturnSnapshot.objects.first().basis, cgt.BASIS_LEGACY)

        declare(self.account, 'RESIDENT', date(2000, 1, 1))
        call_command('capture_cgt_snapshot', '--account', self.account.description)
        self.assertEqual(
            CGTReturnSnapshot.objects.filter(basis=cgt.BASIS_DIVISION_115).count(), 2)

    def test_lodged_marks_the_figures_that_were_filed(self):
        call_command(
            'capture_cgt_snapshot', '--account', self.account.description,
            '--fiscal-year', 'FY2023/24', '--lodged')
        self.assertTrue(CGTReturnSnapshot.objects.get(account=self.account).is_lodged)

    def test_a_dry_run_writes_nothing(self):
        call_command(
            'capture_cgt_snapshot', '--account', self.account.description, '--dry-run')
        self.assertEqual(CGTReturnSnapshot.objects.count(), 0)

    def test_an_unknown_fiscal_year_is_an_error(self):
        with self.assertRaises(CommandError):
            call_command(
                'capture_cgt_snapshot', '--account', self.account.description,
                '--fiscal-year', 'FY1999/00')

    def test_the_snapshot_then_feeds_the_basis_change_report(self):
        """A captured snapshot feeds the basis change report."""
        call_command(
            'capture_cgt_snapshot', '--account', self.account.description,
            '--fiscal-year', 'FY2023/24')
        df = CGTBasisChangeReport(
            account=self.account,
            fiscal_year=FiscalYear.objects.get(name='FY2023/24')).generate()
        # Nothing has changed since the capture, so every line reconciles.
        self.assertTrue((df['status'] == 'UNCHANGED').all())


class CaptureSnapshotButtonTests(TransactionTestCase):
    """The dashboard's Take capital gains snapshot button."""

    def setUp(self):
        self.data = create_golden_master_portfolio()
        self.account = self.data['account']
        self.user = self.account.owner
        self.user.is_staff = True
        self.user.is_superuser = True
        self.user.save()
        self.client.force_login(self.user)
        self.url = reverse('admin:dashboard_capture_snapshot')

    def test_the_button_is_on_the_dashboard(self):
        response = self.client.get(reverse('admin:dashboard'))
        self.assertContains(response, 'Take capital gains snapshot')
        self.assertContains(response, self.url)
        self.assertContains(response, 'No snapshot taken yet')

    def test_it_records_every_year_that_has_a_sale(self):
        response = self.client.post(self.url)
        self.assertEqual(response.status_code, 302)

        years = sorted(
            s.fiscal_year.name
            for s in CGTReturnSnapshot.objects.filter(account=self.account))
        self.assertEqual(years, ['FY2023/24', 'FY2024/25'])

    def test_the_figures_come_from_the_report(self):
        self.client.post(self.url)
        snapshot = CGTReturnSnapshot.objects.get(fiscal_year__name='FY2023/24')
        self.assertEqual(len(snapshot.rows), 1)
        self.assertEqual(
            Decimal(snapshot.totals['total_capital_gain']), Decimal('4414.8816'))

    def test_nothing_is_marked_as_lodged(self):
        """Snapshots from the button are not marked lodged."""
        self.client.post(self.url)
        self.assertFalse(
            CGTReturnSnapshot.objects.filter(account=self.account, is_lodged=True).exists())

    def test_a_lodged_snapshot_taken_the_same_day_is_kept(self):
        """The button records the other years, and leaves the lodged figures as filed."""
        lodged = CGTReturnSnapshot.capture(
            account=self.account, fiscal_year=FiscalYear.objects.get(name='FY2023/24'),
            is_lodged=True)
        CGTReturnSnapshotRow.objects.filter(snapshot=lodged).update(
            capital_gain=Decimal('1234'))

        response = self.client.post(self.url, follow=True)

        lodged.refresh_from_db()
        self.assertTrue(lodged.is_lodged)
        self.assertEqual(
            list(lodged.captured_rows.values_list('capital_gain', flat=True)),
            [Decimal('1234')])
        self.assertTrue(CGTReturnSnapshot.objects.filter(
            account=self.account, fiscal_year__name='FY2024/25').exists())
        self.assertContains(response, 'marked as lodged')

    def test_capturing_over_a_lodged_snapshot_is_refused(self):
        fiscal_year = FiscalYear.objects.get(name='FY2023/24')
        CGTReturnSnapshot.capture(account=self.account, fiscal_year=fiscal_year, is_lodged=True)

        with self.assertRaises(LodgedSnapshotError):
            CGTReturnSnapshot.capture(account=self.account, fiscal_year=fiscal_year)

        self.assertTrue(CGTReturnSnapshot.objects.get(fiscal_year=fiscal_year).is_lodged)

    def test_the_basis_is_recorded_with_the_figures(self):
        self.client.post(self.url)
        self.assertTrue(
            CGTReturnSnapshot.objects.filter(basis=cgt.BASIS_LEGACY).exists())

        declare(self.account, 'RESIDENT', date(2000, 1, 1))
        self.client.post(self.url)
        self.assertTrue(
            CGTReturnSnapshot.objects.filter(basis=cgt.BASIS_DIVISION_115).exists())

    def test_it_says_when_the_figures_rest_on_an_assumption(self):
        response = self.client.post(self.url, follow=True)
        self.assertContains(response, 'residency has not been declared')

    def test_the_dashboard_then_says_when_they_were_recorded(self):
        self.client.post(self.url)
        response = self.client.get(reverse('admin:dashboard'))
        self.assertContains(response, 'Last snapshot')

    def test_a_portfolio_with_no_sales_says_so_rather_than_recording_nothing(self):
        empty = create_account(
            owner=create_user(username='nosales'),
            description='Empty',
            fy_type=create_fiscal_year_type(description='AU Tax Year Empty'))
        empty.owner.is_staff = True
        empty.owner.is_superuser = True
        empty.owner.save()
        self.client.force_login(empty.owner)

        response = self.client.post(self.url, follow=True)
        self.assertContains(response, 'no sales yet')
        self.assertEqual(CGTReturnSnapshot.objects.filter(account=empty).count(), 0)

    def test_a_get_records_nothing(self):
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 405)
        self.assertEqual(CGTReturnSnapshot.objects.count(), 0)

    def test_a_user_who_is_not_staff_is_refused(self):
        outsider = AppUser.objects.create_user(
            username='outsider2', password='x', is_staff=False)
        self.client.force_login(outsider)
        response = self.client.post(self.url)
        self.assertEqual(response.status_code, 302)
        self.assertIn('login', response['Location'])
        self.assertEqual(CGTReturnSnapshot.objects.count(), 0)


class SnapshotColumnPrecisionTests(TransactionTestCase):
    """Snapshot rows are stored at column precision."""

    def setUp(self):
        self.data = create_golden_master_portfolio()
        self.account = self.data['account']
        self.fiscal_year = FiscalYear.objects.get(name='FY2024/25')

    def test_stored_figures_are_money_at_four_places(self):
        snapshot = CGTReturnSnapshot.capture(
            account=self.account, fiscal_year=self.fiscal_year)
        for row in snapshot.captured_rows.all():
            for field in ('proceeds', 'cost_base', 'capital_gain'):
                value = getattr(row, field)
                with self.subTest(field=field):
                    self.assertGreaterEqual(value.amount.as_tuple().exponent, -4)

    def test_the_total_is_a_money_amount(self):
        snapshot = CGTReturnSnapshot.capture(
            account=self.account, fiscal_year=self.fiscal_year)
        total = snapshot.totals['total_capital_gain']
        self.assertEqual(Decimal(total), Decimal('-1577.5054'))
        self.assertGreaterEqual(Decimal(total).as_tuple().exponent, -4)


class InexactAllocationPrecisionTests(TransactionTestCase):
    """A partial sale whose quantities do not divide evenly is reported at four places."""

    def setUp(self):
        self.account = create_account(
            owner=create_user(username='inexact'),
            description='Inexact',
            fy_type=create_fiscal_year_type(description='AU Tax Year Inexact'))
        market = create_market(account=self.account)
        instrument = create_instrument(
            account=self.account, market=market, name='VDHG')

        Buy.objects.create(
            account=self.account, instrument=instrument, date=date(2022, 3, 1),
            quantity=Decimal('745'), unit_price=Money(Decimal('53.17'), 'AUD'),
            total_brokerage=Money(Decimal('9.50'), 'AUD'),
        )
        # 135 of 745 is a repeating fraction, which is the whole point of the fixture.
        Sell.objects.create(
            account=self.account, instrument=instrument, date=date(2024, 5, 20),
            quantity=Decimal('135'), unit_price=Money(Decimal('53.00'), 'AUD'),
            total_brokerage=Money(Decimal('9.50'), 'AUD'), strategy='FIFO',
        )

    def test_reported_money_stops_at_four_decimal_places(self):
        for event in cgt.disposal_events(self.account):
            for field in ('net_proceeds', 'cost_base', 'capital_gain',
                          'gross_proceeds', 'buy_consideration'):
                value = getattr(event, field)
                with self.subTest(field=field):
                    self.assertGreaterEqual(
                        value.amount.as_tuple().exponent, -4, f'{field} = {value}')

    def test_proceeds_less_cost_base_equals_the_reported_gain(self):
        """Net proceeds less cost base equals the reported capital gain exactly."""
        for event in cgt.disposal_events(self.account):
            self.assertEqual(
                event.net_proceeds - event.cost_base, event.capital_gain)

    def test_a_snapshot_total_is_a_money_amount(self):
        fiscal_year = FiscalYear.objects.get(name='FY2023/24')
        snapshot = CGTReturnSnapshot.capture(
            account=self.account, fiscal_year=fiscal_year)

        total = snapshot.totals['total_capital_gain']
        self.assertGreaterEqual(Decimal(total).as_tuple().exponent, -4, total)

    def test_a_figure_carrying_old_precision_compares_as_unchanged(self):
        """An old snapshot figure with full Decimal precision compares as unchanged."""
        drifted = Decimal('7154.835632530120481927710843')
        rounded = Decimal('7154.8356')
        self.assertNotEqual(drifted, rounded)
        self.assertEqual(
            CGTBasisChangeReport._as_decimal(drifted),
            CGTBasisChangeReport._as_decimal(rounded),
        )

    def test_a_snapshot_taken_now_reports_no_change(self):
        fiscal_year = FiscalYear.objects.get(name='FY2023/24')
        CGTReturnSnapshot.capture(account=self.account, fiscal_year=fiscal_year)
        df = CGTBasisChangeReport(
            account=self.account, fiscal_year=fiscal_year).generate()
        self.assertTrue((df['status'] == 'UNCHANGED').all())

    def test_a_real_change_is_still_reported(self):
        """A real change is still reported."""
        fiscal_year = FiscalYear.objects.get(name='FY2023/24')
        snapshot = CGTReturnSnapshot.capture(
            account=self.account, fiscal_year=fiscal_year)

        row = snapshot.captured_rows.first()
        row.capital_gain = Money(row.capital_gain.amount + Decimal('0.01'), 'AUD')
        row.save()

        df = CGTBasisChangeReport(
            account=self.account, fiscal_year=fiscal_year).generate()
        self.assertTrue((df['status'] == 'CHANGED').any())


class ExportButtonTests(TransactionTestCase):
    """The dashboard's Export portfolio button returns the file itself."""

    def setUp(self):
        self.data = create_golden_master_portfolio()
        self.account = self.data['account']
        self.user = self.account.owner
        self.user.is_staff = True
        self.user.is_superuser = True
        self.user.save()
        self.client.force_login(self.user)
        self.url = reverse('admin:dashboard_export')

    def _download(self, **post):
        """POST and consume the response so FileResponse releases the file (Windows)."""
        response = self.client.post(self.url, post or None)
        if hasattr(response, 'streaming_content'):
            response.body = b''.join(response.streaming_content)
        return response

    def test_the_button_is_on_the_dashboard(self):
        response = self.client.get(reverse('admin:dashboard'))
        self.assertContains(response, 'Export portfolio')
        self.assertContains(response, self.url)
        self.assertContains(response, 'Never exported')

    def test_it_returns_the_file_itself(self):
        """The response is the xlsx file as an attachment."""
        response = self._download()

        self.assertEqual(response.status_code, 200)
        self.assertIn('attachment', response['Content-Disposition'])
        self.assertIn('.xlsx', response['Content-Disposition'])
        body = response.body
        # An xlsx is a zip, so it starts with the zip magic number.
        self.assertTrue(body.startswith(b'PK'), 'not an xlsx file')
        self.assertGreater(len(body), 5000)

    def test_the_export_is_still_recorded(self):
        self._download()
        export = DataExport.objects.get(account=self.account)
        self.assertTrue(export.file)
        self.assertFalse(export.include_price_history)

    def test_price_history_is_left_out_unless_asked_for(self):
        """Price history is included only when requested."""
        self._download()
        self.assertFalse(DataExport.objects.get(account=self.account).include_price_history)

        DataExport.objects.all().delete()
        self._download(include_price_history='on')
        self.assertTrue(DataExport.objects.get(account=self.account).include_price_history)

    def test_the_file_loads_back_in(self):
        """The exported file loads back in."""
        response = self._download()

        path = Path(tempfile.mkdtemp()) / 'export.xlsx'
        path.write_bytes(response.body)

        loading.DataLoader(account=self.account, input_file=path)
        self.assertEqual(Buy.objects.filter(account=self.account).count(), 3)
        self.assertEqual(Sell.objects.filter(account=self.account).count(), 2)

    def test_the_dashboard_then_says_when(self):
        self._download()
        response = self.client.get(reverse('admin:dashboard'))
        self.assertContains(response, 'Last exported')

    def test_a_failure_says_so_and_keeps_the_user_on_the_dashboard(self):
        with patch.object(
                excelinterface.ExcelGen, 'save', side_effect=RuntimeError('disk full')):
            response = self.client.post(self.url, follow=True)
        self.assertContains(response, 'Could not build the export')
        self.assertContains(response, 'disk full')

    def test_a_get_exports_nothing(self):
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 405)
        self.assertEqual(DataExport.objects.count(), 0)

    def test_a_user_who_is_not_staff_is_refused(self):
        outsider = AppUser.objects.create_user(
            username='outsider3', password='x', is_staff=False)
        self.client.force_login(outsider)
        response = self.client.post(self.url)
        self.assertEqual(response.status_code, 302)
        self.assertIn('login', response['Location'])
        self.assertEqual(DataExport.objects.count(), 0)


class ManualAllocationImportTests(TransactionTestCase):
    """Importing a file that pins its own sell allocations, once and twice."""

    def _workbook(self, path, allocations=None):
        if allocations is None:
            allocations = [{
                'legacy_id': 'A001', 'lookup_legacy_buy': 'B002',
                'lookup_legacy_sell': 'S001', 'quantity': Decimal('60'),
            }]

        generator = excelinterface.ExcelGen(title='Manual allocations')
        generator.add_table(
            pd.DataFrame([{'code': 'ASX', 'suffix': 'AX'}]), table_name='Market')
        generator.add_table(
            pd.DataFrame([{'name': 'BHP', 'currency': 'AUD', 'market__code': 'ASX'}]),
            table_name='Instrument')
        generator.add_table(
            pd.DataFrame([
                {'legacy_id': 'B001', 'instrument__name': 'BHP',
                 'date': date(2022, 7, 1), 'quantity': Decimal('100'),
                 'unit_price': Decimal('10'), 'unit_price_currency': 'AUD',
                 'total_brokerage': Decimal('0'), 'total_brokerage_currency': 'AUD'},
                {'legacy_id': 'B002', 'instrument__name': 'BHP',
                 'date': date(2023, 7, 1), 'quantity': Decimal('100'),
                 'unit_price': Decimal('12'), 'unit_price_currency': 'AUD',
                 'total_brokerage': Decimal('0'), 'total_brokerage_currency': 'AUD'},
            ]),
            table_name='Buy')
        generator.add_table(
            pd.DataFrame([{
                'legacy_id': 'S001', 'instrument__name': 'BHP',
                'date': date(2024, 7, 1), 'quantity': Decimal('100'),
                'unit_price': Decimal('15'), 'unit_price_currency': 'AUD',
                'total_brokerage': Decimal('0'), 'total_brokerage_currency': 'AUD',
                'strategy': 'MANUAL',
            }]),
            table_name='Sell')
        generator.add_table(pd.DataFrame(allocations), table_name='SellAllocation')
        generator.save(path)
        return path

    def setUp(self):
        self.account = create_account()
        self.path = self._workbook(Path(tempfile.mkdtemp()) / 'manual.xlsx')

    def test_a_pinned_allocation_loads_against_the_buy_it_names(self):
        loading.DataLoader(account=self.account, input_file=self.path)

        allocation = SellAllocation.objects.get(account=self.account)
        self.assertEqual(allocation.quantity, Decimal('60'))
        # Against B002, not the older B001 that FIFO would have taken.
        self.assertEqual(allocation.parcel.buy.legacy_id, 'B002')

    def test_loading_the_same_file_twice_updates_rather_than_failing(self):
        """Loading the same file twice updates rather than fails when a holding is fully sold."""
        # The whole of B002, so nothing of it is left over. A partial sale leaves an
        # unsold remnant and the lookup still finds exactly one parcel, which is why this
        # went unnoticed: it only fails once a holding is completely sold.
        path = self._workbook(
            Path(tempfile.mkdtemp()) / 'full.xlsx',
            allocations=[{
                'legacy_id': 'A001', 'lookup_legacy_buy': 'B002',
                'lookup_legacy_sell': 'S001', 'quantity': Decimal('100'),
            }])
        loading.DataLoader(account=self.account, input_file=path)
        loading.DataLoader(account=self.account, input_file=path)

        self.assertEqual(SellAllocation.objects.filter(account=self.account).count(), 1)
        allocation = SellAllocation.objects.get(account=self.account)
        self.assertEqual(allocation.quantity, Decimal('100'))
        self.assertEqual(allocation.parcel.buy.legacy_id, 'B002')

    def test_an_unknown_buy_says_which_one_and_why(self):
        """An unknown buy legacy id gives an error naming it."""
        path = self._workbook(
            Path(tempfile.mkdtemp()) / 'missing.xlsx',
            allocations=[{
                'legacy_id': 'A001', 'lookup_legacy_buy': 'NOPE',
                'lookup_legacy_sell': 'S001', 'quantity': Decimal('60'),
            }])
        with self.assertRaises(Exception) as caught:
            loading.DataLoader(account=self.account, input_file=path)
        message = str(caught.exception)
        self.assertIn('NOPE', message)
        self.assertIn('no buy with that legacy id', message.lower())

    def test_an_over_allocated_buy_says_so(self):
        """Allocations exceeding what a buy holds give an error saying so."""
        path = self._workbook(
            Path(tempfile.mkdtemp()) / 'over.xlsx',
            allocations=[
                {'legacy_id': 'A001', 'lookup_legacy_buy': 'B002',
                 'lookup_legacy_sell': 'S001', 'quantity': Decimal('100')},
                {'legacy_id': 'A002', 'lookup_legacy_buy': 'B002',
                 'lookup_legacy_sell': 'S001', 'quantity': Decimal('100')},
            ])
        with self.assertRaises(Exception) as caught:
            loading.DataLoader(account=self.account, input_file=path)
        self.assertIn('B002', str(caught.exception))


def import_buy(legacy_id, day, quantity='100', price='10', **extra):
    return {
        'legacy_id': legacy_id, 'instrument__name': 'BHP', 'date': day,
        'quantity': Decimal(quantity), 'unit_price': Decimal(price),
        'unit_price_currency': 'AUD', 'total_brokerage': Decimal('0'),
        'total_brokerage_currency': 'AUD', **extra,
    }


def import_sell(legacy_id, day, quantity, strategy='FIFO', price='20', **extra):
    return {
        'legacy_id': legacy_id, 'instrument__name': 'BHP', 'date': day,
        'quantity': Decimal(quantity), 'unit_price': Decimal(price),
        'unit_price_currency': 'AUD', 'total_brokerage': Decimal('0'),
        'total_brokerage_currency': 'AUD', 'strategy': strategy, **extra,
    }


def import_split(legacy_id, day, before='1', after='2'):
    return {
        'legacy_id': legacy_id, 'instrument__name': 'BHP', 'date': day,
        'quantity_before': Decimal(before), 'quantity_after': Decimal(after),
    }


class ImportIntegrityTests(TransactionTestCase):
    """What an import file can and cannot do to a portfolio."""

    def setUp(self):
        self.account = create_account()

    def _file(self, instruments=None, **tables):
        path = Path(tempfile.mkdtemp()) / 'import.xlsx'
        generator = excelinterface.ExcelGen(title='Import')
        generator.add_table(
            pd.DataFrame([{'code': 'ASX', 'suffix': 'AX'}]), table_name='Market')
        generator.add_table(
            pd.DataFrame(instruments or [
                {'name': 'BHP', 'currency': 'AUD', 'market__code': 'ASX'}]),
            table_name='Instrument')
        for name, rows in tables.items():
            generator.add_table(pd.DataFrame(rows), table_name=name)
        generator.save(path)
        return path

    def _load(self, **tables):
        loading.DataLoader(account=self.account, input_file=self._file(**tables))

    def _held(self):
        return Instrument.objects.get(account=self.account, name='BHP').quantity_held

    # --- Date order ---

    def test_a_sale_after_a_split_is_allocated_in_post_split_units(self):
        self._load(
            Buy=[import_buy('B1', date(2020, 1, 10))],
            ShareSplit=[import_split('X1', date(2021, 1, 1))],
            Sell=[import_sell('S1', date(2022, 1, 1), '150')],
        )
        self.assertEqual(self._held(), Decimal('50'))
        self.assertFalse(Sell.with_unallocated_quantity(self.account).exists())

    def test_a_sale_before_a_split_is_allocated_before_it(self):
        self._load(
            Buy=[import_buy('B1', date(2020, 1, 10))],
            ShareSplit=[import_split('X1', date(2021, 1, 1))],
            Sell=[import_sell('S1', date(2020, 6, 1), '50')],
        )
        self.assertEqual(self._held(), Decimal('100'))

    def test_sales_are_allocated_in_date_order_not_row_order(self):
        self._load(
            Buy=[import_buy('B1', date(2020, 1, 10), price='5'),
                 import_buy('B2', date(2021, 1, 10), price='9')],
            Sell=[import_sell('S2', date(2023, 1, 1), '100'),
                  import_sell('S1', date(2022, 1, 1), '100')],
        )
        first = Sell.objects.get(account=self.account, legacy_id='S1')
        self.assertEqual(first.sale_allocation.get().parcel.buy.legacy_id, 'B1')

    def test_a_pinned_allocation_follows_its_sale_across_a_split(self):
        self._load(
            Buy=[import_buy('B1', date(2020, 1, 10))],
            ShareSplit=[import_split('X1', date(2021, 1, 1))],
            Sell=[import_sell('S1', date(2022, 1, 1), '200', strategy='MANUAL')],
            SellAllocation=[{'legacy_id': 'A1', 'lookup_legacy_buy': 'B1',
                             'lookup_legacy_sell': 'S1', 'quantity': Decimal('200')}],
        )
        self.assertEqual(self._held(), Decimal('0'))

    # --- Pinned allocations ---

    def test_pinned_allocations_need_a_manual_sale(self):
        with self.assertRaisesMessage(ValueError, 'MANUAL'):
            self._load(
                Buy=[import_buy('B1', date(2020, 1, 10))],
                Sell=[import_sell('S1', date(2022, 1, 1), '50')],
                SellAllocation=[{'legacy_id': 'A1', 'lookup_legacy_buy': 'B1',
                                 'lookup_legacy_sell': 'S1', 'quantity': Decimal('50')}],
            )
        self.assertFalse(Sell.objects.filter(account=self.account).exists())

    # --- Loading again ---

    def test_a_blank_cell_leaves_the_stored_value_alone(self):
        buys = [import_buy('B1', date(2020, 1, 10), notes=None),
                import_buy('B2', date(2020, 1, 11), notes='from the file')]
        self._load(Buy=buys)
        Buy.objects.filter(legacy_id='B1').update(notes='added in the admin')

        self._load(Buy=buys)

        self.assertEqual(Buy.objects.get(legacy_id='B1').notes, 'added in the admin')

    def test_a_row_without_a_legacy_id_is_refused_once_the_portfolio_has_rows(self):
        """Loading it again would add it a second time rather than update it."""
        buys = [import_buy(None, date(2020, 1, 10))]
        self._load(Buy=buys)

        with self.assertRaisesMessage(ValueError, 'legacy_id'):
            self._load(Buy=buys)
        self.assertEqual(Buy.objects.filter(account=self.account).count(), 1)

    def test_a_legacy_id_used_twice_in_a_sheet_is_refused(self):
        with self.assertRaisesMessage(ValueError, 'B1'):
            self._load(Buy=[import_buy('B1', date(2020, 1, 10)),
                            import_buy('B1', date(2021, 1, 10))])

    def test_a_changed_quantity_on_a_second_load_is_refused(self):
        self._load(Buy=[import_buy('B1', date(2020, 1, 10))])
        with self.assertRaisesMessage(ValueError, 'Delete it and enter it again'):
            self._load(Buy=[import_buy('B1', date(2020, 1, 10), quantity='200')])

    # --- Cell values ---

    def test_a_strategy_may_be_given_by_its_label_or_in_lower_case(self):
        self._load(
            Buy=[import_buy('B1', date(2020, 1, 10))],
            Sell=[import_sell('S1', date(2022, 1, 1), '10', strategy='First-in, First-out'),
                  import_sell('S2', date(2022, 1, 2), '10', strategy='lifo')],
        )
        self.assertEqual(
            dict(Sell.objects.values_list('legacy_id', 'strategy')),
            {'S1': 'FIFO', 'S2': 'LIFO'})

    def test_an_unknown_strategy_is_refused(self):
        with self.assertRaisesMessage(ValueError, 'first in'):
            self._load(
                Buy=[import_buy('B1', date(2020, 1, 10))],
                Sell=[import_sell('S1', date(2022, 1, 1), '10', strategy='first in')],
            )

    def test_a_numeric_legacy_id_does_not_gain_a_decimal_point(self):
        """A blank in the column makes pandas read the rest as floats."""
        self._load(Buy=[import_buy(1001, date(2020, 1, 10)),
                        import_buy(None, date(2020, 1, 11))])
        self.assertTrue(Buy.objects.filter(legacy_id='1001').exists())

    @patch('share_dinkum_app.models.yfinanceinterface.get_exchange_rate',
           return_value=Decimal('1.5'))
    def test_a_blank_currency_is_the_instruments(self, mock_get_rate):
        self._load(
            instruments=[{'name': 'BHP', 'currency': 'USD', 'market__code': 'ASX'}],
            Buy=[import_buy('B1', date(2020, 1, 10),
                            unit_price_currency=None, total_brokerage_currency=None)],
        )
        buy = Buy.objects.get(legacy_id='B1')
        self.assertEqual(str(buy.unit_price.currency), 'USD')
        self.assertEqual(str(buy.total_brokerage.currency), 'USD')

    def test_a_date_typed_as_text_is_read_as_a_date(self):
        self._load(Buy=[import_buy('B1', '2020-01-10')])
        self.assertEqual(Buy.objects.get(legacy_id='B1').date, date(2020, 1, 10))


class LegalFormSourceTests(TransactionTestCase):
    """`Instrument.save()` records who set the legal form."""

    def setUp(self):
        self.account = create_account()
        self.market = create_market(account=self.account)

    def _instrument(self, name='BHP', **kwargs):
        return Instrument.objects.create(
            account=self.account, market=self.market, name=name,
            currency=DEFAULT_CURRENCY, **kwargs)

    def test_an_untouched_instrument_is_not_classified(self):
        instrument = self._instrument()
        self.assertEqual(instrument.legal_form, 'UNKNOWN')
        self.assertFalse(instrument.is_classified)

    def test_setting_the_legal_form_confirms_it(self):
        """Setting the legal form marks it as confirmed."""
        instrument = self._instrument()
        instrument.legal_form = 'COMPANY'
        instrument.save()

        instrument.refresh_from_db()
        self.assertEqual(instrument.legal_form_source, 'USER')
        self.assertTrue(instrument.is_classified)

    def test_creating_with_a_legal_form_counts_as_confirmed(self):
        """Creating an instrument with a legal form marks it confirmed."""
        instrument = self._instrument(name='VAS', legal_form='UNIT_TRUST')
        instrument.refresh_from_db()
        self.assertEqual(instrument.legal_form_source, 'USER')
        self.assertTrue(instrument.is_classified)

    def test_a_suggestion_does_not_count_as_confirmed(self):
        """A suggestion saved with source SUGGESTED stays unconfirmed."""
        instrument = self._instrument()
        instrument.legal_form = 'COMPANY'
        instrument.legal_form_source = 'SUGGESTED'
        instrument.save()

        instrument.refresh_from_db()
        self.assertEqual(instrument.legal_form_source, 'SUGGESTED')
        self.assertFalse(instrument.is_classified)

    def test_a_user_correcting_a_suggestion_confirms_it(self):
        """Correcting a suggestion confirms it."""
        instrument = self._instrument()
        instrument.legal_form = 'COMPANY'
        instrument.legal_form_source = 'SUGGESTED'
        instrument.save()

        instrument.legal_form = 'UNIT_TRUST'
        instrument.save()

        instrument.refresh_from_db()
        self.assertEqual(instrument.legal_form_source, 'USER')
        self.assertTrue(instrument.is_classified)

    def test_agreeing_with_a_suggestion_by_re_saving_does_not_confirm_it(self):
        """Re-saving an unchanged suggestion does not confirm it."""
        instrument = self._instrument()
        instrument.legal_form = 'COMPANY'
        instrument.legal_form_source = 'SUGGESTED'
        instrument.save()

        instrument.description = 'BHP Group'
        instrument.save()

        instrument.refresh_from_db()
        self.assertEqual(instrument.legal_form_source, 'SUGGESTED')

    def test_the_suggest_command_still_marks_its_own_work_as_suggested(self):
        instrument = self._instrument(name='VAS')
        Distribution.objects.create(
            account=self.account, instrument=instrument, date=date(2024, 1, 15),
            quantity=Decimal('100'),
            distribution_amount_per_share=Money(Decimal('0.50'), 'AUD'),
        )
        call_command(
            'suggest_instrument_classification', '--account', self.account.description)

        instrument.refresh_from_db()
        self.assertEqual(instrument.legal_form, 'UNIT_TRUST')
        self.assertEqual(instrument.legal_form_source, 'SUGGESTED')
        self.assertFalse(instrument.is_classified)


class Modelling2027RegimeSettingTests(TransactionTestCase):
    """The 2027 regime is an account setting."""

    def setUp(self):
        self.data = create_cutover_portfolio()
        self.account = self.data['account']
        declare(self.account, 'RESIDENT', date(2010, 1, 1))

    def test_it_is_off_by_default(self):
        self.assertFalse(self.account.model_2027_regime)
        events = cgt.disposal_events(self.account)
        self.assertEqual(len(events), 1, 'the old regime produces one unsplit row')

    def test_ticking_it_models_the_new_regime(self):
        enable_2027_regime(self, self.account)
        events = cgt.disposal_events(self.account)
        self.assertEqual(len(events), 2, 'a straddling disposal splits at the cutover')

    def test_it_is_editable_in_the_admin(self):
        """model_2027_regime is editable in the admin."""
        from django.contrib.admin.sites import AdminSite
        from share_dinkum_app.admin import AccountAdmin

        admin_instance = AccountAdmin(Account, AdminSite())
        form = admin_instance.get_form(None)
        self.assertIn('model_2027_regime', form.base_fields)

    def test_one_portfolio_can_model_ahead_while_another_does_not(self):
        """One portfolio can model the 2027 regime while another does not."""
        other = create_cutover_portfolio(suffix='b')
        declare(other['account'], 'RESIDENT', date(2010, 1, 1))
        enable_2027_regime(self, other['account'])

        self.assertEqual(len(cgt.disposal_events(self.account)), 1)
        self.assertEqual(len(cgt.disposal_events(other['account'])), 2)


class PostCutoverWarningTests(TransactionTestCase):
    """Schedule warnings for years with disposals from 1 July 2027, whichever the setting."""

    def setUp(self):
        self.data = create_cutover_portfolio()
        self.account = self.data['account']
        self.account.taxpayer_type = 'INDIVIDUAL'
        self.account.save()
        declare(self.account, 'RESIDENT', date(2010, 1, 1))
        # Taken from the disposal itself rather than written out, so the test follows the
        # fixture if its sell date ever moves.
        self.year = cgt.disposal_events(self.account)[0].fiscal_year

    def _warnings(self):
        return ' '.join(cgt.build_schedule(self.account, self.year).warnings)

    def test_with_the_setting_off_it_says_how_to_turn_it_on(self):
        joined = self._warnings()
        self.assertIn('being worked out under the rules that applied before it', joined)
        self.assertIn('Tick "Model 2027 regime" on the portfolio', joined)
        self.assertIn('Accounts in the admin', joined)

    def test_with_the_setting_on_it_says_the_figures_are_projections(self):
        enable_2027_regime(self, self.account)
        joined = self._warnings()
        self.assertIn('projections rather than settled amounts', joined)
        self.assertIn('Untick "Model 2027 regime"', joined)

    def test_a_year_before_the_cutover_says_nothing_either_way(self):
        """A year before the cutover gets no 2027 warning either way."""
        enable_2027_regime(self, self.account)
        earlier = cgt.build_schedule(self.account, 'FY2023/24')
        joined = ' '.join(earlier.warnings)
        self.assertNotIn('Model 2027 regime', joined)
        self.assertNotIn('projections rather than settled amounts', joined)


class ScheduleFiscalYearScopeTests(TransactionTestCase):
    """The schedule looks a year up by name within the account's own fiscal year type."""

    def test_another_types_year_of_the_same_name_is_not_matched(self):
        from share_dinkum_app.cgt.schedule import _fiscal_year

        # Created first, so a lookup by name alone would find it. It ends 30 September 2024.
        other_type = create_fiscal_year_type(description='October Year', start_month=10)
        FiscalYear.objects.create(fiscal_year_type=other_type, start_year=2023)
        account = create_account()
        FiscalYear.objects.create(fiscal_year_type=account.fiscal_year_type, start_year=2023)

        year = _fiscal_year(account, 'FY2023/24')

        self.assertEqual(year.fiscal_year_type, account.fiscal_year_type)
        self.assertEqual(year.end_date, date(2024, 6, 30))


class YearInProgressIsADraftTests(TransactionTestCase):
    """A fiscal year that has not ended is a draft."""

    def setUp(self):
        self.data = create_golden_master_portfolio()
        self.account = self.data['account']
        self.account.taxpayer_type = 'INDIVIDUAL'
        self.account.save()
        declare(self.account, 'RESIDENT', date(2000, 1, 1))
        instrument = self.data['instrument']
        instrument.legal_form = 'COMPANY'
        instrument.legal_form_source = 'USER'
        instrument.save()
        instrument.market.country = 'AU'
        instrument.market.save()

    def _year_covering(self, when):
        fiscal_year, _ = self.account.fiscal_year_type.classify_date(when)
        return fiscal_year

    def test_a_finished_year_with_nothing_outstanding_is_final(self):
        schedule = cgt.build_schedule(self.account, 'FY2023/24')
        self.assertEqual(schedule.warnings, [])
        self.assertFalse(schedule.is_draft)

    def test_the_current_year_is_a_draft(self):
        current = self._year_covering(date.today())
        schedule = cgt.build_schedule(self.account, current.name)

        self.assertTrue(schedule.is_draft)
        joined = ' '.join(schedule.warnings)
        self.assertIn('has not ended', joined)
        # Naming the end date, because "not finished" invites "finished when?".
        self.assertIn(f'{current.end_date:%d %B %Y}', joined)

    def test_the_last_day_of_a_year_is_still_open(self):
        """A year is still open on its last day."""
        from share_dinkum_app.cgt.schedule import _year_still_running

        current = self._year_covering(date.today())
        with patch('share_dinkum_app.cgt.schedule.date') as fake:
            fake.today.return_value = current.end_date
            self.assertEqual(_year_still_running(self.account, current.name), current.end_date)

    def test_the_day_after_a_year_ends_it_is_closed(self):
        from share_dinkum_app.cgt.schedule import _year_still_running

        current = self._year_covering(date.today())
        with patch('share_dinkum_app.cgt.schedule.date') as fake:
            fake.today.return_value = current.end_date + timedelta(days=1)
            self.assertIsNone(_year_still_running(self.account, current.name))

    def test_the_all_years_view_says_nothing_about_time(self):
        """The all-years view gets no year-in-progress warning."""
        from share_dinkum_app.cgt.schedule import _year_still_running

        self.assertIsNone(_year_still_running(self.account, None))


class FullBackupTests(TransactionTestCase):
    """The full backup: database and documents, in the shared backup folder."""

    def setUp(self):
        from share_dinkum_app import backup

        self.backup = backup
        self.root = Path(tempfile.mkdtemp())
        self.database = self.root / 'db.sqlite3'
        self.media = self.root / 'media'
        self.destination = self.root / 'backups'

        connection = sqlite3.connect(self.database)
        connection.execute('create table t (a int)')
        connection.execute('insert into t values (42)')
        connection.commit()
        connection.close()

        (self.media / 'AFI').mkdir(parents=True)
        (self.media / 'AFI' / 'statement.pdf').write_bytes(b'%PDF-1.4 not really')

    def test_it_copies_the_database_and_the_documents(self):
        result = self.backup.make_backup(self.database, self.media, self.destination)

        self.assertEqual(result['media_files'], 1)
        self.assertGreater(result['database_bytes'], 0)
        self.assertTrue((result['path'] / 'db.sqlite3').exists())
        self.assertTrue((result['path'] / 'media' / 'AFI' / 'statement.pdf').exists())

    def test_the_copied_database_is_usable(self):
        """The copied database opens and holds the data."""
        result = self.backup.make_backup(self.database, self.media, self.destination)

        connection = sqlite3.connect(result['path'] / 'db.sqlite3')
        self.assertEqual(connection.execute('select a from t').fetchone()[0], 42)
        connection.close()

    def test_nothing_to_back_up_returns_none(self):
        empty = Path(tempfile.mkdtemp())
        self.assertIsNone(self.backup.make_backup(
            empty / 'missing.sqlite3', empty / 'missing', self.destination))

    def test_the_latest_backup_is_found_by_name(self):
        """The latest backup is found by folder name."""
        for name in ('2026-01-01T090000', '2026-09-05T141921', '2026-03-02T120000'):
            (self.destination / 'main' / name).mkdir(parents=True)
        self.assertEqual(
            self.backup.latest_backup(self.destination).name, '2026-09-05T141921')

    def test_no_backups_yet(self):
        self.assertIsNone(self.backup.latest_backup(self.destination))

    def test_backups_go_into_a_named_set(self):
        result = self.backup.make_backup(self.database, self.media, self.destination)
        self.assertEqual(result['path'].parent.name, 'main')

    def test_old_backups_are_pruned(self):
        """Backups beyond the retention limit are pruned."""
        for index in range(7):
            (self.destination / 'main' / f'2026-01-0{index + 1}T090000').mkdir(parents=True)

        result = self.backup.make_backup(
            self.database, self.media, self.destination, keep=5)

        kept = self.backup.list_backups(self.destination)
        self.assertEqual(len(kept), 5)
        self.assertIn(result['path'].name, kept, 'the new one must survive its own pruning')

    def test_backups_in_the_old_flat_layout_are_still_found(self):
        """Backups in the old flat layout are still found."""
        (self.destination / '2026-09-05T141921').mkdir(parents=True)
        latest = self.backup.latest_backup(self.destination)
        self.assertEqual(latest.name, '2026-09-05T141921')
        self.assertEqual(len(self.backup.legacy_backups(self.destination)), 1)

    def test_the_old_flat_layout_is_never_pruned(self):
        """Old-layout backups are never pruned."""
        for index in range(7):
            (self.destination / f'2026-01-0{index + 1}T090000').mkdir(parents=True)

        self.backup.make_backup(self.database, self.media, self.destination, keep=1)

        self.assertEqual(len(self.backup.legacy_backups(self.destination)), 7)


class FullBackupButtonTests(TransactionTestCase):
    """The dashboard's Full backup button."""

    def setUp(self):
        self.data = create_golden_master_portfolio()
        self.account = self.data['account']
        self.user = self.account.owner
        self.user.is_staff = True
        self.user.is_superuser = True
        self.user.save()
        self.client.force_login(self.user)
        self.url = reverse('admin:dashboard_full_backup')
        self.root = Path(tempfile.mkdtemp())

    def test_it_writes_to_disk_rather_than_downloading(self):
        """The backup is written to disk, not returned as a download.

        Only checks a directory is written: the test database is in memory.
        """
        with patch.object(dashboard, '_backup_root', return_value=self.root):
            response = self.client.post(self.url, follow=True)

        self.assertEqual(response.status_code, 200)
        # A file response would carry a Content-Disposition; this one redirects to the page.
        self.assertNotIn('Content-Disposition', response)

        written = list(self.root.iterdir())
        self.assertEqual(len(written), 1, 'one timestamped backup directory')
        self.assertTrue(written[0].is_dir())

    def test_the_message_says_where_it_went(self):
        """The success message says where the backup went."""
        with patch.object(dashboard, '_backup_root', return_value=self.root):
            response = self.client.post(self.url, follow=True)

        body = response.content.decode()
        self.assertIn('Backed up to', body)
        self.assertIn('document(s)', body)

    def test_the_two_data_actions_describe_themselves_differently(self):
        """The export and backup actions describe themselves differently."""
        actions = {a.name: a for a in dashboard.DASHBOARD_ACTIONS}
        export = actions['export'].description
        full = actions['full_backup'].description

        self.assertIn('load back into an empty portfolio', export)
        self.assertIn('does not contain them', export)
        self.assertIn('complete database', full)
        self.assertIn('attached documents', full)
        self.assertNotIn('This is your backup', export)


class CGTScheduleExportTests(TransactionTestCase):
    """The CGT schedule workbook export."""

    def setUp(self):
        self.data = create_golden_master_portfolio()
        self.account = self.data['account']
        self.user = self.account.owner
        self.user.is_staff = True
        self.user.is_superuser = True
        self.user.save()
        self.client.force_login(self.user)
        self.url = reverse('admin:dashboard_export_cgt_schedule')
        declare(self.account, 'RESIDENT', date(2000, 1, 1))

    def _workbook(self):
        path = Path(tempfile.mkdtemp()) / 'schedule.xlsx'
        reports.cgt_schedule_workbook(self.account, path)
        return path

    def test_the_workbook_has_a_sheet_for_each_part_of_the_answer(self):
        tables = set(pd.read_excel(self._workbook(), sheet_name='Index')['table_name'])
        self.assertEqual(
            tables, {'ReturnSchedule', 'Summary', 'ScheduleLines', 'Events', 'Warnings'})

    def test_money_is_written_as_numbers_a_spreadsheet_can_add_up(self):
        """Money is written as numbers."""
        path = self._workbook()
        index = pd.read_excel(path, sheet_name='Index')
        sheet = index.set_index('table_name').loc['Summary', 'sheet_name']
        summary = pd.read_excel(path, sheet_name=f'{int(sheet):02d}')
        self.assertTrue(
            pd.api.types.is_numeric_dtype(summary['net_capital_gain']),
            'net_capital_gain must be numeric, not text')

    def test_a_draft_year_is_exported_and_says_so(self):
        """A draft year is exported, flagged, with its warnings."""
        # Undo the setUp declaration, so the schedule is a draft again.
        ResidencyPeriod.objects.filter(account=self.account).delete()
        self.account.taxpayer_type = 'UNDECLARED'
        self.account.save()

        path = self._workbook()
        index = pd.read_excel(path, sheet_name='Index').set_index('table_name')
        summary = pd.read_excel(path, sheet_name=f"{int(index.loc['Summary', 'sheet_name']):02d}")
        warnings = pd.read_excel(path, sheet_name=f"{int(index.loc['Warnings', 'sheet_name']):02d}")

        self.assertTrue(summary['is_draft'].any(), 'the draft years should be marked')
        self.assertGreater(len(warnings), 0, 'and the reasons should be in the file')

    def test_the_return_schedule_is_laid_out_as_the_form_is(self):
        """The return schedule has a row per form line and a column per year."""
        frame = reports.cgt_return_schedule_frame(self.account, ['FY2023/24'])
        self.assertEqual(list(frame.columns), ['line', 'kind', 'FY2023/24'])

        lines = list(frame['line'])
        self.assertIn('Total current year capital gains', lines)
        self.assertIn('Net capital gain', lines)
        self.assertIn('Capital gains from trusts (including managed funds)', lines)
        for category in choices.CGTAssetCategory.reportable():
            self.assertIn(category.label, lines)

    def test_the_categories_add_up_to_the_stated_total(self):
        """The per-category gains add up to the stated total."""
        frame = reports.cgt_return_schedule_frame(self.account, ['FY2023/24'])
        indexed = frame.set_index(frame['line'].str.strip())

        gains = indexed[indexed.index == 'Capital gain']['FY2023/24'].fillna(0).sum()
        trust = indexed.loc['Capital gains from trusts (including managed funds)', 'FY2023/24']
        unclassified = indexed.loc[
            'Unclassified (not reportable until the holding is classified)', 'FY2023/24']
        total = indexed.loc['Total current year capital gains', 'FY2023/24']

        self.assertGreater(total, 0, 'this fixture should have gains to check')
        self.assertAlmostEqual(
            gains + (trust or 0) + (unclassified or 0), total, places=4)

    def test_a_draft_year_is_marked_in_its_own_column(self):
        """Draft years are marked in their own row."""
        expected = reports.CGTScheduleReport(
            account=self.account, fiscal_year='FY2023/24').is_draft
        frame = reports.cgt_return_schedule_frame(self.account, ['FY2023/24'])
        row = frame[frame['line'] == 'Draft (year not final)'].iloc[0]
        self.assertEqual(row['FY2023/24'], 'yes' if expected else 'no')

    def test_the_workbook_carries_the_return_schedule(self):
        tables = set(pd.read_excel(self._workbook(), sheet_name='Index')['table_name'])
        self.assertIn('ReturnSchedule', tables)

    def test_the_button_returns_a_workbook(self):
        response = self.client.post(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response['Content-Type'])
        self.assertIn('attachment;', response['Content-Disposition'])
        self.assertIn('CGT_Schedule', response['Content-Disposition'])
        self.assertTrue(response.content[:2] == b'PK', 'an xlsx is a zip')

    def test_the_button_is_on_the_dashboard_under_tax(self):
        body = self.client.get(reverse('admin:dashboard')).content.decode()
        self.assertIn('Export Australian CGT report', body)
        self.assertIn(self.url, body)

    def test_no_temporary_file_is_left_behind(self):
        """No temporary file is left behind."""
        before = set(Path(tempfile.gettempdir()).glob('*.xlsx'))
        self.client.post(self.url)
        self.assertEqual(set(Path(tempfile.gettempdir()).glob('*.xlsx')) - before, set())


class DashboardActionRegistryTests(TransactionTestCase):
    """DASHBOARD_ACTIONS drives the URLs, the page and the button order."""

    def setUp(self):
        self.data = create_golden_master_portfolio()
        self.account = self.data['account']
        self.user = self.account.owner
        self.user.is_staff = True
        self.user.is_superuser = True
        self.user.save()
        self.client.force_login(self.user)

    def test_every_declared_action_has_a_live_url(self):
        for action in dashboard.DASHBOARD_ACTIONS:
            url = reverse(f'admin:{action.url_name}')
            # POST, because every action changes something and is registered @require_POST.
            response = self.client.post(url)
            self.assertNotEqual(
                response.status_code, 404,
                f'{action.name} is declared but its URL is not wired up')

    def test_the_dashboard_renders_every_action(self):
        response = self.client.get(reverse('admin:dashboard'))
        body = response.content.decode()
        for action in dashboard.DASHBOARD_ACTIONS:
            self.assertIn(action.label, body)
            self.assertIn(reverse(f'admin:{action.url_name}'), body)
            self.assertIn(action.description, body)

    def test_groups_appear_in_declaration_order(self):
        """Groups appear in declaration order."""
        groups = dashboard.action_groups(self.account)
        self.assertEqual([g['label'] for g in groups], ['Prices', 'Tax', 'Data'])

        # Matched on the heading element, not the bare word: "Data" also appears in the app
        # list and in the data-* attributes, well before the panel.
        body = self.client.get(reverse('admin:dashboard')).content.decode()
        positions = [body.index(f'action-group-heading">{g["label"]}') for g in groups]
        self.assertEqual(positions, sorted(positions))

    def test_an_action_status_is_rendered(self):
        """Each action's status line is rendered."""
        groups = dashboard.action_groups(self.account)
        export = next(a for g in groups for a in g['actions'] if a['name'] == 'export')
        self.assertEqual(export['status_text'], 'Never exported.')
        self.assertContains(self.client.get(reverse('admin:dashboard')), 'Never exported.')

    def test_an_action_option_is_rendered_as_a_checkbox(self):
        body = self.client.get(reverse('admin:dashboard')).content.decode()
        self.assertIn('name="include_price_history"', body)

    def test_the_admin_site_is_ours(self):
        """admin.site is a ShareDinkumAdminSite."""
        from share_dinkum_app.admin_site import ShareDinkumAdminSite

        site = admin.site._wrapped if hasattr(admin.site, '_wrapped') else admin.site
        self.assertIsInstance(site, ShareDinkumAdminSite)

    def test_model_registration_survived_the_swap(self):
        """Models stay registered after the admin site swap."""
        response = self.client.get(reverse('admin:share_dinkum_app_instrument_changelist'))
        self.assertEqual(response.status_code, 200)


class PostSalePriceChaseTests(TransactionTestCase):
    """When a price refresh stops chasing a fully sold instrument."""

    def setUp(self):
        self.account = create_account()
        self.instrument = create_instrument(account=self.account)
        self.buy = Buy.objects.create(
            account=self.account, instrument=self.instrument, date=date(2020, 1, 10),
            quantity=Decimal('100'), unit_price=Money(50, 'AUD'),
            total_brokerage=Money(10, 'AUD'))

    def _sell_everything(self, sell_date, recorded_on=None):
        sell = Sell.objects.create(
            account=self.account, instrument=self.instrument, date=sell_date,
            quantity=Decimal('100'), unit_price=Money(60, 'AUD'),
            total_brokerage=Money(10, 'AUD'))
        if recorded_on is not None:
            # created_at is auto_now_add, so it has to be written past the model to
            # stand in for a disposal that was entered long after it happened.
            Sell.objects.filter(pk=sell.pk).update(created_at=recorded_on)
            sell.refresh_from_db()
        return sell

    def _refreshed(self):
        """How many instruments a price refresh fetches."""
        with patch.object(Instrument, 'update_price_history') as fetch:
            self.account.update_all_price_history()
        return fetch.call_count

    def test_a_recent_sale_is_still_chased(self):
        self._sell_everything(date.today() - timedelta(days=2))
        self.assertEqual(self._refreshed(), 1)

    def test_an_old_sale_with_no_price_since_is_given_up_on(self):
        """An old sale with no price since is given up on."""
        long_ago = timezone.now() - timedelta(days=400)
        self._sell_everything(date.today() - timedelta(days=400), recorded_on=long_ago)
        self.assertEqual(self._refreshed(), 0)

    def test_a_sale_entered_months_late_is_still_chased(self):
        """A sale entered late is chased from when it was entered."""
        self._sell_everything(date.today() - timedelta(days=120))
        self.assertEqual(self._refreshed(), 1)

    def test_a_price_after_the_sale_ends_the_chase(self):
        """A price on or after the sale ends the chase."""
        sell_date = date.today() - timedelta(days=2)
        self._sell_everything(sell_date)
        InstrumentPriceHistory.objects.create(
            account=self.account, instrument=self.instrument, date=sell_date,
            open=Decimal('60'), high=Decimal('60'), low=Decimal('60'),
            close=Decimal('60'), volume=0, stock_splits=Decimal('0'))
        self.assertEqual(self._refreshed(), 0)

    def test_a_holding_still_open_is_always_chased(self):
        self.assertEqual(self._refreshed(), 1)


class InlineFieldBudgetTests(TransactionTestCase):
    """Inlines are dropped when a change page would exceed Django's field limit."""

    def setUp(self):
        from django.contrib.admin.sites import AdminSite
        from share_dinkum_app.admin import GenericModelAdmin

        self.account = create_account()
        self.instrument = create_instrument(account=self.account)
        self.admin = GenericModelAdmin(Instrument, AdminSite())
        for index in range(12):
            Buy.objects.create(
                account=self.account, instrument=self.instrument,
                date=date(2020, 1, 1) + timedelta(days=index),
                quantity=Decimal('10'), unit_price=Money(50, 'AUD'),
                total_brokerage=Money(10, 'AUD'))

    def _request(self):
        from django.test import RequestFactory
        request = RequestFactory().get('/')
        request.user = AppUser.objects.first()
        return request

    def _models(self):
        return {i.model.__name__
                for i in self.admin.get_inline_instances(self._request(), self.instrument)}

    def test_inlines_are_shown_within_the_budget(self):
        self.assertIn('Buy', self._models())

    def test_an_inline_over_the_budget_is_dropped_rather_than_breaking_the_save(self):
        """An inline over the budget is dropped, so the page still saves."""
        self.admin.INLINE_FIELD_BUDGET = 10
        self.assertNotIn('Buy', self._models())

    def test_the_budget_is_measured_in_fields_rather_than_rows(self):
        """The inline budget counts fields, not rows."""
        editable = sum(1 for f in Buy._meta.fields if f.editable)
        self.admin.INLINE_FIELD_BUDGET = editable * 12
        self.assertNotIn('Buy', self._models())

        self.admin.INLINE_FIELD_BUDGET = editable * 13
        self.assertIn('Buy', self._models())


class CostBaseAgreesWithStatementTests(TransactionTestCase):
    """Checking a statement's stated cost base movement against its linked adjustment."""

    def setUp(self):
        self.account = create_account()
        self.instrument = create_instrument(account=self.account)
        self.end = date(2024, 6, 30)

    def _pair(self, recorded, components):
        adjustment = CostBaseAdjustment.objects.create(
            account=self.account, instrument=self.instrument,
            financial_year_end_date=self.end,
            cost_base_increase=Money(Decimal(recorded), 'AUD'))
        statement = AttributionStatement.objects.create(
            account=self.account, instrument=self.instrument,
            financial_year_end_date=self.end, cost_base_adjustment=adjustment)
        for component, amount in components:
            AttributionComponent.objects.create(
                account=self.account, statement=statement, component=component,
                amount=Money(Decimal(amount), 'AUD'))
        return statement

    def test_a_matching_shortfall_agrees(self):
        statement = self._pair('251.84', [('COSTBASE_INCREASE', '251.84')])
        self.assertEqual(statement.stated_cost_base_movement, Decimal('251.84'))
        self.assertTrue(statement.cost_base_agrees)

    def test_an_excess_is_a_reduction(self):
        statement = self._pair('-86.98', [('COSTBASE_DECREASE', '86.98')])
        self.assertEqual(statement.stated_cost_base_movement, Decimal('-86.98'))
        self.assertTrue(statement.cost_base_agrees)

    def test_equal_legs_net_to_nil_rather_than_disagreeing(self):
        """Equal AMIT increase and decrease net to nil and agree with a nil adjustment."""
        statement = self._pair('0.00', [('COSTBASE_INCREASE', '1958.03'),
                                        ('COSTBASE_DECREASE', '1958.03')])
        self.assertEqual(statement.stated_cost_base_movement, Decimal('0'))
        self.assertTrue(statement.cost_base_agrees)

    def test_a_pre_amit_tax_deferred_amount_reduces_the_cost_base(self):
        statement = self._pair('-23.30', [('TAX_DEFERRED', '23.30')])
        self.assertEqual(statement.stated_cost_base_movement, Decimal('-23.30'))
        self.assertTrue(statement.cost_base_agrees)

    def test_a_non_attributable_amount_reduces_the_cost_base(self):
        statement = self._pair('-671.54', [('NON_ATTRIBUTABLE', '671.54')])
        self.assertTrue(statement.cost_base_agrees)

    def test_the_amit_pair_takes_precedence_over_other_lines(self):
        """The AMIT net amount takes precedence over other cost base lines."""
        statement = self._pair('122.70', [('COSTBASE_INCREASE', '122.70'),
                                          ('NON_ATTRIBUTABLE', '66.53')])
        self.assertEqual(statement.stated_cost_base_movement, Decimal('122.70'))
        self.assertTrue(statement.cost_base_agrees)

    def test_a_disagreement_is_reported(self):
        statement = self._pair('100.00', [('COSTBASE_INCREASE', '251.84')])
        self.assertFalse(statement.cost_base_agrees)

    def test_nothing_to_compare_is_not_a_pass(self):
        """With nothing to compare, cost_base_agrees is None."""
        no_component = self._pair('100.00', [])
        self.assertIsNone(no_component.stated_cost_base_movement)
        self.assertIsNone(no_component.cost_base_agrees)

        unlinked = AttributionStatement.objects.create(
            account=self.account, instrument=create_instrument(
                account=self.account, market=self.instrument.market, name='OTH'),
            financial_year_end_date=self.end)
        self.assertIsNone(unlinked.cost_base_agrees)


class ReverseOneToOneInlineTests(TransactionTestCase):
    """A missing reverse one-to-one does not break the CostBaseAdjustment change page."""

    def setUp(self):
        from django.contrib.admin.sites import AdminSite
        from share_dinkum_app.admin import GenericModelAdmin

        self.account = create_account()
        self.instrument = create_instrument(account=self.account)
        self.admin = GenericModelAdmin(CostBaseAdjustment, AdminSite())
        self.adjustment = CostBaseAdjustment.objects.create(
            account=self.account, instrument=self.instrument,
            financial_year_end_date=date(2024, 6, 30),
            cost_base_increase=Money(Decimal('100'), 'AUD'),
        )

    def _request(self):
        from django.test import RequestFactory
        request = RequestFactory().get('/')
        request.user = AppUser.objects.first()
        return request

    def test_the_change_page_builds_without_a_linked_statement(self):
        inlines = self.admin.get_inline_instances(self._request(), self.adjustment)
        models_shown = {inline.model for inline in inlines}
        self.assertIn(AttributionStatement, models_shown)

    def test_the_change_page_builds_with_one(self):
        AttributionStatement.objects.create(
            account=self.account, instrument=self.instrument,
            financial_year_end_date=date(2024, 6, 30),
            cost_base_adjustment=self.adjustment,
        )
        inlines = self.admin.get_inline_instances(self._request(), self.adjustment)
        self.assertIn(AttributionStatement, {inline.model for inline in inlines})


class ConfirmLegalFormInAdminTests(TransactionTestCase):
    """The admin's Confirm legal form tick and bulk action."""

    def setUp(self):
        from django.contrib.admin.sites import AdminSite
        from share_dinkum_app.admin import InstrumentAdmin

        self.account = create_account()
        self.market = create_market(account=self.account)
        self.admin = InstrumentAdmin(Instrument, AdminSite())
        self.user = AppUser.objects.first()

    def _suggested(self, name='VAS', legal_form='UNIT_TRUST'):
        instrument = Instrument.objects.create(
            account=self.account, market=self.market, name=name,
            currency=DEFAULT_CURRENCY)
        instrument.legal_form = legal_form
        instrument.legal_form_source = 'SUGGESTED'
        instrument.save()
        return instrument

    def _request(self):
        """A request the admin can attach messages to."""
        from django.contrib.messages.storage.fallback import FallbackStorage
        from django.test import RequestFactory

        request = RequestFactory().post('/')
        request.user = self.user
        request.session = {}
        request._messages = FallbackStorage(request)
        return request

    def test_the_action_confirms_a_suggestion_without_changing_it(self):
        instrument = self._suggested()
        self.admin.confirm_legal_form_action(
            self._request(), Instrument.objects.filter(pk=instrument.pk))

        instrument.refresh_from_db()
        self.assertEqual(instrument.legal_form, 'UNIT_TRUST')
        self.assertEqual(instrument.legal_form_source, 'USER')
        self.assertTrue(instrument.is_classified)

    def test_the_action_confirms_in_bulk(self):
        """The bulk action confirms several instruments at once."""
        names = ['VAS', 'VGS', 'STW']
        for name in names:
            self._suggested(name=name)

        self.admin.confirm_legal_form_action(
            self._request(), Instrument.objects.filter(name__in=names))

        for name in names:
            instrument = Instrument.objects.get(name=name)
            self.assertTrue(instrument.is_classified, f'{name} was not confirmed')

    def test_the_action_will_not_confirm_an_instrument_with_no_legal_form(self):
        """The bulk action does not confirm an instrument with no legal form."""
        instrument = Instrument.objects.create(
            account=self.account, market=self.market, name='???',
            currency=DEFAULT_CURRENCY)

        self.admin.confirm_legal_form_action(
            self._request(), Instrument.objects.filter(pk=instrument.pk))

        instrument.refresh_from_db()
        self.assertEqual(instrument.legal_form, 'UNKNOWN')
        self.assertFalse(instrument.is_classified)

    def test_the_tick_confirms_on_the_change_form(self):
        instrument = self._suggested()
        form = self.admin.get_form(self._request(), instrument)(
            instance=instrument,
            data={'account': self.account.pk, 'market': self.market.pk,
                  'name': 'VAS', 'currency': DEFAULT_CURRENCY,
                  'legal_form': 'UNIT_TRUST', 'confirm_legal_form': 'on'},
        )
        self.assertTrue(form.is_valid(), form.errors)
        self.admin.save_model(self._request(), form.instance, form, change=True)

        instrument.refresh_from_db()
        self.assertEqual(instrument.legal_form_source, 'USER')

    def test_leaving_the_tick_clear_leaves_the_suggestion_alone(self):
        """Saving with the tick clear leaves a suggestion unconfirmed."""
        instrument = self._suggested()
        form = self.admin.get_form(self._request(), instrument)(
            instance=instrument,
            data={'account': self.account.pk, 'market': self.market.pk,
                  'name': 'VAS', 'currency': DEFAULT_CURRENCY,
                  'legal_form': 'UNIT_TRUST', 'description': 'Vanguard Australian Shares'},
        )
        self.assertTrue(form.is_valid(), form.errors)
        self.admin.save_model(self._request(), form.instance, form, change=True)

        instrument.refresh_from_db()
        self.assertEqual(instrument.legal_form_source, 'SUGGESTED')
        self.assertFalse(instrument.is_classified)

    def test_the_tap_override_options_say_what_they_do(self):
        """The TAP override options are labelled by what they do."""
        instrument = self._suggested()
        form = self.admin.get_form(self._request(), instrument)(instance=instrument)
        rendered = str(form['is_taxable_australian_property_override'])

        self.assertIn('Unset', rendered)
        self.assertNotIn('Unknown', rendered)
        self.assertIn('derive it per parcel', rendered)
        self.assertIn('overriding the departure deeming', rendered)

    def test_the_tap_override_still_round_trips_all_three_states(self):
        """The TAP override still round-trips all three states."""
        from share_dinkum_app.admin import UnsetNullBooleanSelect

        widget = UnsetNullBooleanSelect()
        for submitted, expected in [('unknown', None), ('true', True), ('false', False)]:
            self.assertIs(
                widget.value_from_datadict({'f': submitted}, {}, 'f'), expected,
                f'{submitted} did not round-trip')

    def test_the_tick_is_offered_only_where_there_is_something_to_confirm(self):
        suggested = self._suggested()
        self.assertIn('confirm_legal_form', self.admin.get_fields(self._request(), suggested))

        confirmed = self._suggested(name='VGS')
        confirmed.legal_form_source = 'USER'
        confirmed.save()
        self.assertNotIn(
            'confirm_legal_form', self.admin.get_fields(self._request(), confirmed))

        unknown = Instrument.objects.create(
            account=self.account, market=self.market, name='???',
            currency=DEFAULT_CURRENCY)
        self.assertNotIn('confirm_legal_form', self.admin.get_fields(self._request(), unknown))
        self.assertNotIn('confirm_legal_form', self.admin.get_fields(self._request(), None))


class ClassificationClearsTheDraftTests(TransactionTestCase):
    """Confirming legal forms clears the schedule's draft warning, end to end."""

    def test_a_confirmed_instrument_clears_the_draft_warning(self):
        """Confirming the legal form clears the draft warning."""
        data = create_golden_master_portfolio()
        account = data['account']
        account.taxpayer_type = 'INDIVIDUAL'
        account.save()
        declare(account, 'RESIDENT', date(2000, 1, 1))
        data['instrument'].market.country = 'AU'
        data['instrument'].market.save()

        schedule = cgt.build_schedule(account, 'FY2023/24')
        self.assertTrue(any('suggested rather than confirmed' in w or 'not been classified' in w
                            for w in schedule.warnings))

        data['instrument'].legal_form = 'COMPANY'
        data['instrument'].save()

        schedule = cgt.build_schedule(account, 'FY2023/24')
        self.assertEqual(schedule.warnings, [])
        self.assertFalse(schedule.is_draft)


class AssetCategoryOverrideTests(TransactionTestCase):
    """The asset category override accepts only real schedule categories."""

    def setUp(self):
        self.account = create_account()
        self.market = create_market(account=self.account)
        self.market.country = 'AU'
        self.market.save()
        self.instrument = create_instrument(
            account=self.account, market=self.market, name='BHP')
        self.instrument.legal_form = 'COMPANY'
        self.instrument.save()

    def test_the_field_offers_only_the_real_categories(self):
        field = Instrument._meta.get_field('cgt_asset_category_override')
        # The eight real boxes. UNCLASSIFIED is an answer the application reaches on its
        # own when it does not know, not something to be selected as an override.
        self.assertEqual(
            [value for value, _label in field.choices],
            [c.value for c in CGTAssetCategory.reportable()])
        self.assertEqual(len(field.choices), 8)
        self.assertNotIn(CGTAssetCategory.UNCLASSIFIED, dict(field.choices))

    def test_without_an_override_the_category_is_derived(self):
        self.assertEqual(
            cgt.asset_category(self.instrument),
            CGTAssetCategory.AU_LISTED_SHARES)

    def test_a_valid_override_wins(self):
        self.instrument.cgt_asset_category_override = CGTAssetCategory.OTHER_ASSETS
        self.assertEqual(
            cgt.asset_category(self.instrument), CGTAssetCategory.OTHER_ASSETS)

    def test_a_form_refuses_anything_else(self):
        self.instrument.cgt_asset_category_override = 'Shares in Aus listed companys'
        with self.assertRaises(ValidationError):
            self.instrument.full_clean()

    def test_an_invented_category_is_treated_as_unknown_not_printed(self):
        """An invalid category override is treated as unclassified."""
        self.instrument.cgt_asset_category_override = 'Shares in Aus listed companys'
        self.assertEqual(
            cgt.asset_category(self.instrument), CGTAssetCategory.UNCLASSIFIED)

    def test_an_invalid_override_does_not_silently_fall_back_to_the_derivation(self):
        """An invalid override does not fall back to the derived category."""
        self.instrument.cgt_asset_category_override = 'nonsense'
        self.assertNotEqual(
            cgt.asset_category(self.instrument),
            CGTAssetCategory.AU_LISTED_SHARES)

    def test_an_empty_override_is_not_an_override(self):
        for empty in (None, ''):
            with self.subTest(empty=empty):
                self.instrument.cgt_asset_category_override = empty
                self.assertEqual(
                    cgt.asset_category(self.instrument),
                    CGTAssetCategory.AU_LISTED_SHARES)


class VocabularyTests(TransactionTestCase):
    """Fixed vocabularies live in choices.py and are used everywhere."""

    def setUp(self):
        self.account = create_account()

    def test_an_unrecognised_taxpayer_type_is_treated_as_undeclared(self):
        """An unrecognised taxpayer type is treated as undeclared."""
        self.account.taxpayer_type = 'DECEASED_ESTATE'
        self.assertEqual(
            cgt.discount.taxpayer_type_of(self.account), choices.TaxpayerType.UNDECLARED)

    def test_a_missing_rate_raises_rather_than_defaulting(self):
        """`base_rate` raises for a taxpayer type with no rate, rather than defaulting to 50%."""
        from share_dinkum_app.cgt.discount import _RATE_BY_TAXPAYER_TYPE, base_rate
        self.account.taxpayer_type = choices.TaxpayerType.SMSF

        with patch.dict(_RATE_BY_TAXPAYER_TYPE, clear=False):
            del _RATE_BY_TAXPAYER_TYPE[choices.TaxpayerType.SMSF]
            with self.assertRaises(KeyError):
                base_rate(self.account)

        # And with the entry present it is the superannuation rate, not the full discount.
        self.assertEqual(base_rate(self.account), Decimal(1) / Decimal(3))

    def test_every_declared_taxpayer_type_has_a_rate(self):
        """Every taxpayer type has a rate."""
        from share_dinkum_app.cgt.discount import _RATE_BY_TAXPAYER_TYPE
        self.assertEqual(
            set(_RATE_BY_TAXPAYER_TYPE), set(choices.TaxpayerType))

    def test_the_asset_category_stores_a_code_and_reads_as_the_ato_wording(self):
        """The asset category stores a code and displays the ATO wording."""
        self.assertEqual(
            choices.CGTAssetCategory.AU_LISTED_SHARES.value, 'AU_LISTED_SHARES')
        self.assertEqual(
            choices.CGTAssetCategory.AU_LISTED_SHARES.label,
            'Shares in Australian listed companies')

    def test_choice_values_survived_the_refactor_unchanged(self):
        """Stored choice values did not change."""
        self.assertEqual(
            [c.value for c in choices.ResidencyStatus],
            ['RESIDENT', 'FOREIGN', 'TEMPORARY'])
        self.assertEqual(
            [c.value for c in choices.SellStrategy],
            ['FIFO', 'LIFO', 'MIN_CGT', 'MANUAL'])
        self.assertEqual(
            [c.value for c in choices.LegalFormSource],
            ['DEFAULT', 'SUGGESTED', 'USER'])
        self.assertEqual([c.value for c in choices.CGTBasis], ['LEGACY', 'DIVISION_115'])

    def test_the_models_expose_the_shared_vocabulary(self):
        """The models use the shared vocabularies."""
        self.assertEqual(
            [v for v, _ in Account._meta.get_field('taxpayer_type').choices],
            [c.value for c in choices.TaxpayerType])
        self.assertEqual(
            [v for v, _ in ResidencyPeriod._meta.get_field('status').choices],
            [c.value for c in choices.ResidencyStatus])


class VocabularyPortfolioTests(TransactionTestCase):
    """Vocabulary cases that need their own portfolio."""

    def test_an_unrecognised_type_is_flagged_rather_than_silently_discounted(self):
        """An unrecognised taxpayer type gets 50% and a schedule warning."""
        data = create_golden_master_portfolio()
        account = data['account']
        Account.objects.filter(pk=account.pk).update(taxpayer_type='DECEASED_ESTATE')
        account.refresh_from_db()

        schedule = cgt.build_schedule(account, 'FY2023/24')
        self.assertIn(
            'does not say who owns this portfolio', ' '.join(schedule.warnings))

    def test_the_event_report_shows_the_wording_not_the_code(self):
        """The event report shows the ATO wording, not the code."""
        data = create_golden_master_portfolio()
        data['instrument'].legal_form = choices.LegalForm.COMPANY
        data['instrument'].save()
        data['instrument'].market.country = 'AU'
        data['instrument'].market.save()

        df = CGTEventReport(account=data['account']).generate()
        categories = set(str(v) for v in df['asset_category'])
        self.assertIn('Shares in Australian listed companies', categories)
        self.assertNotIn('AU_LISTED_SHARES', categories)


class ConsecutiveAbsenceTests(TransactionTestCase):
    """s104-165(3): back-to-back non-resident periods are one absence for the I1 election."""

    def setUp(self):
        self.account = create_account()
        self.instrument = create_instrument(account=self.account)
        declare(self.account, 'RESIDENT', date(2000, 1, 1), date(2021, 6, 30))

    def _status(self, sold, acquired=date(2015, 1, 1)):
        return cgt.parcel_tap_status(self.account, self.instrument, acquired, sold)

    def test_the_election_carries_into_a_second_period_of_the_same_absence(self):
        declare(self.account, 'FOREIGN', date(2021, 7, 1), date(2023, 12, 31), i1=True)
        declare(self.account, 'TEMPORARY', date(2024, 1, 1))
        self.assertEqual(self._status(date(2025, 1, 1)), cgt.tap.TAP)

    def test_an_election_recorded_on_a_later_period_covers_the_whole_absence(self):
        declare(self.account, 'FOREIGN', date(2021, 7, 1), date(2023, 12, 31))
        declare(self.account, 'FOREIGN', date(2024, 1, 1), i1=True)
        self.assertEqual(self._status(date(2022, 1, 1)), cgt.tap.TAP)

    def test_a_return_to_residency_ends_the_absence(self):
        declare(self.account, 'FOREIGN', date(2021, 7, 1), date(2022, 12, 31), i1=True)
        declare(self.account, 'RESIDENT', date(2023, 1, 1), date(2023, 12, 31))
        declare(self.account, 'FOREIGN', date(2024, 1, 1))
        self.assertEqual(self._status(date(2025, 1, 1)), cgt.tap.NTAP)


class PreDepartureDisposalTests(TransactionTestCase):
    """A sale made before departure is outside the I1 deeming."""

    def setUp(self):
        self.account = create_account()
        self.instrument = create_instrument(account=self.account)
        declare(self.account, 'RESIDENT', date(2000, 1, 1), date(2021, 6, 30))
        declare(self.account, 'FOREIGN', date(2021, 7, 1), i1=True)
        self.declared = cgt.residency.periods(self.account)

    def _status(self, acquired, sold):
        return cgt.parcel_tap_status(
            self.account, self.instrument, acquired, sold, declared=self.declared)

    def test_a_share_sold_before_departure_is_not_deemed_taxable_australian_property(self):
        self.assertEqual(self._status(date(2015, 1, 1), date(2020, 1, 1)), cgt.tap.NTAP)

    def test_the_question_does_not_change_the_outcome_for_a_resident(self):
        """A resident's gain is not disregarded, whatever its TAP status."""
        disregarded, reason = cgt.disregard(
            self.account, cgt.tap.NTAP, date(2020, 1, 1), declared=self.declared)
        self.assertFalse(disregarded)
        self.assertIsNone(reason)

    def test_a_share_still_held_at_departure_is_deemed_taxable_australian_property(self):
        self.assertEqual(self._status(date(2015, 1, 1), date(2026, 1, 1)), cgt.tap.TAP)

    def test_a_share_bought_after_departure_is_not(self):
        self.assertEqual(self._status(date(2022, 1, 1), date(2026, 1, 1)), cgt.tap.NTAP)

    def test_a_sale_on_the_day_of_departure_is_inside_the_deeming(self):
        """A sale on the departure day is inside the deeming."""
        self.assertEqual(self._status(date(2015, 1, 1), date(2021, 7, 1)), cgt.tap.TAP)

    def test_a_sale_the_day_before_departure_is_outside_it(self):
        self.assertEqual(self._status(date(2015, 1, 1), date(2021, 6, 30)), cgt.tap.NTAP)


class UnallocatedSaleWarningTests(TransactionTestCase):
    """Units of a sale no parcel was found for are not in any gain, so the schedule says so."""

    def _oversold(self):
        account = create_account()
        instrument = create_instrument(account=account, name='OVR')
        Buy.objects.create(
            account=account, instrument=instrument, date=date(2023, 1, 10),
            quantity=Decimal('100'), unit_price=Money(5, 'AUD'), total_brokerage=Money(0, 'AUD'),
        )
        sell = Sell.objects.create(
            account=account, instrument=instrument, date=date(2023, 9, 1),
            quantity=Decimal('150'), unit_price=Money(10, 'AUD'),
            total_brokerage=Money(0, 'AUD'), strategy='FIFO',
        )
        return account, sell

    def test_a_sale_larger_than_the_holding_is_listed(self):
        account, sell = self._oversold()

        self.assertEqual(list(Sell.with_unallocated_quantity(account)), [sell])
        warnings = cgt.build_schedule(account, 'FY2023/24').warnings
        self.assertTrue(
            any('not allocated to any parcel' in w and 'OVR' in w and '(50 of 150 units)' in w
                for w in warnings), warnings)

    def test_the_warning_belongs_to_the_year_of_the_sale(self):
        account, _ = self._oversold()
        warnings = cgt.build_schedule(account, 'FY2022/23').warnings
        self.assertFalse(any('not allocated to any parcel' in w for w in warnings))

    def test_a_fully_allocated_sale_is_not_listed(self):
        account = create_golden_master_portfolio()['account']
        self.assertFalse(Sell.with_unallocated_quantity(account).exists())


class CarryForwardYearScopeTests(TransactionTestCase):
    """Carried-forward losses apply only to later years."""

    def setUp(self):
        self.data = create_golden_master_portfolio()
        self.account = self.data['account']
        self.earlier = FiscalYear.objects.get(name='FY2023/24')
        self.later = FiscalYear.objects.get(name='FY2024/25')

    def _record(self, fiscal_year, amount):
        return CapitalLossCarryForward.objects.create(
            account=self.account, fiscal_year=fiscal_year,
            amount=Money(Decimal(amount), 'AUD'), is_opening_balance=True)

    def test_a_loss_is_applied_to_a_later_year(self):
        self._record(self.earlier, '1000')
        schedule = cgt.build_schedule(self.account, self.later)
        # FY2024/25 is all losses in this fixture, so nothing absorbs it and it rolls on.
        self.assertEqual(schedule.prior_year_losses_applied.amount, Decimal('0'))
        self.assertEqual(
            schedule.losses_carried_forward.amount, Decimal('1577.5054') + Decimal('1000'))

    def test_a_loss_from_a_later_year_is_not_applied_to_an_earlier_one(self):
        """A later year's loss is not applied to an earlier year."""
        self._record(self.later, '1000')
        schedule = cgt.build_schedule(self.account, self.earlier)
        self.assertEqual(schedule.prior_year_losses_applied.amount, Decimal('0'))
        self.assertEqual(schedule.net_capital_gain.amount, Decimal('2207.44080'))

    def test_a_loss_from_the_same_year_is_not_double_counted(self):
        """A carry-forward recorded for the same year is not applied to it."""
        self._record(self.earlier, '1000')
        schedule = cgt.build_schedule(self.account, self.earlier)
        self.assertEqual(schedule.prior_year_losses_applied.amount, Decimal('0'))
        self.assertEqual(schedule.net_capital_gain.amount, Decimal('2207.44080'))

    def test_an_earlier_loss_reduces_a_later_gain(self):
        older = FiscalYear.objects.filter(start_year__lt=self.earlier.start_year).first()
        if older is None:
            self.skipTest('no earlier fiscal year in the fixture')
        self._record(older, '1000')
        schedule = cgt.build_schedule(self.account, self.earlier)
        self.assertEqual(schedule.prior_year_losses_applied.amount, Decimal('1000'))
        self.assertEqual(schedule.net_capital_gain.amount, Decimal('1707.44080'))

    def test_an_explicit_override_still_wins(self):
        """An explicit prior_year_losses override is used as given."""
        self._record(self.later, '5000')
        schedule = cgt.build_schedule(
            self.account, self.earlier, prior_year_losses=Money(Decimal('1000'), 'AUD'))
        self.assertEqual(schedule.prior_year_losses_applied.amount, Decimal('1000'))

    def _record_before_the_gain_year(self, amount):
        older = FiscalYear.objects.filter(start_year__lt=self.earlier.start_year).first()
        if older is None:
            self.skipTest('no earlier fiscal year in the fixture')
        self._record(older, amount)

    def test_a_loss_used_up_in_one_year_is_not_applied_again_the_next(self):
        """The FY2023/24 gain absorbs the whole loss, so none of it reaches FY2024/25."""
        self._record_before_the_gain_year('1000')

        later = cgt.build_schedule(self.account, self.later)

        self.assertEqual(later.losses_carried_forward.amount, Decimal('1577.5054'))

    def test_only_what_is_left_of_a_loss_rolls_on(self):
        """A loss larger than the FY2023/24 gain carries only its remainder into FY2024/25."""
        self._record_before_the_gain_year('5000')

        earlier = cgt.build_schedule(self.account, self.earlier)
        later = cgt.build_schedule(self.account, self.later)

        self.assertEqual(earlier.prior_year_losses_applied.amount, Decimal('4414.8816'))
        self.assertEqual(
            later.losses_carried_forward.amount,
            Decimal('1577.5054') + Decimal('5000') - Decimal('4414.8816'))


# =============================================================================
# Error sweep regressions
# =============================================================================


def _superuser_request():
    from django.test import RequestFactory
    request = RequestFactory().post('/')
    request.user = AppUser.objects.create_superuser('sweep', 'sweep@example.com', 'x')
    return request


class AdminSavesEveryModelTests(TransactionTestCase):
    """The admin saved each record twice, the first time passing `user`, which a model with no
    save() of its own rejects."""

    def test_a_cpi_quarter_can_be_added(self):
        entry = CPIIndex(quarter_start_date=date(2027, 7, 1), index_number=Decimal('140.2'))
        admin.site._registry[CPIIndex].save_model(
            _superuser_request(), entry, form=None, change=False)
        self.assertTrue(CPIIndex.objects.filter(quarter_start_date=date(2027, 7, 1)).exists())

    def test_rates_cannot_be_added_by_hand(self):
        """Its account is not editable, so a rate added in the admin could never be saved."""
        self.assertFalse(
            admin.site._registry[ExchangeRate].has_add_permission(_superuser_request()))


@patch('share_dinkum_app.models.yfinanceinterface.get_exchange_rate', return_value=None)
class ExchangeRateCorrectedInAdminTests(TransactionTestCase):
    """A rate corrected by hand is a real rate, and what was converted at it is redone."""

    def _usd_buy(self):
        account = create_account()
        instrument = create_instrument(account=account, name='IBIT', currency='USD')
        return Buy.objects.create(
            account=account, instrument=instrument, date=date(2024, 1, 15),
            quantity=Decimal('10'), unit_price=Money(100, 'USD'), total_brokerage=Money(0, 'USD'),
        )

    def _correct(self, rate, multiplier):
        model_admin = admin.site._registry[ExchangeRate]
        request = _superuser_request()
        form = model_admin.get_form(request, rate)(data={
            'convert_from': 'USD', 'convert_to': 'AUD', 'date': rate.date.isoformat(),
            'exchange_rate_multiplier': multiplier,
        }, instance=rate)
        self.assertTrue(form.is_valid(), form.errors)
        model_admin.save_model(request, form.save(commit=False), form, change=True)

    def test_a_stand_in_corrected_by_hand_stops_being_one(self, mock_get_rate):
        buy = self._usd_buy()
        rate = buy.exchange_rate
        self.assertTrue(rate.is_placeholder)

        self._correct(rate, '1.4')

        rate.refresh_from_db()
        self.assertFalse(rate.is_placeholder)
        parcel = Parcel.objects.get(buy=buy)
        self.assertEqual(parcel.calculated_total_cost_base.amount, Decimal('1400'))

    def test_a_fetched_rate_corrected_by_hand_is_worked_through(self, mock_get_rate):
        mock_get_rate.return_value = Decimal('1.5')
        buy = self._usd_buy()
        self.assertFalse(buy.exchange_rate.is_placeholder)

        self._correct(buy.exchange_rate, '1.6')

        buy.refresh_from_db()
        self.assertEqual(buy.calculated_unit_price_converted.amount, Decimal('160'))
        self.assertEqual(
            Parcel.objects.get(buy=buy).calculated_total_cost_base.amount, Decimal('1600'))


@patch('share_dinkum_app.models.yfinanceinterface.get_exchange_rate', return_value=Decimal('1.5'))
class ForeignSaleEntryTests(TransactionTestCase):
    """A foreign-currency sale is given its rate on save, as a buy is."""

    def _holding(self):
        account = create_account()
        instrument = create_instrument(account=account, name='IBIT', currency='USD')
        Buy.objects.create(
            account=account, instrument=instrument, date=date(2024, 1, 10),
            quantity=Decimal('100'), unit_price=Money(50, 'USD'), total_brokerage=Money(0, 'USD'),
        )
        return account, instrument

    def test_it_validates_without_a_rate_and_is_given_one(self, mock_get_rate):
        account, instrument = self._holding()
        sell = Sell(
            account=account, instrument=instrument, date=date(2025, 3, 10),
            quantity=Decimal('40'), unit_price=Money(60, 'USD'), total_brokerage=Money(10, 'USD'),
            strategy='FIFO',
        )
        sell.full_clean()
        sell.save()
        self.assertEqual(sell.exchange_rate.exchange_rate_multiplier, Decimal('1.5'))

    def test_a_sale_with_no_instrument_is_a_form_error(self, mock_get_rate):
        sell = Sell(
            account=create_account(), date=date(2025, 3, 10), quantity=Decimal('40'),
            unit_price=Money(60, 'AUD'), total_brokerage=Money(0, 'AUD'), strategy='FIFO',
        )
        with self.assertRaises(ValidationError) as caught:
            sell.full_clean()
        self.assertIn('instrument', caught.exception.message_dict)

    def test_a_worthless_foreign_sale_is_still_converted(self, mock_get_rate):
        """Its amounts are all zero, but its proceeds still meet an AUD cost base."""
        account, instrument = self._holding()
        sell = Sell.objects.create(
            account=account, instrument=instrument, date=date(2025, 3, 10),
            quantity=Decimal('100'), unit_price=Money(0, 'USD'), total_brokerage=Money(0, 'USD'),
            strategy='FIFO',
        )
        self.assertIsNotNone(sell.exchange_rate)
        gain = sell.sale_allocation.get().calculated_total_capital_gain
        self.assertEqual(gain, Money(Decimal('-7500'), 'AUD'))


class CaptureValuationOnAGivenDateTests(TransactionTestCase):
    """`--date` values the day given. It used to value 30 June 2027 whatever it was told."""

    def setUp(self):
        self.account = create_account()
        self.instrument = create_instrument(account=self.account, name='CUT')
        Buy.objects.create(
            account=self.account, instrument=self.instrument, date=date(2020, 1, 15),
            quantity=Decimal('1000'), unit_price=Money(10, 'AUD'), total_brokerage=Money(0, 'AUD'),
        )
        for day, close in ((date(2021, 7, 1), '12.25'), (date(2027, 6, 30), '15.5')):
            InstrumentPriceHistory.objects.create(
                account=self.account, instrument=self.instrument, date=day, open=Decimal(close),
                high=Decimal(close), low=Decimal(close), close=Decimal(close), volume=1000,
                stock_splits=Decimal('0'),
            )

    def _capture(self, *args):
        call_command(
            'capture_cutover_valuations', '--account', self.account.description, *args,
            stdout=io.StringIO())

    def test_the_date_given_is_the_date_valued(self):
        self._capture('--date', '2021-07-01')
        valuation = InstrumentValuation.objects.get()
        self.assertEqual(valuation.valuation_date, date(2021, 7, 1))
        self.assertEqual(valuation.unit_value, Money(Decimal('12.25'), 'AUD'))
        self.assertEqual(valuation.purpose, 'OTHER')

    def test_its_purpose_can_be_given(self):
        self._capture('--date', '2021-07-01', '--purpose', 'DEPARTURE')
        self.assertEqual(InstrumentValuation.objects.get().purpose, 'DEPARTURE')


class TradesOnASplitDateTests(TransactionTestCase):
    """A split is dated on its ex-date: trades that day are already in post-split units."""

    SPLIT_DAY = date(2021, 1, 4)

    def setUp(self):
        self.account = create_account()
        self.instrument = create_instrument(account=self.account)
        Buy.objects.create(
            account=self.account, instrument=self.instrument, date=date(2020, 1, 10),
            quantity=Decimal('100'), unit_price=Money(10, 'AUD'), total_brokerage=Money(0, 'AUD'),
        )

    def _split(self):
        return ShareSplit.objects.create(
            account=self.account, instrument=self.instrument, date=self.SPLIT_DAY,
            quantity_before=Decimal('1'), quantity_after=Decimal('2'),
        )

    def _buy_on_the_split_day(self):
        return Buy.objects.create(
            account=self.account, instrument=self.instrument, date=self.SPLIT_DAY,
            quantity=Decimal('50'), unit_price=Money(5, 'AUD'), total_brokerage=Money(0, 'AUD'),
        )

    def _sale_on_the_split_day(self, quantity):
        return Sell(
            account=self.account, instrument=self.instrument, date=self.SPLIT_DAY,
            quantity=Decimal(quantity), unit_price=Money(6, 'AUD'),
            total_brokerage=Money(0, 'AUD'), strategy='FIFO',
        )

    def _held(self):
        return Instrument.objects.get(pk=self.instrument.pk).quantity_held

    def test_a_buy_on_the_split_date_is_not_split(self):
        buy = self._buy_on_the_split_day()
        self._split()
        self.assertEqual(
            Parcel.objects.get(buy=buy, deactivation_date__isnull=True).parcel_quantity,
            Decimal('50'))
        self.assertEqual(self._held(), Decimal('250'))

    def test_a_buy_on_the_split_date_can_be_entered_after_the_split(self):
        self._split()
        self._buy_on_the_split_day()
        self.assertEqual(self._held(), Decimal('250'))

    def test_a_split_entered_after_a_same_day_sale_is_refused(self):
        self._sale_on_the_split_day('40').save()
        split = ShareSplit(
            account=self.account, instrument=self.instrument, date=self.SPLIT_DAY,
            quantity_before=Decimal('1'), quantity_after=Decimal('2'))
        with self.assertRaisesMessage(ValidationError, 'on the same day'):
            split.full_clean()
        with self.assertRaisesMessage(ValueError, 'on the same day'):
            split.save()

    def test_a_sale_on_the_split_date_sells_split_units(self):
        self._split()
        self._sale_on_the_split_day('150').save()
        self.assertEqual(self._held(), Decimal('50'))


class TrustOnlyYearInWorkbookTests(TransactionTestCase):
    """A year whose only capital gains a trust attributed is still in the CGT workbook."""

    def test_the_year_is_included(self):
        account = create_account()
        instrument = create_instrument(account=account, name='VGS')
        statement = AttributionStatement.objects.create(
            account=account, instrument=instrument, financial_year_end_date=date(2025, 6, 30))
        AttributionComponent.objects.create(
            account=account, statement=statement, component='DISCOUNTED_TAP',
            amount=Money(Decimal('100'), 'AUD'))
        path = Path(tempfile.mkdtemp()) / 'schedule.xlsx'
        self.addCleanup(shutil.rmtree, path.parent, ignore_errors=True)

        reports.cgt_schedule_workbook(account, path)

        index = pd.read_excel(path, sheet_name='Index')
        sheet = index.loc[index['table_name'] == 'Summary', 'sheet_name'].iloc[0]
        summary = pd.read_excel(path, sheet_name=str(sheet).zfill(2))
        self.assertEqual(list(summary['fiscal_year']), ['FY2024/25'])
        self.assertEqual(summary['total_current_year_capital_gains'].iloc[0], 200)


def _rate_for(convert_from, convert_to, exchange_date=None):
    """A stand-in for the market: USD 1.5 up to 1 March 2024 and 1.6 after; GBP 2.0."""
    if convert_from == 'GBP':
        return Decimal('2.0')
    if exchange_date is not None and exchange_date <= date(2024, 3, 1):
        return Decimal('1.5')
    return Decimal('1.6')


@patch('share_dinkum_app.models.yfinanceinterface.get_exchange_rate', side_effect=_rate_for)
class ExchangeRateFollowsTheRecordTests(TransactionTestCase):
    """A dividend's date and currency can be corrected, and its rate has to follow."""

    def _dividend(self, account=None, **amounts):
        account = account or create_account()
        instrument = create_instrument(account=account, name='AAPL', currency='USD')
        return Dividend.objects.create(
            account=account, instrument=instrument, date=date(2024, 3, 1),
            quantity=Decimal('100'), dividend_type='FOREIGN', **amounts)

    def test_moving_it_to_another_date_converts_it_at_that_date(self, mock_get_rate):
        dividend = self._dividend(unfranked_amount_per_share=Money(Decimal('0.5'), 'USD'))
        self.assertEqual(dividend.total_dividend_converted, Money(Decimal('75'), 'AUD'))

        dividend.date = date(2024, 4, 2)
        dividend.save()

        dividend.refresh_from_db()
        self.assertEqual(dividend.exchange_rate.date, date(2024, 4, 2))
        self.assertEqual(dividend.total_dividend_converted, Money(Decimal('80'), 'AUD'))

    def test_changing_its_currency_changes_its_rate(self, mock_get_rate):
        dividend = self._dividend(unfranked_amount_per_share=Money(Decimal('0.5'), 'USD'))

        dividend.unfranked_amount_per_share = Money(Decimal('0.5'), 'GBP')
        dividend.save()

        self.assertEqual(str(dividend.exchange_rate.convert_from), 'GBP')
        self.assertEqual(dividend.total_dividend_converted, Money(Decimal('100'), 'AUD'))

    def test_changing_it_to_the_accounts_currency_drops_its_rate(self, mock_get_rate):
        dividend = self._dividend(unfranked_amount_per_share=Money(Decimal('0.5'), 'USD'))

        dividend.unfranked_amount_per_share = Money(Decimal('0.5'), 'AUD')
        dividend.save()

        self.assertIsNone(dividend.exchange_rate)
        self.assertEqual(dividend.total_dividend_converted, Money(Decimal('50'), 'AUD'))

    def test_zero_amounts_left_in_the_default_currency_do_not_decide(self, mock_get_rate):
        """In a US dollar portfolio the unused amounts still default to AUD 0."""
        account = create_account(currency='USD')
        dividend = self._dividend(
            account=account, unfranked_amount_per_share=Money(Decimal('0.5'), 'USD'))

        self.assertIsNone(dividend.exchange_rate)
        self.assertEqual(dividend.total_dividend_converted, Money(Decimal('50'), 'USD'))

    def test_a_rate_chosen_for_a_trade_is_kept(self, mock_get_rate):
        account = create_account()
        instrument = create_instrument(account=account, name='AAPL', currency='USD')
        chosen = create_exchange_rate(
            account, 'USD', 'AUD', rate=Decimal('1.7'), exchange_date=date(2024, 1, 5))
        buy = Buy.objects.create(
            account=account, instrument=instrument, date=date(2024, 1, 10),
            quantity=Decimal('10'), unit_price=Money(100, 'USD'),
            total_brokerage=Money(0, 'USD'), exchange_rate=chosen,
        )

        buy.save()

        buy.refresh_from_db()
        self.assertEqual(buy.exchange_rate, chosen)


class PositiveQuantityTests(TransactionTestCase):
    """Quantities are divided by, so zero failed deep inside a signal and a negative made a
    negative holding. Both are refused with a reason, by the form and on save."""

    def setUp(self):
        self.account = create_account()
        self.instrument = create_instrument(account=self.account)

    def _buy(self, quantity):
        return Buy(
            account=self.account, instrument=self.instrument, date=date(2024, 1, 10),
            quantity=Decimal(quantity), unit_price=Money(10, 'AUD'), total_brokerage=Money(0, 'AUD'))

    def _split(self, before, after):
        return ShareSplit(
            account=self.account, instrument=self.instrument, date=date(2024, 3, 1),
            quantity_before=Decimal(before), quantity_after=Decimal(after))

    def _refused(self, record, field):
        with self.assertRaises(ValidationError) as caught:
            record.full_clean()
        self.assertIn(field, caught.exception.message_dict)
        with self.assertRaisesMessage(ValueError, 'must be more than zero'):
            record.save()

    def test_a_buy_of_nothing_is_refused(self):
        self._refused(self._buy('0'), 'quantity')
        self.assertFalse(Buy.objects.exists())

    def test_a_negative_buy_is_refused(self):
        self._refused(self._buy('-5'), 'quantity')

    def test_a_split_from_nothing_is_refused(self):
        self._refused(self._split('0', '2'), 'quantity_before')

    def test_a_split_to_nothing_is_refused(self):
        self._refused(self._split('1', '0'), 'quantity_after')

    def test_a_negative_dividend_quantity_is_a_form_error(self):
        dividend = Dividend(
            account=self.account, instrument=self.instrument, date=date(2024, 3, 1),
            quantity=Decimal('-1'))
        with self.assertRaises(ValidationError) as caught:
            dividend.full_clean()
        self.assertIn('quantity', caught.exception.message_dict)


class DashboardHoldingsAcrossASplitTests(TransactionTestCase):
    """The charts count a holding in each day's units, so a split is a step on its ex-date."""

    def setUp(self):
        today = date.today()
        self.bought, self.split, self.sold = (
            today - timedelta(days=30), today - timedelta(days=20), today - timedelta(days=10))
        self.account = create_account()
        instrument = create_instrument(account=self.account, name='SPL')
        Buy.objects.create(
            account=self.account, instrument=instrument, date=self.bought,
            quantity=Decimal('100'), unit_price=Money(10, 'AUD'), total_brokerage=Money(0, 'AUD'),
        )
        ShareSplit.objects.create(
            account=self.account, instrument=instrument, date=self.split,
            quantity_before=Decimal('1'), quantity_after=Decimal('2'),
        )
        Sell.objects.create(
            account=self.account, instrument=instrument, date=self.sold,
            quantity=Decimal('150'), unit_price=Money(5, 'AUD'), total_brokerage=Money(0, 'AUD'),
            strategy='FIFO',
        )
        # As traded: 10 a share before the split, 5 after it.
        for day, close in ((self.bought, '10'), (self.split, '5')):
            InstrumentPriceHistory.objects.create(
                account=self.account, instrument=instrument, date=day, open=Decimal(close),
                high=Decimal(close), low=Decimal(close), close=Decimal(close), volume=0,
                stock_splits=Decimal('0'),
            )

    def _context(self):
        from django.test import RequestFactory
        request = RequestFactory().get('/')
        request.user = AppUser.objects.get(pk=self.account.owner_id)
        update = {'current_version': '0', 'latest_version': None, 'release_url': '',
                  'update_available': False}
        with patch.object(dashboard.version, 'check_for_update', return_value=update):
            return dashboard.prepare_dashboard_context(request, {})

    def _on(self, context, series, day):
        labels = context[f'{series}_chart_labels']
        return context[f'{series}_chart_datasets'][0]['data'][labels.index(day.isoformat())]

    def test_the_split_steps_the_holding_up_and_the_sale_leaves_it_positive(self):
        context = self._context()
        self.assertEqual(self._on(context, 'area', self.split - timedelta(days=1)), 100)
        self.assertEqual(self._on(context, 'area', self.split), 200)
        self.assertEqual(self._on(context, 'area', self.sold), 50)
        self.assertGreaterEqual(min(context['area_chart_datasets'][0]['data']), 0)

    def test_the_value_does_not_jump_at_the_split(self):
        context = self._context()
        self.assertEqual(self._on(context, 'value', self.split - timedelta(days=1)), 1000)
        self.assertEqual(self._on(context, 'value', self.split), 1000)


class StoredPricesAreAsTradedTests(TransactionTestCase):
    """A stored day is replaced by a fresh fetch, and valuations read as-traded prices."""

    def setUp(self):
        self.account = create_account()
        self.instrument = create_instrument(account=self.account, name='VAS')
        Buy.objects.create(
            account=self.account, instrument=self.instrument, date=date(2020, 1, 15),
            quantity=Decimal('100'), unit_price=Money(10, 'AUD'), total_brokerage=Money(0, 'AUD'),
        )

    def _store(self, day, close):
        return InstrumentPriceHistory.objects.create(
            account=self.account, instrument=self.instrument, date=day, open=Decimal(close),
            high=Decimal(close), low=Decimal(close), close=Decimal(close), volume=0,
            stock_splits=Decimal('0'),
        )

    def _fetched(self, day, close):
        return pd.DataFrame([{
            'instrument': self.instrument, 'date': day, 'open': close, 'high': close,
            'low': close, 'close': close, 'volume': 10, 'stock_splits': 0.0,
        }])

    def test_a_day_already_stored_is_replaced(self):
        """It used to be kept: a price fetched mid-session, or adjusted, stayed for good."""
        day = date(2027, 6, 30)
        self._store(day, '98.50')
        with patch('share_dinkum_app.models.yfinanceinterface.get_instrument_price_history',
                   return_value=self._fetched(day, 100.25)), \
                patch('share_dinkum_app.models.yfinanceinterface.get_current_price',
                      return_value=None):
            stored = self.instrument.update_price_history(start_date=day, end_date=day)

        self.assertEqual(stored, 1)
        self.assertEqual(
            InstrumentPriceHistory.objects.get(instrument=self.instrument, date=day).close,
            Decimal('100.25'))

    def test_a_value_from_before_a_split_is_brought_into_todays_units(self):
        self._store(date(2027, 6, 30), '20')
        ShareSplit.objects.create(
            account=self.account, instrument=self.instrument, date=date(2028, 1, 10),
            quantity_before=Decimal('1'), quantity_after=Decimal('2'),
        )
        parcel = Parcel.objects.get(deactivation_date__isnull=True)

        value, source = parcel.market_value_at(date(2027, 6, 30))

        self.assertEqual(parcel.parcel_quantity, Decimal('200'))
        self.assertEqual(value, Money(Decimal('2000'), 'AUD'))

    def _refetch(self, fetched):
        output = io.StringIO()
        with patch.object(Instrument, 'update_price_history', return_value=fetched) as update:
            call_command(
                'refetch_price_history', '--account', self.account.description, stdout=output)
        return output.getvalue(), update

    def test_the_refetch_starts_from_the_earliest_stored_day(self):
        self._store(date(2021, 3, 1), '11')
        self._store(date(2024, 3, 1), '12')

        _, update = self._refetch(fetched=800)

        update.assert_called_once_with(start_date=date(2021, 3, 1))

    def test_what_the_provider_no_longer_has_is_kept_and_named(self):
        self._store(date(2021, 3, 1), '11')
        InstrumentValuation.objects.create(
            account=self.account, instrument=self.instrument, valuation_date=date(2021, 3, 1),
            unit_value=Money(Decimal('11'), 'AUD'), source='PRICE_HISTORY')

        output, _ = self._refetch(fetched=0)

        self.assertTrue(InstrumentPriceHistory.objects.filter(instrument=self.instrument).exists())
        self.assertIn('still adjusted', output)
        self.assertIn('VAS', output)
        self.assertIn('copied from a stored price', output)

    def test_a_dry_run_fetches_nothing(self):
        self._store(date(2021, 3, 1), '11')
        output = io.StringIO()
        with patch.object(Instrument, 'update_price_history') as update:
            call_command('refetch_price_history', '--dry-run', stdout=output)
        update.assert_not_called()
        self.assertIn('would fetch from 2021-03-01', output.getvalue())


class UnrecordedLossWarningTests(TransactionTestCase):
    """A loss the application worked out reaches no later year until it is recorded."""

    def setUp(self):
        self.account = create_golden_master_portfolio()['account']
        self.loss_year = FiscalYear.objects.get(name='FY2024/25')  # all losses in the fixture
        self.account.fiscal_year_type.classify_date(date(2025, 8, 1))

    def _warnings(self):
        return ' '.join(cgt.build_schedule(self.account, 'FY2025/26').warnings)

    def test_a_later_year_says_it_is_missing(self):
        warnings = self._warnings()
        self.assertIn('not been recorded as carried forward', warnings)
        self.assertIn('FY2024/25: AUD 1,577.51', warnings)

    def test_recording_it_clears_the_warning(self):
        CapitalLossCarryForward.objects.create(
            account=self.account, fiscal_year=self.loss_year,
            amount=Money(Decimal('1577.51'), 'AUD'))
        self.assertNotIn('FY2024/25:', self._warnings())


class CostBaseAdjustmentAcrossASplitTests(TransactionTestCase):
    """An adjustment is spread by units held, counted in the same units for every parcel.

    Two parcels of 100 units each: one sold on 30 December 2023, after 183 days of FY2023/24,
    and one held all 366. The one held all year takes two thirds of the adjustment, whichever
    side of a 2-for-1 split the sale and the adjustment fall.
    """

    YEAR_END = date(2024, 6, 30)

    def setUp(self):
        self.account = create_account()
        self.instrument = create_instrument(account=self.account, name='SPL')
        self.sold = self._buy(date(2023, 1, 10))
        self.held = self._buy(date(2023, 1, 11))
        Sell.objects.create(
            account=self.account, instrument=self.instrument, date=date(2023, 12, 30),
            quantity=Decimal('100'), unit_price=Money(12, 'AUD'), total_brokerage=Money(0, 'AUD'),
            strategy='FIFO',
        )

    def _buy(self, day):
        return Buy.objects.create(
            account=self.account, instrument=self.instrument, date=day,
            quantity=Decimal('100'), unit_price=Money(10, 'AUD'), total_brokerage=Money(0, 'AUD'),
        )

    def _split(self, day):
        ShareSplit.objects.create(
            account=self.account, instrument=self.instrument, date=day,
            quantity_before=Decimal('1'), quantity_after=Decimal('2'),
        )

    def _allocated(self):
        CostBaseAdjustment.objects.create(
            account=self.account, instrument=self.instrument,
            financial_year_end_date=self.YEAR_END, cost_base_increase=Money(300, 'AUD'),
        )
        return {
            buy: sum(allocation.cost_base_increase.amount
                     for allocation in CostBaseAdjustmentAllocation.objects.filter(
                         parcel__buy=buy, is_active=True))
            for buy in (self.sold, self.held)
        }

    def test_a_split_during_the_year(self):
        self._split(date(2024, 3, 1))
        allocated = self._allocated()
        self.assertEqual(allocated[self.sold], Decimal('100'))
        self.assertEqual(allocated[self.held], Decimal('200'))

    def test_a_split_after_the_year_but_before_the_adjustment_is_entered(self):
        self._split(date(2024, 8, 1))
        allocated = self._allocated()
        self.assertEqual(allocated[self.sold], Decimal('100'))
        self.assertEqual(allocated[self.held], Decimal('200'))


class NegativeCostBaseCheckTests(TransactionTestCase):
    """A cost base taken below zero by decreases is flagged, since E10 is not worked out."""

    def test_it_is_found(self):
        from share_dinkum_app import data_checks

        account = create_account()
        instrument = create_instrument(account=account, name='VAS')
        Buy.objects.create(
            account=account, instrument=instrument, date=date(2023, 1, 10),
            quantity=Decimal('10'), unit_price=Money(1, 'AUD'), total_brokerage=Money(0, 'AUD'),
        )
        self.assertNotIn('negative_cost_base_parcels', {f.key for f in data_checks.run(account)})

        CostBaseAdjustment.objects.create(
            account=account, instrument=instrument, financial_year_end_date=date(2023, 6, 30),
            cost_base_increase=Money(-50, 'AUD'),
        )

        finding = {f.key: f for f in data_checks.run(account)}['negative_cost_base_parcels']
        self.assertEqual(finding.count, 1)
        self.assertTrue(finding.affects_gains)
        self.assertFalse(finding.repairable)


@patch('share_dinkum_app.yfinanceinterface.yf')
class ExchangeRateOnADayWithNoQuoteTests(TestCase):
    """A weekend or holiday takes the last rate before it, as a stand-in's replacement does."""

    def test_the_last_rate_on_or_before_the_day_is_used(self, mock_yf):
        mock_ticker = MagicMock()
        mock_yf.Ticker.return_value = mock_ticker
        mock_ticker.history.return_value = pd.DataFrame(
            {'Close': [1.50, 1.52]},
            index=pd.DatetimeIndex([pd.Timestamp('2024-06-27'), pd.Timestamp('2024-06-28')]))

        # 30 June 2024 was a Sunday.
        rate = yfinanceinterface.get_exchange_rate('USD', 'AUD', exchange_date=date(2024, 6, 30))

        self.assertEqual(rate.quantize(Decimal('0.01')), Decimal('1.52'))
        mock_ticker.history.assert_called_once_with(start='2024-06-23', end='2024-07-01')
